import base64
import hashlib
import json
import os
import shutil
import socket
import struct
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request

from openpyxl import load_workbook


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP = os.path.join(ROOT, "src", "freight_app.py")
DEFAULT_RULES = os.path.join(ROOT, "config", "freight_rules.default.json")
WEBSOCKET_GUID = "258EAFA5-E914-47DA-95CA-C5AB0DC85B11"


def free_port():
    with socket.socket() as server:
        server.bind(("127.0.0.1", 0))
        return server.getsockname()[1]


def event(group_id, message_id, text):
    return {
        "post_type": "message",
        "message_type": "group",
        "group_id": group_id,
        "message_id": message_id,
        "time": time.time(),
        "self_id": 1765614718,
        "user_id": 2000000001,
        "sender": {"nickname": f"模拟用户{group_id}"},
        "message": [{"type": "text", "data": {"text": text}}],
    }


def websocket_text_frame(value):
    payload = json.dumps(value, ensure_ascii=False).encode("utf-8")
    if len(payload) < 126:
        header = bytes([0x81, len(payload)])
    else:
        header = bytes([0x81, 126]) + struct.pack("!H", len(payload))
    return header + payload


class FakeNapCat:
    def __init__(self):
        self.port = free_port()
        self.connection_count = 0
        self._stop = threading.Event()
        self._connection_lock = threading.Lock()
        self._active_connection = None
        self._server = socket.socket()
        self._server.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        self._server.bind(("127.0.0.1", self.port))
        self._server.listen(5)
        self._server.settimeout(0.5)
        self._thread = threading.Thread(target=self._run, daemon=True)

    def start(self):
        self._thread.start()

    def stop(self):
        self._stop.set()
        with self._connection_lock:
            if self._active_connection is not None:
                try:
                    self._active_connection.close()
                except OSError:
                    pass
        self._server.close()
        self._thread.join(timeout=3)

    def send_event(self, value, timeout=5):
        deadline = time.time() + timeout
        frame = websocket_text_frame(value)
        while time.time() < deadline:
            with self._connection_lock:
                connection = self._active_connection
                if connection is not None:
                    try:
                        connection.sendall(frame)
                        return
                    except OSError:
                        pass
            time.sleep(0.05)
        raise TimeoutError("FakeNapCat没有可用连接")

    def _run(self):
        while not self._stop.is_set():
            try:
                connection, _address = self._server.accept()
            except (OSError, socket.timeout):
                continue
            threading.Thread(
                target=self._handle,
                args=(connection,),
                daemon=True,
            ).start()

    def _handle(self, connection):
        connection.settimeout(3)
        request = b""
        try:
            while b"\r\n\r\n" not in request and len(request) < 16384:
                block = connection.recv(4096)
                if not block:
                    return
                request += block
            headers = request.decode("latin-1").split("\r\n")
            key = ""
            for header in headers:
                if header.lower().startswith("sec-websocket-key:"):
                    key = header.split(":", 1)[1].strip()
                    break
            if not key:
                return
            accept = base64.b64encode(hashlib.sha1(
                (key + WEBSOCKET_GUID).encode("ascii")
            ).digest()).decode("ascii")
            connection.sendall((
                "HTTP/1.1 101 Switching Protocols\r\n"
                "Upgrade: websocket\r\n"
                "Connection: Upgrade\r\n"
                f"Sec-WebSocket-Accept: {accept}\r\n\r\n"
            ).encode("ascii"))
            with self._connection_lock:
                self._active_connection = connection
            self.connection_count += 1
            connection_number = self.connection_count
            if connection_number == 1:
                messages = [
                    event(1001, "first-1", "南宁到佛山大板235"),
                    event(1002, "first-2", "南宁到佛山红板236"),
                ]
            else:
                messages = [
                    event(1001, "second-1", "明天南宁到杭州红板240"),
                    event(1002, "second-2", "南宁到黄冈武穴大板160"),
                ]
            for message in messages:
                connection.sendall(websocket_text_frame(message))
                time.sleep(0.05)
            if connection_number == 1:
                time.sleep(0.2)
                return
            connection.settimeout(0.5)
            deadline = time.time() + 15
            while time.time() < deadline and not self._stop.is_set():
                try:
                    if not connection.recv(4096):
                        return
                except socket.timeout:
                    continue
        except OSError:
            pass
        finally:
            with self._connection_lock:
                if self._active_connection is connection:
                    self._active_connection = None
            try:
                connection.close()
            except OSError:
                pass


def request_json(url):
    with urllib.request.urlopen(url, timeout=3) as response:
        return json.loads(response.read().decode("utf-8"))


def request_json_response(url, method="GET", payload=None):
    data = None if payload is None else json.dumps(payload, ensure_ascii=False).encode("utf-8")
    base = url.split("/api/", 1)[0]
    csrf = ""
    if method == "POST":
        with urllib.request.urlopen(base + "/", timeout=3) as response:
            csrf = response.headers.get("X-Freight-CSRF-Token", "")
    request = urllib.request.Request(
        url,
        data=data,
        method=method,
        headers={
            "Content-Type": "application/json",
            "Origin": base,
            "X-Freight-CSRF": csrf,
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=3) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        return exc.code, json.loads(exc.read().decode("utf-8"))


def main():
    fake_napcat = FakeNapCat()
    fake_napcat.start()
    with tempfile.TemporaryDirectory(
        prefix="freight-ws-integration-", ignore_cleanup_errors=True
    ) as directory:
        dashboard_port = free_port()
        rules_path = os.path.join(directory, "freight_rules.json")
        shutil.copy2(DEFAULT_RULES, rules_path)
        config_path = os.path.join(directory, "qq_live_config.json")
        config = {
            "ws_url": f"ws://127.0.0.1:{fake_napcat.port}",
            "access_token": "",
            "group_ids": ["1001", "1002"],
            "group_names": {"1001": "一群", "1002": "二群"},
            "group_default_origins": {"1001": "南宁", "1002": "南宁"},
            "output_root": os.path.join(directory, "data"),
            "rules_file": rules_path,
            "accept_self_messages": True,
            "ocr_enabled": False,
            "status_dashboard": {
                "enabled": True,
                "host": "127.0.0.1",
                "port": dashboard_port,
            },
            "backup_retention_days": 1,
            "time_check_urls": ["http://127.0.0.1:1"],
            "reconnect_seconds": 1,
            "heartbeat_seconds": 5,
            "excel_batch_seconds": 2,
            "excel_full_refresh_seconds": 10,
            "rebuild_on_start": False,
        }
        with open(config_path, "w", encoding="utf-8") as stream:
            json.dump(config, stream, ensure_ascii=False, indent=2)

        process = subprocess.Popen(
            [sys.executable, APP, "--mode", "qq-live", "--config", config_path],
            cwd=ROOT,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        try:
            deadline = time.time() + 30
            final_status = None
            recent = None
            while time.time() < deadline:
                if process.poll() is not None:
                    raise AssertionError("采集器在重连测试中提前退出")
                try:
                    final_status = request_json(
                        f"http://127.0.0.1:{dashboard_port}/api/status"
                    )
                    recent = request_json(
                        f"http://127.0.0.1:{dashboard_port}/api/recent?limit=20"
                    )
                except OSError:
                    time.sleep(0.2)
                    continue
                groups = final_status.get("groups", {})
                processed = [
                    groups.get(group_id, {}).get("processed_events", 0)
                    for group_id in ("1001", "1002")
                ]
                kinds = {item["kind"] for item in recent.get("items", [])}
                if (
                    fake_napcat.connection_count >= 2
                    and final_status.get("connection") == "connected"
                    and final_status.get("websocket_reconnect_count", 0) >= 1
                    and processed == [2, 2]
                    and {"正式识别", "待生效", "不合格"}.issubset(kinds)
                ):
                    break
                time.sleep(0.2)
            else:
                raise AssertionError({
                    "connections": fake_napcat.connection_count,
                    "status": final_status,
                    "recent": recent,
                })

            heartbeat_deadline = time.time() + 8
            while time.time() < heartbeat_deadline:
                final_status = request_json(
                    f"http://127.0.0.1:{dashboard_port}/api/status"
                )
                if final_status.get("websocket_last_ping_at"):
                    break
                time.sleep(0.2)
            assert final_status.get("websocket_last_ping_at"), final_status
            assert final_status["groups"]["1001"]["queue_depth"] == 0
            assert final_status["groups"]["1002"]["queue_depth"] == 0

            for index in range(40):
                fake_napcat.send_event(event(
                    1001,
                    f"before-config-reload-{index}",
                    f"南宁到佛山大板{300 + index}",
                ))

            received_deadline = time.time() + 10
            while time.time() < received_deadline:
                final_status = request_json(
                    f"http://127.0.0.1:{dashboard_port}/api/status"
                )
                if final_status["groups"]["1001"].get("received_events", 0) >= 42:
                    break
                time.sleep(0.1)
            else:
                raise AssertionError("配置重载前的消息未全部进入群队列")

            configuration = request_json(
                f"http://127.0.0.1:{dashboard_port}/api/config"
            )
            target_group = next(
                group for group in configuration["groups"] if group["id"] == "1001"
            )
            target_group["name"] = "一群已改名"
            target_group["default_origin"] = "贵港"
            configuration["destinations"].append({
                "name": "肇庆",
                "aliases": ["肇庆", "四会"],
            })
            code, apply_result = request_json_response(
                f"http://127.0.0.1:{dashboard_port}/api/config",
                method="POST",
                payload=configuration,
            )
            assert code == 200 and apply_result["ok"], apply_result

            reload_deadline = time.time() + 35
            reloaded_config = None
            reloaded_status = None
            while time.time() < reload_deadline:
                if process.poll() is not None:
                    raise AssertionError("直接启动的采集器在配置重载后没有自行恢复")
                try:
                    code, candidate = request_json_response(
                        f"http://127.0.0.1:{dashboard_port}/api/config"
                    )
                    if code != 200:
                        time.sleep(0.2)
                        continue
                    reloaded_status = request_json(
                        f"http://127.0.0.1:{dashboard_port}/api/status"
                    )
                except OSError:
                    time.sleep(0.2)
                    continue
                changed_group = next(
                    group for group in candidate["groups"] if group["id"] == "1001"
                )
                if (
                    fake_napcat.connection_count >= 3
                    and reloaded_status.get("connection") == "connected"
                    and changed_group["name"] == "一群已改名"
                    and changed_group["default_origin"] == "贵港"
                    and any(row["name"] == "肇庆" for row in candidate["destinations"])
                ):
                    reloaded_config = candidate
                    break
                time.sleep(0.2)
            assert reloaded_config is not None, {
                "connections": fake_napcat.connection_count,
                "status": reloaded_status,
            }

            stable_output = os.path.join(directory, "data", "QQ群_1001")
            report_path = os.path.join(stable_output, "物流报价汇总.xlsx")
            assert os.path.exists(report_path), report_path
            workbook = load_workbook(report_path, read_only=True, data_only=True)
            try:
                active_rows = workbook["报价明细"].max_row - 1
                pending_rows = workbook["待生效运价"].max_row - 1
            finally:
                workbook.close()
            assert active_rows >= 41, active_rows
            assert pending_rows >= 1, pending_rows

            fake_napcat.send_event(event(1001, "after-config-reload", "四会大板245"))
            route_deadline = time.time() + 15
            route_item = None
            while time.time() < route_deadline:
                recent = request_json(
                    f"http://127.0.0.1:{dashboard_port}/api/recent?limit=200"
                )
                route_item = next((
                    item for item in recent.get("items", [])
                    if item.get("route") == "贵港到四会"
                ), None)
                if route_item:
                    break
                time.sleep(0.2)
            assert route_item and route_item["group_name"] == "一群已改名", route_item
        finally:
            process.terminate()
            try:
                process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait(timeout=5)
            fake_napcat.stop()

    print(json.dumps({
        "real_websocket_handshake": True,
        "disconnect_detected": True,
        "automatic_reconnect": True,
        "heartbeat_ping": True,
        "messages_before_and_after_reconnect": True,
        "three_way_status_api": True,
        "configuration_reload_without_supervisor": True,
        "queues_drained_before_reload": True,
        "renamed_group_kept_stable_output": True,
        "new_configuration_processed_after_reload": True,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
