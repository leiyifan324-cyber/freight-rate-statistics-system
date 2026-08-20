import importlib.util
import json
import os
import sys
import tempfile
import warnings
from datetime import date

from openpyxl import load_workbook


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(ROOT, "src")
sys.path.insert(0, SRC_DIR)
PARSER_PATH = os.path.join(SRC_DIR, "qq_freight_parser.py")


def load_freight_module():
    spec = importlib.util.spec_from_file_location(
        "freight_subcategory_statistics_test",
        PARSER_PATH,
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def record(record_date, cargo, price):
    return {
        "日期": record_date,
        "发布人": "测试员",
        "始发地": "南宁",
        "目的地": "佛山",
        "目的城市": "佛山",
        "货物小类": cargo,
        "货物大类": "",
        "报价文本": str(price),
        "最低报价": price,
        "最高报价": price,
        "平均报价": price,
        "是否整车价": "是" if price > 1000 else "否",
        "原始消息": f"南宁到佛山{cargo}{price}",
    }


def read_daily_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["每日统计"]
        headers = [cell.value for cell in sheet[1]]
        return [
            dict(zip(headers, row))
            for row in sheet.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in row)
        ]
    finally:
        workbook.close()


def main():
    warnings.filterwarnings("ignore", message="Glyph .* missing from font")
    freight = load_freight_module()
    freight.CARGO_TYPES = {
        "大板": {"aliases": ["大板"], "category": "板材"},
        "美人松": {"aliases": ["美人松"], "category": "木材"},
        "白松": {"aliases": ["白松"], "category": "木材"},
    }

    assert freight.spring_festival_date(2024) == date(2024, 2, 10)
    assert freight.spring_festival_date(2025) == date(2025, 1, 29)
    assert freight.spring_festival_date(2026) == date(2026, 2, 17)
    assert freight.spring_festival_date(2027) == date(2027, 2, 6)
    assert freight.spring_festival_year_for_date(date(2026, 2, 16)) == 2025
    assert freight.spring_festival_year_for_date(date(2026, 2, 17)) == 2026

    records = [
        record("2024-02-10", "美人松", 200),
        record("2025-01-28", "美人松", 220),
        record("2025-01-29", "美人松", 240),
        record("2026-02-16", "美人松", 250),
        record("2026-02-17", "美人松", 260),
        record("2027-02-06", "美人松", 280),
        record("2025-01-29", "白松", 300),
        record("2026-02-17", "大板", 400),
        record("2026-03-01", "美人松", 3000),
    ]
    daily = freight.summarize_daily_quote_by_subcategory(records)
    assert not any(row["平均运价"] == 3000 for row in daily)
    history = freight.merge_daily_history(daily)

    with tempfile.TemporaryDirectory(prefix="freight-subcategory-") as temp_root:
        previous_directory = os.getcwd()
        try:
            os.chdir(temp_root)
            os.makedirs(freight.CHART_FOLDER, exist_ok=True)
            existing_big_chart = os.path.join(
                freight.CHART_FOLDER,
                "南宁到佛山_木材_月价格走势图.png",
            )
            with open(existing_big_chart, "wb") as stream:
                stream.write(b"existing-big-category-chart")

            as_of = date(2027, 2, 7)
            summary = freight.write_small_category_outputs(history, as_of)
            assert summary["current_spring_year"] == 2027
            assert summary["subcategory_count"] == 3, summary
            assert os.path.isfile(existing_big_chart)

            beauty_dir = os.path.join(
                freight.SMALL_CATEGORY_OUTPUT_FOLDER,
                "木材",
                "美人松",
            )
            current_excel = os.path.join(beauty_dir, "美人松_每日统计.xlsx")
            current_rows = read_daily_rows(current_excel)
            assert [row["日期"] for row in current_rows] == ["2027-02-06"]
            assert current_rows[0]["春节年度"] == "2027春节年度"
            assert current_rows[0]["年度开始"] == "2027-02-06"
            assert current_rows[0]["年度结束"] == "2028-01-25"

            archive_2026 = os.path.join(
                beauty_dir,
                freight.SMALL_CATEGORY_ARCHIVE_FOLDER,
                "2026春节年度",
            )
            archive_excel = os.path.join(
                archive_2026,
                "美人松_2026春节年度_每日统计.xlsx",
            )
            archive_rows = read_daily_rows(archive_excel)
            assert [row["日期"] for row in archive_rows] == ["2026-02-17"]
            assert any(name.endswith("周价格走势图.png") for name in os.listdir(archive_2026))
            assert any(name.endswith("月价格走势图.png") for name in os.listdir(archive_2026))
            assert any(name.endswith("年价格走势图.png") for name in os.listdir(archive_2026))

            state_path = os.path.join(
                archive_2026,
                freight.SMALL_CATEGORY_ARCHIVE_STATE,
            )
            first_state = freight.read_json_object(state_path)
            first_state_mtime = os.stat(state_path).st_mtime_ns
            second_summary = freight.write_small_category_outputs(history, as_of)
            assert second_summary["archive_skipped_count"] >= 1
            assert os.stat(state_path).st_mtime_ns == first_state_mtime

            late_daily = freight.summarize_daily_quote_by_subcategory([
                *records,
                record("2026-03-02", "美人松", 270),
            ])
            late_history = freight.merge_daily_history(late_daily)
            freight.write_small_category_outputs(late_history, as_of)
            second_state = freight.read_json_object(state_path)
            assert second_state["fingerprint"] != first_state["fingerprint"]
            assert [row["日期"] for row in read_daily_rows(archive_excel)] == [
                "2026-02-17",
                "2026-03-02",
            ]
        finally:
            os.chdir(previous_directory)

    print(json.dumps({
        "spring_festival_dates_verified": True,
        "offline_multi_year_archive_recovered": True,
        "small_archive_idempotent": True,
        "late_old_record_rebuilds_archive": True,
        "current_excel_contains_current_spring_year_only": True,
        "large_category_chart_preserved": True,
        "full_truck_price_excluded": True,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
