"""Isolated year controls; never opens/closes a production statistical year."""
import hashlib
import json
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
import freight_archive as archive
import qq_freight_parser as parser
from freight_data import FreightDataService
from freight_runtime import FreightDatabase, RuntimeStatus, start_status_dashboard


class ManualArchiveTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='freight-year-')
        self.root = Path(self.tmp.name)
        self.today = patch.object(parser, 'trusted_today', return_value=date(2026, 3, 1))
        self.clock = self.today.start()
        self.profiles = {}
        for key in ('1001', '1002'):
            folder = self.root / key
            folder.mkdir()
            self.profiles[key] = {'group_id':key, 'group_name':'演示群'+key,
                'output_dir':str(folder), 'database_file':str(folder/'运价数据.db'),
                'input_file':str(folder/parser.INPUT_FILE), 'default_origin':'南宁',
                'data_lifecycle':{}, 'backup_enabled':False}
            FreightDatabase(self.profiles[key]['database_file'])
        self.service = FreightDataService(self.profiles)
        self.years = self.service.archive
        self.db = FreightDatabase(self.profiles['1001']['database_file'])
        self.add('2025-12-31', 80)
        self.add('2026-01-01', 100)
        self.add('2026-01-01', 200)
        self.add('2026-01-02', 300)
        self.add('2026-02-01', 400)
        self.add('2026-02-01', 2000)
        self.add('2026-03-01', 900)
        self.add('2026-03-02', 950)

    def tearDown(self):
        self.join_worker()
        self.today.stop()
        self.tmp.cleanup()

    def add(self, day, price, group='1001'):
        record = parser.parse_freight_line(f'南宁到佛山 大板 {price}', day, '测试人')
        self.assertIsNotNone(record)
        key = parser.build_record_content_dedup_key(record)
        FreightDatabase(self.profiles[group]['database_file']).insert_records([(key,record)])
        return key, record

    def body(self, **kwargs):
        return {'version':self.years.status()['version'], 'confirm':True, **kwargs}

    def open_year(self, name='2025—2026统计年度', start='2026-01-01'):
        return self.years.open(self.body(name=name, start_day=start))

    def close_year(self, end='2026-02-28'):
        self.years.close(self.body(end_day=end))
        self.join_worker()
        self.assertEqual(self.years.status()['current']['state'], 'closed')

    def join_worker(self):
        if self.years._worker:
            self.years._worker.join(30)
            self.assertFalse(self.years._worker.is_alive())

    def test_legacy_status_does_not_enable_or_create_database(self):
        self.assertEqual(self.years.status()['mode'], 'legacy')
        self.assertFalse(self.years.path.exists())
        self.assertEqual(len(archive.active_records(self.profiles['1001'],self.db.fetch_records())),8)

    def test_first_open_current_query_and_no_spring_cutoff(self):
        self.open_year()
        records = archive.active_records(self.profiles['1001'],self.db.fetch_records())
        self.assertEqual(len(records),7)
        stats = self.service.statistics({'group_id':'1001'})
        self.assertEqual(stats['quote_count'],5)  # Historical December, future and full-truck are excluded.
        self.assertIn('手动年度', stats['note'])
        notes = [r['year_note'] for r in self.service.query({'group_id':'1001'})['items']]
        self.assertEqual(notes.count('本统计年度'),7)
        with self.assertRaises(ValueError):
            self.open_year('重复开启')

    def test_fixed_snapshot_mean_empty_group_and_reports(self):
        self.open_year()
        preview = self.years.preview_close({'end_day':'2026-02-28'})
        self.assertEqual([g['count'] for g in preview['groups']], [5,0])
        self.close_year()
        details = self.years.details({'cycle_id':'1'})
        a = details['groups'][0]['annual'][0]
        # Jan daily means 150,300 -> 225; Feb 400 -> annual (225+400)/2.
        self.assertEqual(a['平均运价'],312.5)
        self.assertEqual(a['报价数量'],4)
        self.assertEqual(a['数据月份数'],2)
        self.assertEqual(details['groups'][0]['count'],5)
        self.assertEqual(details['groups'][1]['annual'],[])
        path = self.root/'1001'/archive.ARCHIVE_FOLDER/'1'/'年度统计汇总.xlsx'
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True)
        self.assertEqual(len(book.sheetnames),10)
        self.assertEqual(book['报价明细'].max_row,6)
        book.close()
        with self.db.session() as con:
            con.execute("DELETE FROM freight_records WHERE effective_date='2026-01-01'")
        self.add('2026-02-25',555)
        with patch.object(parser,'PRICE_THRESHOLD',50):
            self.assertEqual(self.years.details({'cycle_id':'1'}),details)
        parser.rebuild_qq_group_output(self.profiles['1001'])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)
        self.assertEqual(self.service.statistics({'group_id':'1001'})['quote_count'],0)
        self.assertTrue(self.service.files({'group_id':'1001'})['report_state']['manual'])
        # Windows CI may expose TEMP using an 8.3 alias (RUNNER~1). Verify
        # physical file identity, not two spellings of the same valid path.
        self.assertTrue(self.service.file_path({'group_id':'1001','path':details['groups'][0]['path']}).samefile(path))

    def test_pause_and_later_open_excludes_previous_and_future_receipts(self):
        self.open_year()
        self.close_year()
        self.add('2026-03-01',111)
        self.add('2026-03-05',222)  # Received while paused; must not become eligible later.
        self.clock.return_value=date(2026,3,5)
        self.open_year('下一统计年度','2026-03-05')
        self.assertEqual(self.service.statistics({'group_id':'1001'})['quote_count'],0)
        self.add('2026-03-05',333)
        self.assertEqual(self.service.statistics({'group_id':'1001'})['quote_count'],1)
        self.assertEqual(self.service.statistics({'group_id':'1001'})['items'][0]['平均运价'],333)
        self.close_year('2026-03-05')
        self.assertEqual(self.years.details({'cycle_id':'2'})['groups'][0]['count'],1)
        self.assertEqual(self.db.count_records(),11)  # No deletion by archive or pause.

    def test_validation_and_stale_confirms(self):
        for body in (self.body(name='',start_day='2026-01-01'),
                     self.body(name='x',start_day='2026-03-02'),
                     self.body(name='x',start_day='2026-2-01'),
                     {'version':0,'name':'x','start_day':'2026-01-01'},
                     self.body(name='x\n',start_day='2026-01-01',version=True)):
            with self.assertRaises(ValueError):
                self.years.open(body)
        self.open_year()
        for end in ('2025-12-31','2026-03-02','not-a-date'):
            with self.assertRaises(ValueError):
                self.years.preview_close({'end_day':end})
        with self.assertRaises(ValueError):
            self.years.close({'version':0,'confirm':True,'end_day':'2026-02-28'})
        self.assertEqual(self.years.status()['current']['state'],'active')
        self.close_year('2026-03-01')
        for start in ('2026-03-01','2026-02-20','2026-03-02'):
            with self.assertRaises(ValueError):
                self.open_year('next',start)
        self.clock.return_value=date(2026,3,2)
        with self.assertRaises(ValueError):
            self.open_year('2025—2026统计年度','2026-03-02')

    def test_concurrent_close_creates_exactly_one_snapshot(self):
        self.open_year()
        body=self.body(end_day='2026-02-28')
        results=[]
        def close_once():
            try:
                self.years.close(body)
                results.append('ok')
            except ValueError:
                results.append('conflict')
        threads=[threading.Thread(target=close_once) for _ in range(2)]
        for t in threads:t.start()
        for t in threads:t.join(30)
        self.join_worker()
        self.assertCountEqual(results,['ok','conflict'])
        with self.years.session() as con:
            self.assertEqual(con.execute('SELECT count(*) FROM snapshots').fetchone()[0],5)

    def test_pending_received_messages_must_drain_before_boundary(self):
        self.db.enqueue_inbound_event('queued', {'message':'南宁到佛山 大板 700'})
        try:
            with self.assertRaisesRegex(ValueError,'排队'):
                self.open_year()
            self.assertEqual(self.years.status()['mode'],'legacy')
            with self.db.session() as con:
                con.execute('DELETE FROM inbound_events')
            self.open_year()
            self.db.enqueue_inbound_event('queued2', {'message':'test'})
            with self.assertRaisesRegex(ValueError,'排队'):
                self.years.close(self.body(end_day='2026-02-28'))
            self.assertEqual(self.years.status()['current']['state'],'active')
        finally:
            self.db.close_inbound_event_queue()

    def test_failure_retry_reuses_snapshot(self):
        self.open_year()
        with patch.object(archive,'write_reports',side_effect=PermissionError('locked')):
            self.years.close(self.body(end_day='2026-02-28'))
            self.join_worker()
        self.assertEqual(self.years.status()['current']['state'],'error')
        with self.assertRaises(ValueError):self.open_year('next','2026-03-01')
        self.add('2026-02-28',666)
        self.years.retry(self.body())
        self.join_worker()
        self.assertEqual(self.years.status()['current']['state'],'closed')
        self.assertEqual(self.years.details({'cycle_id':'1'})['groups'][0]['count'],5)
        with self.assertRaises(ValueError):self.years.retry(self.body())

    def test_restart_resumes_and_duplicate_worker_does_not_rewrite(self):
        self.open_year()
        with patch.object(self.years,'resume'):
            self.years.close(self.body(end_day='2026-02-28'))
        self.assertEqual(self.years.status()['current']['state'],'closing')
        self.years=archive.ArchiveService(self.profiles)
        other=archive.ArchiveService(self.profiles)
        self.years.resume()
        other.resume()
        self.join_worker()
        if other._worker:other._worker.join(30)
        self.assertEqual(self.years.status()['current']['state'],'closed')
        self.assertEqual(self.years.status()['version'],3)

    def test_empty_year_and_injection_are_plain_text(self):
        name='=1+1 <img src=x onerror=alert(1)>'
        self.open_year(name,'2026-03-01')
        self.close_year('2026-03-01')
        from openpyxl import load_workbook
        path=self.root/'1002'/archive.ARCHIVE_FOLDER/'1'/'年度统计汇总.xlsx'
        book=load_workbook(path)
        self.assertEqual(book['年度说明']['A2'].value,name)
        self.assertEqual(book['年度说明']['A2'].data_type,'s')
        book.close()
        csv=(path.parent/'年度说明.csv').read_text(encoding='utf-8-sig')
        self.assertIn("'=1+1",csv)

    def test_one_seven_thirty_365_days(self):
        for n in (1,7,30,365):
            start=date(2025,1,1)
            rows=[]
            for i in range(n):
                day=(start+timedelta(days=i)).isoformat()
                r=parser.parse_freight_line('南宁到佛山 大板 100',day,'test')
                rows.append(r)
            cycle={'name':str(n)+'天','start_day':start.isoformat(),'end_day':rows[-1]['日期'],'state':'closing'}
            tables=archive.tables_for(rows,cycle,archive.context_now(),rows[-1]['日期'])
            a=tables['大类年度汇总'][0]
            self.assertEqual(a['平均运价'],100)
            self.assertEqual(a['报价数量'],n)
            self.assertEqual(a['数据天数'],n)
            weeks=[r['周份'] for r in tables['大类周均']]
            self.assertEqual(weeks,list(range(1,(n-1)//7+2)))

    def test_http_post_security_and_download(self):
        status=RuntimeStatus(str(self.root/'status'))
        server=start_status_dashboard(status,port=0,data_service=self.service)
        base=f'http://127.0.0.1:{server.server_port}'
        try:
            with urllib.request.urlopen(base+'/?view=archive') as response:
                html=response.read().decode()
                csrf=response.headers['X-Freight-CSRF-Token']
            self.assertIn('id="archiveView"',html)
            self.assertIn('年度归档',html)
            with urllib.request.urlopen(base+'/api/data/archive-status') as response:
                self.assertEqual(json.load(response)['mode'],'legacy')
            body=json.dumps(self.body(name='测试',start_day='2026-01-01')).encode()
            for headers in ({},{'Origin':'https://attacker.invalid','X-Freight-CSRF':csrf}):
                req=urllib.request.Request(base+'/api/data/archive-open',data=body,
                    headers={'Content-Type':'application/json',**headers})
                with self.assertRaises(urllib.error.HTTPError) as error:urllib.request.urlopen(req)
                self.assertEqual(error.exception.code,403)
                error.exception.close()
            req=urllib.request.Request(base+'/api/data/archive-open',data=body,
                headers={'Content-Type':'application/json','Origin':base,'X-Freight-CSRF':csrf})
            with urllib.request.urlopen(req) as response:
                self.assertIn('message',json.load(response))
            self.assertEqual(self.years.status()['current']['state'],'active')
        finally:
            server.shutdown()
            server.server_close()
            status.close()


if __name__=='__main__':
    unittest.main(verbosity=2)
