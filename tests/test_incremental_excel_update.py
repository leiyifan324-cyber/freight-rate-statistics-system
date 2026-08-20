import importlib.util
import json
import os
import sys
import tempfile
import time

from openpyxl import load_workbook


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(ROOT, "src")
sys.path.insert(0, SRC_DIR)


PARSER_PATH = os.path.join(SRC_DIR, "qq_freight_parser.py")
RULES_PATH = os.path.join(ROOT, "config", "freight_rules.default.json")


def load_freight_module():
    spec = importlib.util.spec_from_file_location("freight_incremental_test", PARSER_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def make_event(message_id, text):
    return {
        "post_type": "message",
        "message_type": "group",
        "group_id": 123456789,
        "message_id": message_id,
        "time": time.time(),
        "self_id": 1765614718,
        "user_id": 2000000001,
        "sender": {"nickname": "增量测试"},
        "message": [{"type": "text", "data": {"text": text}}],
    }


def workbook_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            "detail": workbook["报价明细"].max_row,
            "pending": workbook["待生效运价"].max_row,
            "daily": workbook["历史每日平均运价"].max_row,
        }
    finally:
        workbook.close()


def main():
    freight = load_freight_module()
    freight.load_rules_config(RULES_PATH)

    with tempfile.TemporaryDirectory(prefix="freight-excel-append-") as temp_root:
        config = {
            "group_ids": {"123456789"},
            "group_names": {"123456789": "测试"},
            "group_default_origins": {"123456789": "南宁"},
            "output_root": temp_root,
            "rules_file": RULES_PATH,
            "backup_enabled": False,
            "backup_retention_days": 30,
            "config_path": "",
        }
        profiles = freight.build_qq_group_profiles(config)
        profile = profiles["123456789"]
        ingestor = freight.OneBotFreightIngestor(profiles)
        freight.rebuild_qq_group_output(profile)
        report_path = os.path.join(profile["output_dir"], freight.REPORT_XLSX)
        initial_rows = workbook_rows(report_path)
        assert initial_rows["detail"] == 1
        assert initial_rows["pending"] == 1

        current = ingestor.ingest(make_event("append-1", "南宁到佛山大板235"))
        current_summary = freight.append_qq_group_excel_records(
            profile,
            current["_inserted_records"],
        )
        assert current_summary["excel_update_mode"] == "append"
        assert current_summary["aggregate_state"] == "pending_refresh"

        future = ingestor.ingest(make_event("append-2", "明天南宁到杭州红板240"))
        future_summary = freight.append_qq_group_excel_records(
            profile,
            future["_inserted_records"],
        )
        assert future_summary["pending_records"] == 1

        second = ingestor.ingest(make_event("append-3", "南宁到佛山大板236"))
        freight.append_qq_group_excel_records(profile, second["_inserted_records"])
        appended_rows = workbook_rows(report_path)
        assert appended_rows["detail"] == 3, appended_rows
        assert appended_rows["pending"] == 2, appended_rows
        assert appended_rows["daily"] == 1, appended_rows

        overlap_summary = freight.append_qq_group_excel_records(
            profile,
            current["_inserted_records"],
        )
        assert overlap_summary["excel_append_count"] == 0
        assert workbook_rows(report_path) == appended_rows

        duplicate = ingestor.ingest(make_event("append-4", "南宁到佛山大板236"))
        assert duplicate["status"] == "duplicate_content", duplicate
        metadata_summary = freight.append_qq_group_excel_records(
            profile,
            duplicate["_inserted_records"],
        )
        assert metadata_summary["excel_update_mode"] == "metadata_only"
        assert workbook_rows(report_path) == appended_rows

        full_summary = freight.rebuild_qq_group_output(profile)
        refreshed_rows = workbook_rows(report_path)
        assert full_summary["active_records"] == 2
        assert full_summary["pending_records"] == 1
        assert refreshed_rows["detail"] == 3
        assert refreshed_rows["pending"] == 2
        assert refreshed_rows["daily"] >= 2

    print(json.dumps({
        "detail_rows_appended": True,
        "pending_rows_appended": True,
        "duplicates_not_appended": True,
        "full_and_incremental_overlap_deduped": True,
        "aggregates_deferred_then_refreshed": True,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
