import importlib.util
import json
import os
import sys
import tempfile
import time
from datetime import date

import pandas as pd
from openpyxl import load_workbook


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(ROOT, "src")
sys.path.insert(0, SRC_DIR)

from freight_runtime import DASHBOARD_HTML, FreightConfigurationManager  # noqa: E402


PARSER_PATH = os.path.join(SRC_DIR, "qq_freight_parser.py")


def load_freight_module():
    spec = importlib.util.spec_from_file_location("freight_cargo_rules_test", PARSER_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def write_json(path, value):
    with open(path, "w", encoding="utf-8") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)


def cargo_rules():
    return {
        "default_origin": "南宁",
        "origin_groups": {
            "南宁": ["南宁", "南宁市"],
            "贵港": ["贵港", "贵港市"],
        },
        "destination_groups": {
            "佛山": ["佛山", "南海"],
            "杭州": ["杭州", "余杭"],
        },
        "board_types": ["大板", "红板", "木方", "钢管"],
        "default_cargo_subcategory": "大板",
        "default_cargo_category": "板材",
        "cargo_types": [
            {"name": "大板", "aliases": ["大板"], "category": ""},
            {"name": "红板", "aliases": ["红板", "红板材"], "category": ""},
            {"name": "木方", "aliases": ["木方", "建筑木方"], "category": "木材"},
            {"name": "钢管", "aliases": ["钢管"], "category": "钢材"},
        ],
        "cargo_route_scopes": [
            {
                "level": "category",
                "cargo": "板材",
                "routes": [{"origin": "南宁", "destination": "佛山"}],
            },
            {
                "level": "subcategory",
                "cargo": "红板",
                "routes": [{"origin": "南宁", "destination": "杭州"}],
            },
            {
                "level": "subcategory",
                "cargo": "木方",
                "routes": [{"origin": "贵港", "destination": "杭州"}],
            },
        ],
        "price_threshold": 1000,
        "invalid_line_keywords": ["求车"],
    }


def make_event(message_id, text):
    return {
        "post_type": "message",
        "message_type": "group",
        "group_id": 123456,
        "message_id": message_id,
        "time": time.time(),
        "self_id": 1,
        "user_id": 2,
        "sender": {"nickname": "货物规则测试"},
        "message": [{"type": "text", "data": {"text": text}}],
    }


def main():
    freight = load_freight_module()
    with tempfile.TemporaryDirectory(prefix="freight-cargo-rules-") as directory:
        rules_path = os.path.join(directory, "freight_rules.json")
        write_json(rules_path, cargo_rules())
        freight.load_rules_config(rules_path)

        # 最长别名匹配后仍回写标准小类；留空的大类继承默认“板材”。
        red = freight.parse_freight_line(
            "南宁到杭州红板材120", "2026-08-19", "测试人", "南宁"
        )
        assert red["货物小类"] == "红板"
        assert red["货物大类"] == "板材"
        assert freight.validate_cargo_route(red)[0]

        # 小类规则覆盖大类规则：红板只允许杭州，即使板材大类允许佛山。
        red_foshan = freight.parse_freight_line(
            "南宁到佛山红板110", "2026-08-19", "测试人", "南宁"
        )
        allowed, reason = freight.validate_cargo_route(red_foshan)
        assert not allowed and "货物小类“红板”" in reason

        board_foshan = freight.parse_freight_line(
            "南宁到佛山大板100", "2026-08-19", "测试人", "南宁"
        )
        assert freight.validate_cargo_route(board_foshan)[0]
        board_hangzhou = freight.parse_freight_line(
            "南宁到杭州大板100", "2026-08-19", "测试人", "南宁"
        )
        assert not freight.validate_cargo_route(board_hangzhou)[0]

        # 没有专门筛选规则的钢材不受板材/木材规则影响，可走任意现有线路。
        steel = freight.parse_freight_line(
            "贵港到佛山钢管300", "2026-08-19", "测试人", "南宁"
        )
        assert freight.validate_cargo_route(steel)[0]

        formatted, reason = freight.format_freight_line_with_reason(
            "南宁到佛山235", "南宁"
        )
        assert formatted == "南宁到佛山 大板 235" and reason == ""

        # 实际接收器将越界货物线路记入不合格，而不是静默丢弃或写入数据库。
        output_root = os.path.join(directory, "output")
        profiles = freight.build_qq_group_profiles({
            "group_ids": {"123456"},
            "group_names": {"123456": "测试群"},
            "group_default_origins": {"123456": "南宁"},
            "output_root": output_root,
            "rules_file": rules_path,
        })
        ingestor = freight.OneBotFreightIngestor(profiles)
        rejected = ingestor.ingest(make_event("scope-reject", "南宁到杭州大板100"))
        assert rejected["status"] == "rejected", rejected
        assert "未开放线路" in rejected["reason"], rejected
        assert ingestor.databases["123456"].count_records() == 0
        accepted = ingestor.ingest(make_event("scope-accept", "贵港到佛山钢管300"))
        assert accepted["status"] == "added" and accepted["inserted_count"] == 1

        # 相同日期和线路的不同大类必须分别统计，不能混算平均价。
        steel_same_route = freight.parse_freight_line(
            "南宁到佛山钢管300", "2026-08-19", "测试人", "南宁"
        )
        summary = freight.summarize_daily_quote([board_foshan, steel_same_route])
        assert {row["货物"] for row in summary} == {"板材", "钢材"}
        assert len(summary) == 2
        daily_df = pd.DataFrame(summary)
        daily_df["日期排序"] = pd.to_datetime(daily_df["日期"])
        weekly = freight.aggregate_weekly_statistics(
            daily_df, date(2026, 8, 19), date(2026, 8, 31)
        )
        monthly = freight.aggregate_monthly_statistics(
            daily_df, date(2026, 8, 19), date(2026, 8, 31)
        )
        yearly = freight.aggregate_yearly_statistics(
            monthly, date(2026, 8, 19), date(2026, 8, 31)
        )
        assert set(weekly["货物"]) == {"板材", "钢材"}
        assert set(monthly["货物"]) == {"板材", "钢材"}
        assert set(yearly["货物"]) == {"板材", "钢材"}

        report_path = os.path.join(directory, "货物分类输出.xlsx")
        history = freight.enrich_history_indicators(pd.DataFrame(summary))
        freight.save_excel(
            [board_foshan, steel_same_route],
            history.to_dict(orient="records"),
            weekly.to_dict(orient="records"),
            monthly.to_dict(orient="records"),
            yearly.to_dict(orient="records"),
            report_path,
        )
        workbook = load_workbook(report_path, read_only=True, data_only=True)
        daily_sheet = workbook["历史每日平均运价"]
        headers = [cell.value for cell in next(daily_sheet.iter_rows(min_row=1, max_row=1))]
        cargo_column = headers.index("货物") + 1
        output_categories = {
            row[cargo_column - 1].value
            for row in daily_sheet.iter_rows(min_row=2)
        }
        workbook.close()
        assert output_categories == {"板材", "钢材"}

        # 管理接口完整保存结构化货物规则，并拒绝不存在的线路。
        live_path = os.path.join(directory, "qq_live_config.json")
        write_json(live_path, {
            "ws_url": "ws://127.0.0.1:3001",
            "group_ids": ["123456"],
            "group_names": {"123456": "测试群"},
            "group_default_origins": {"123456": "南宁"},
        })
        manager = FreightConfigurationManager(live_path, rules_path)
        snapshot = manager.snapshot()
        assert snapshot["default_cargo_subcategory"] == "大板"
        assert snapshot["cargo_types"][1]["aliases"] == ["红板", "红板材"]
        result = manager.apply(snapshot)
        assert result["ok"]
        saved = manager.snapshot()
        assert saved["cargo_route_scopes"] == snapshot["cargo_route_scopes"]

        invalid = manager.snapshot()
        invalid["cargo_route_scopes"][0]["routes"][0]["destination"] = "不存在"
        try:
            manager.apply(invalid)
            raise AssertionError("不存在的筛选线路本应被拒绝")
        except ValueError as exc:
            assert "不存在的目的地" in str(exc)

    assert 'id="cargoRows"' in DASHBOARD_HTML
    assert 'id="cargoScopeRows"' in DASHBOARD_HTML
    assert 'id="routeDialog"' in DASHBOARD_HTML
    assert "小类规则优先于大类规则" in DASHBOARD_HTML
    print("CARGO_RULES_TESTS_OK")


if __name__ == "__main__":
    main()
