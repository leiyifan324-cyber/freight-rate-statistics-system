"""Local data browser, exact-record deletion and generated report previews."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import secrets
import threading
import time
from datetime import date, datetime
from pathlib import Path
from typing import Callable

from freight_runtime import FreightDatabase


class DataConflictError(ValueError):
    pass


def bounded_int(value: object, default: int, maximum: int) -> int:
    if value is None or value == "":
        return default
    if isinstance(value, bool) or not str(value).isdigit():
        raise ValueError("页码和每页条数必须是正整数。")
    number = int(str(value))
    if not 1 <= number <= maximum:
        raise ValueError(f"数值必须在1到{maximum}之间。")
    return number


class FreightDataService:
    def __init__(self, profiles: dict, status=None) -> None:
        self.profiles = profiles
        self.status = status
        self.refresh_callback: Callable | None = None
        self._previews: dict[str, dict] = {}
        self._lock = threading.RLock()
        from freight_archive import ArchiveService
        self.archive = ArchiveService(profiles, lambda key, **kwargs: self.refresh_callback(key, **kwargs) if self.refresh_callback else False)
        from freight_collection import CollectionControl
        from freight_daily import DailyReportService
        self.collection = CollectionControl(self.archive.path.parent) if self.archive.path else None
        self.daily = DailyReportService(self)

    def collection_status(self) -> dict:
        return self.collection.snapshot() if self.collection else {'state': 'unavailable', 'enabled': False, 'pending': 0}

    def collection_change(self, body: dict) -> dict:
        if self.collection is None:
            raise ValueError('请先配置QQ群。')
        return self.collection.change(body)

    def profile(self, group_id: str) -> dict:
        if str(group_id) not in self.profiles:
            raise ValueError("请先选择一个已配置的QQ群。")
        return self.profiles[str(group_id)]

    def options(self) -> dict:
        import qq_freight_parser as parser

        return {
            "groups": [{"id": key, "name": value["group_name"]}
                       for key, value in self.profiles.items()],
            "origins": list(parser.ORIGIN_GROUPS),
            "destinations": list(parser.DEST_GROUPS),
            "cargo": list(parser.CARGO_TYPES),
            "today": parser.trusted_today().isoformat(),
        }

    def _selection(self, params: dict) -> tuple[FreightDatabase, str, list, bool]:
        import qq_freight_parser as parser

        profile = self.profile(params.get("group_id", ""))
        kind = params.get("kind", "all")
        if kind not in {"all", "active", "pending", "rejected"}:
            raise ValueError("数据分类无效。")
        rejected = kind == "rejected"
        if rejected:
            sql = """SELECT CAST(id AS TEXT) AS id, substr(message_time,1,10) AS day,
                '' AS origin, '' AS destination, '' AS city, '' AS cargo,
                NULL AS average, '' AS price, sender, message_text AS detail,
                reason, created_at, '' AS record_json FROM rejected_messages"""
        else:
            sql = """SELECT dedup_key AS id, effective_date AS day,
                json_extract(record_json,'$.始发地') AS origin,
                json_extract(record_json,'$.目的地') AS destination,
                json_extract(record_json,'$.目的城市') AS city,
                json_extract(record_json,'$.货物小类') AS cargo,
                json_extract(record_json,'$.平均报价') AS average,
                json_extract(record_json,'$.报价文本') AS price,
                json_extract(record_json,'$.发布人') AS sender,
                json_extract(record_json,'$.原始消息') AS detail,
                '' AS reason, created_at, record_json FROM freight_records"""
        where, values = [], []
        if kind in {"active", "pending"}:
            where.append("day <= ?" if kind == "active" else "day > ?")
            values.append(parser.trusted_today().isoformat())
        for key, operator in (("start", ">="), ("end", "<=")):
            value = str(params.get(key, "")).strip()
            if value:
                try:
                    if date.fromisoformat(value).isoformat() != value:
                        raise ValueError
                except ValueError as exc:
                    raise ValueError("日期应为 YYYY-MM-DD。") from exc
                where.append(f"day {operator} ?")
                values.append(value)
        if params.get("start") and params.get("end") and params["start"] > params["end"]:
            raise ValueError("开始日期不能晚于结束日期。")
        for key, column in (("origin", "origin"), ("destination", "city"), ("cargo", "cargo")):
            value = str(params.get(key, "")).strip()
            if value and not rejected:
                if len(value) > 100:
                    raise ValueError("筛选内容过长。")
                where.append(f"{column} = ?")
                values.append(value)
        for key, operator in (("min_price", ">="), ("max_price", "<=")):
            value = str(params.get(key, "")).strip()
            if value and not rejected:
                try:
                    number = float(value)
                    if not math.isfinite(number):
                        raise ValueError
                except ValueError as exc:
                    raise ValueError("价格必须是有限数字。") from exc
                where.append(f"CAST(average AS REAL) {operator} ?")
                values.append(number)
        if params.get("min_price") and params.get("max_price") and not rejected:
            if float(params["min_price"]) > float(params["max_price"]):
                raise ValueError("最低价格不能大于最高价格。")
        keyword = str(params.get("q", "")).strip()
        if keyword:
            if len(keyword) > 200:
                raise ValueError("搜索关键词最多200个字符。")
            where.append("instr(lower(coalesce(detail,'') || ' ' || coalesce(sender,'') || ' ' || reason),lower(?)) > 0")
            values.append(keyword)
        selected = f"SELECT * FROM ({sql})"
        if where:
            selected += " WHERE " + " AND ".join(where)
        return FreightDatabase(profile["database_file"]), selected, values, rejected

    def query(self, params: dict) -> dict:
        import qq_freight_parser as parser

        database, sql, values, rejected = self._selection(params)
        page = bounded_int(params.get("page"), 1, 1000000)
        size = bounded_int(params.get("page_size"), 25, 100)
        with database.session() as connection:
            connection.execute("BEGIN")
            total = connection.execute(f"SELECT COUNT(*) FROM ({sql})", values).fetchone()[0]
            page = min(page, max(1, math.ceil(total / size)))
            rows = connection.execute(sql + " ORDER BY day DESC, created_at DESC, id DESC LIMIT ? OFFSET ?",
                                      [*values, size, (page - 1) * size]).fetchall()
        items = []
        from freight_archive import active_records, read_year
        year_state = read_year(self.profile(params['group_id']).get('archive_database'))
        current_keys = None
        if not rejected and year_state['mode']=='manual':
            current_keys = {parser.build_record_content_dedup_key(r) for r in active_records(self.profile(params['group_id']), [json.loads(r['record_json']) for r in rows])}
        for row in rows:
            item = dict(row)
            record = item.pop("record_json")
            item["type"] = "rejected" if rejected else "freight"
            item["kind"] = "不合格" if rejected else (
                "待生效" if str(item["day"]) > parser.trusted_today().isoformat() else "正式识别")
            if not rejected:
                normalized = parser.normalize_record_cargo(json.loads(record))
                allowed, reason = parser.validate_cargo_route(normalized)
                item["category"] = normalized.get("货物大类", "")
                item["eligible"] = allowed
                item["reason"] = reason if not allowed else ""
                if current_keys is not None:
                    item['year_note'] = '本统计年度' if item['id'] in current_keys else '不计入当前年度（历史 / 暂停期）'
            items.append(item)
        return {"items": items, "total": total, "page": page, "page_size": size,
                "pages": max(1, math.ceil(total / size)), "report_state": self.report_state(params["group_id"])}

    def statistics(self, params: dict) -> dict:
        import qq_freight_parser as parser

        database, sql, values, rejected = self._selection(params)
        level = params.get("level", "category")
        if level not in {"category", "subcategory"}:
            raise ValueError("统计层级无效。")
        records = []
        if not rejected:
            with database.session() as connection:
                for row in connection.execute(sql, values):
                    record = parser.normalize_record_cargo(json.loads(row["record_json"]))
                    if str(record.get("日期", "")) <= parser.trusted_today().isoformat() and parser.validate_cargo_route(record)[0]:
                        records.append(record)
        summarize = parser.summarize_daily_quote if level == "category" else parser.summarize_daily_quote_by_subcategory
        from freight_archive import active_records, read_year
        profile = self.profile(params['group_id'])
        state = read_year(profile.get('archive_database'))
        records = active_records(profile, records)
        items = list(reversed(summarize(records)))
        page = bounded_int(params.get("page"), 1, 1000000)
        size = bounded_int(params.get("page_size"), 25, 100)
        total = len(items)
        page = min(page, max(1, math.ceil(total / size)))
        return {"items": items[(page - 1) * size:page * size], "total": total,
                "page": page, "pages": max(1, math.ceil(total / size)),
                "quote_count": sum(item["报价数量"] for item in items),
                "price_threshold": parser.PRICE_THRESHOLD,
                "note": ('手动年度：仅显示当前统计年度；暂停时为空，旧年度请到“年度归档”查看快照。 ' if state['mode']=='manual' else '') + "按生效日期统计；待生效、线路规则不允许及超过价格阈值的报价不参与日均价。"}

    def report_state(self, group_id: str) -> dict:
        database = FreightDatabase(self.profile(group_id)["database_file"])
        with database.session() as connection:
            dirty = connection.execute("SELECT value FROM metadata WHERE key='reports_dirty'").fetchone()
        state = ((self.status.snapshot().get("groups", {}).get(group_id, {})) if self.status else {})
        return {"pending": bool(dirty), "state": state.get("excel_state", ""),
                "error": state.get("excel_error", ""), "updated_at": state.get("excel_last_updated_at", "")}

    @staticmethod
    def _rows(connection, kind: str, ids: list[str]) -> list[dict]:
        table, key = ("freight_records", "dedup_key") if kind == "freight" else ("rejected_messages", "id")
        marks = ",".join("?" for _ in ids)
        return [dict(row) for row in connection.execute(
            f"SELECT * FROM {table} WHERE {key} IN ({marks}) ORDER BY {key}", ids)]

    @staticmethod
    def _digest(rows: list[dict]) -> str:
        return hashlib.sha256(json.dumps(rows, ensure_ascii=False, sort_keys=True).encode()).hexdigest()

    def preview_delete(self, body: dict) -> dict:
        group_id, kind, raw_ids = str(body.get("group_id", "")), body.get("type"), body.get("ids")
        profile = self.profile(group_id)
        if kind not in {"freight", "rejected"} or not isinstance(raw_ids, list) or not 1 <= len(raw_ids) <= 100:
            raise ValueError("请选择1到100条同类别记录。")
        if any(not isinstance(value, str) or not value or len(value) > 200 for value in raw_ids):
            raise ValueError("记录标识无效。")
        ids = sorted(set(raw_ids))
        database = FreightDatabase(profile["database_file"])
        with database.session() as connection:
            rows = self._rows(connection, kind, ids)
        if len(rows) != len(ids):
            raise DataConflictError("部分记录已发生变化，请刷新后重新选择。")
        token = secrets.token_urlsafe(32)
        with self._lock:
            self._previews = {key: value for key, value in self._previews.items() if value["expires"] > time.monotonic()}
            if len(self._previews) >= 100:
                self._previews.pop(next(iter(self._previews)))
            self._previews[token] = {"group_id": group_id, "type": kind, "ids": ids,
                                     "digest": self._digest(rows), "expires": time.monotonic() + 300}
        return {"token": token, "count": len(rows), "group_name": profile["group_name"],
                "type": kind, "expires_seconds": 300,
                "message": "删除后更新当前Excel、CSV；已生成的每日日报保留快照并提示需重新生成。删除内容保存在本机删除记录中。"}

    def delete(self, body: dict) -> dict:
        import qq_freight_parser as parser

        if body.get("confirm") is not True:
            raise ValueError("请先预览并确认删除。")
        token = str(body.get("token", ""))
        with self._lock:
            preview = self._previews.get(token)
            if not preview or preview["expires"] < time.monotonic():
                raise DataConflictError("删除确认已过期，请重新选择记录。")
            if self.refresh_callback is None:
                raise DataConflictError("报表服务正在启动，请稍后再删除。")
            group_id, kind, ids = preview["group_id"], preview["type"], preview["ids"]
            database = FreightDatabase(self.profile(group_id)["database_file"])
            # Serialize with report writes. SQLite independently serializes competing writers.
            with parser.group_output_lock(self.profile(group_id)), database.session() as connection:
                connection.execute("BEGIN IMMEDIATE")
                rows = self._rows(connection, kind, ids)
                if self._digest(rows) != preview["digest"]:
                    raise DataConflictError("记录已被修改或清理，请刷新后重新确认。")
                batch_id = secrets.token_hex(12)
                deleted_at = datetime.now().astimezone().isoformat(timespec="seconds")
                for row in rows:
                    record_id = str(row["dedup_key"] if kind == "freight" else row["id"])
                    connection.execute("""INSERT INTO deleted_data(kind,record_id,row_json,batch_id,deleted_at)
                        VALUES(?,?,?,?,?)""", (kind, record_id, json.dumps(row, ensure_ascii=False), batch_id, deleted_at))
                table, key = ("freight_records", "dedup_key") if kind == "freight" else ("rejected_messages", "id")
                connection.execute(f"DELETE FROM {table} WHERE {key} IN ({','.join('?' for _ in ids)})", ids)
                connection.execute("INSERT OR REPLACE INTO metadata VALUES ('txt_migration_done','1')")
                connection.execute("INSERT OR REPLACE INTO metadata VALUES ('reports_dirty',?)", (batch_id,))
            self._previews.pop(token, None)
        queued = self.refresh_callback(group_id, force_full=True)
        return {"deleted": len(rows), "batch_id": batch_id, "report_pending": True,
                "message": f"已删除{len(rows)}条记录。" + ("报表正在重新生成。" if queued else "报表将在服务恢复后重新生成。")}

    def files(self, params: dict) -> dict:
        profile = self.profile(params.get("group_id", ""))
        root = Path(profile["output_dir"]).resolve()
        items = []
        def visit(folder: Path) -> None:
            if len(items) >= 3000 or not folder.exists():
                return
            for child in sorted(folder.iterdir()):
                if (child.is_symlink() or getattr(child, "is_junction", lambda: False)() or not child.resolve().is_relative_to(root)
                        or child.name.startswith(".") or child.name in {"备份", "删除备份", "logs"}):
                    continue
                if child.is_dir():
                    visit(child)
                elif child.suffix.lower() in {".csv", ".xlsx", ".png"} and "待更新" not in child.name:
                    info = child.stat()
                    items.append({"path": child.relative_to(root).as_posix(), "size": info.st_size,
                                  "modified": datetime.fromtimestamp(info.st_mtime).isoformat(timespec="seconds"),
                                  "type": child.suffix[1:].lower()})
                if len(items) >= 3000:
                    break
        visit(root)
        from freight_archive import read_year
        manual = read_year(profile.get('archive_database'))['mode']=='manual'
        if manual:
            items.sort(key=lambda row: (not row['path'].startswith('手动年度'),row['path']))
        state = self.report_state(params['group_id'])
        state['manual'] = manual
        return {"items": items, "report_state": state, "truncated": len(items) >= 3000}

    def file_path(self, params: dict) -> Path:
        root = Path(self.profile(params.get("group_id", ""))["output_dir"]).resolve()
        relative = Path(str(params.get("path", "")))
        if relative.is_absolute() or any(part in {"..", "备份", "删除备份", "logs"} or part.startswith(".") for part in relative.parts):
            raise ValueError("报表路径无效。")
        target = (root / relative).resolve()
        if not target.is_relative_to(root) or target.suffix.lower() not in {".csv", ".xlsx", ".png"} or "待更新" in target.name:
            raise ValueError("只允许访问当前群的公开报表文件。")
        if not target.is_file():
            raise ValueError("报表暂不存在，请刷新列表。")
        return target

    def preview_file(self, params: dict) -> dict:
        path = self.file_path(params)
        if path.suffix.lower() == ".png":
            return {"type": "png"}
        page = bounded_int(params.get("page"), 1, 1000000)
        size = bounded_int(params.get("page_size"), 50, 100)
        workbook = None
        sheets, sheet = [], ""
        stream = None
        try:
            if path.suffix.lower() == ".xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=True)
                sheets = workbook.sheetnames
                sheet = str(params.get("sheet") or sheets[0])
                if sheet not in sheets:
                    raise ValueError("工作表不存在。")
                iterator = workbook[sheet].iter_rows(values_only=True)
            else:
                stream = path.open(encoding="utf-8-sig", newline="")
                iterator = iter(csv.reader(stream))
            headers = [str(value or "") for value in next(iterator, [])]
            items, total = [], 0
            for index, row in enumerate(iterator):
                total = index + 1
                if (page - 1) * size <= index < page * size:
                    items.append([value if isinstance(value, (int, float)) else str(value or "") for value in row])
            return {"headers": headers, "items": items, "total": total, "page": page,
                    "pages": max(1, math.ceil(total / size)), "sheets": sheets, "sheet": sheet, "type": path.suffix[1:]}
        finally:
            if workbook:
                workbook.close()
            if stream:
                stream.close()
