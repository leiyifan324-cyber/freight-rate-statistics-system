"""User-controlled statistical years, immutable closing snapshots and local reports."""
from __future__ import annotations

import csv
import json
import math
import os
import sqlite3
import threading
import uuid
from collections import defaultdict
from contextlib import ExitStack, closing, contextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import mean


ARCHIVE_DB = '统计年度管理.db'
CURRENT_FOLDER = '手动年度统计/当前'
ARCHIVE_FOLDER = '手动年度归档'


def store_path(profiles: dict) -> Path | None:
    if not profiles:
        return None
    first = next(iter(profiles.values()))
    return Path(first.get('archive_database') or (Path(first['output_dir']).parent / ARCHIVE_DB)).resolve()


def parse_day(value: object) -> str:
    text = str(value or '')
    try:
        if date.fromisoformat(text).isoformat() != text:
            raise ValueError
    except ValueError as exc:
        raise ValueError('请选择有效日期（YYYY-MM-DD）。') from exc
    return text


def read_year(path: str | Path | None) -> dict:
    if not path or not Path(path).is_file():
        return {'mode':'legacy','version':0,'cycles':[],'current':None}
    with closing(sqlite3.connect(Path(path).resolve().as_uri()+'?mode=ro', uri=True, timeout=15)) as con:
        con.row_factory = sqlite3.Row
        con.execute('BEGIN')  # State and version must describe the same committed snapshot.
        cycles = [dict(r) for r in con.execute('SELECT * FROM cycles ORDER BY id DESC')]
        version = con.execute("SELECT value FROM settings WHERE key='version'").fetchone()
    return {'mode':'manual' if cycles else 'legacy','version':int(version[0]) if version else 0,
            'cycles':cycles,'current':cycles[0] if cycles else None}


def context_now() -> dict:
    import qq_freight_parser as parser
    return {'threshold':parser.PRICE_THRESHOLD,'categories':{c:parser.cargo_category(c) for c in parser.CARGO_TYPES},
            'default_category':parser.DEFAULT_CARGO_CATEGORY}


def selected_rows(profile: dict, cycle: dict, con: sqlite3.Connection) -> list[dict]:
    """Selection by effective date; exclusions make paused/pre-open receipts stay out."""
    from freight_runtime import FreightDatabase
    with FreightDatabase(profile['database_file']).session() as source:
        rows = [dict(r) for r in source.execute('SELECT dedup_key,record_json FROM freight_records WHERE effective_date>=? ORDER BY effective_date,dedup_key', (cycle['start_day'],))]
    excluded = {r[0] for r in con.execute('SELECT record_key FROM exclusions WHERE cycle_id=? AND group_id=?', (cycle['id'],profile['group_id']))}
    return [r for r in rows if r['dedup_key'] not in excluded]


def active_records(profile: dict, records: list[dict]) -> list[dict]:
    import qq_freight_parser as parser
    state = read_year(profile.get('archive_database'))
    if state['mode'] != 'manual':
        return records
    cycle = state['current']
    if cycle['state'] != 'active':
        return []
    with closing(sqlite3.connect(Path(profile['archive_database']).resolve().as_uri()+'?mode=ro',uri=True)) as con:
        excluded = {r[0] for r in con.execute('SELECT record_key FROM exclusions WHERE cycle_id=? AND group_id=?', (cycle['id'],profile['group_id']))}
    return [r for r in records if str(r.get('日期',''))>=cycle['start_day'] and parser.build_record_content_dedup_key(r) not in excluded]


def tables_for(records: list[dict], cycle: dict, context: dict, as_of: str) -> dict[str,list[dict]]:
    """Daily quote mean -> monthly daily mean -> annual monthly mean (existing convention)."""
    end = cycle.get('end_day') or as_of
    start = cycle['start_day']
    records = [r for r in records if start <= r.get('日期','') <= end]
    tables = {'年度说明':[{'统计年度':cycle['name'],'开始日期':start,'截止日期':end,
              '状态':'已归档快照' if cycle['state']!='active' else '统计中',
              '说明':'按生效日期；年度均价为有数据月份均价的平均值；暂停期不补入；超过阈值的整车价不计均价',
              '整车价阈值':context['threshold']}], '报价明细':records}
    for level in ('大类','小类'):
        grouped = defaultdict(list)
        for r in records:
            if not isinstance(r.get('平均报价'),(int,float)) or not math.isfinite(r['平均报价']) or r['平均报价']>context['threshold']:
                continue
            cargo = r.get('货物小类','')
            if level=='大类':
                cargo = context['categories'].get(cargo,context['default_category'])
            grouped[(r['日期'],r['始发地'],r['目的城市'],cargo)].append(r['平均报价'])
        daily = [{'日期':k[0],'始发地':k[1],'目的城市':k[2],'货物':k[3],
                  '平均运价':round(mean(v),3),'最低运价':min(v),'最高运价':max(v),'报价数量':len(v)} for k,v in sorted(grouped.items())]
        tables[level+'日均'] = daily
        monthly = []
        for period_type in ('周','月'):
            periods = defaultdict(list)
            for row in daily:
                label = row['日期'][:7] if period_type=='月' else (date.fromisoformat(row['日期'])-date.fromisoformat(start)).days//7+1
                periods[(label,row['始发地'],row['目的城市'],row['货物'])].append(row)
            values = [{period_type+'份':k[0],'始发地':k[1],'目的城市':k[2],'货物':k[3],
                       '平均运价':round(mean(r['平均运价'] for r in v),3),
                       '最低运价':min(r['最低运价'] for r in v),'最高运价':max(r['最高运价'] for r in v),
                       '数据天数':len(v),'报价数量':sum(r['报价数量'] for r in v)} for k,v in sorted(periods.items())]
            tables[level+period_type+'均'] = values
            if period_type=='月':
                monthly = values
        annual = defaultdict(list)
        for row in monthly:
            annual[(row['始发地'],row['目的城市'],row['货物'])].append(row)
        tables[level+'年度汇总'] = [{'统计年度':cycle['name'],'开始日期':start,'截止日期':end,
            '始发地':k[0],'目的城市':k[1],'货物':k[2],'平均运价':round(mean(r['平均运价'] for r in v),3),
            '最低运价':min(r['最低运价'] for r in v),'最高运价':max(r['最高运价'] for r in v),
            '数据月份数':len(v),'数据天数':sum(r['数据天数'] for r in v),
            '报价数量':sum(r['报价数量'] for r in v)} for k,v in sorted(annual.items())]
    return tables


def write_reports(folder: Path, tables: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    folder.mkdir(parents=True,exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in tables.items():
        sheet = workbook.create_sheet(name)
        headers = list(rows[0]) if rows else ['说明']
        values = [headers]+[[r.get(k,'') for k in headers] for r in rows]
        if not rows:
            values.append(['本统计范围内暂无数据'])
        for row in values:
            sheet.append(row)
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value,str):
                    cell.data_type = 's'  # Never evaluate source messages or user labels as formulas.
                elif isinstance(cell.value,float):
                    cell.number_format = '0.000'
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF',bold=True)
            cell.fill = PatternFill('solid',fgColor='235E66')
        for index, header in enumerate(headers,1):
            sheet.column_dimensions[get_column_letter(index)].width = min(42,max(16,len(header)*2+4))
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        target = folder/(name+'.csv')
        temp = folder/('.'+uuid.uuid4().hex+'.csv')
        try:
            with temp.open('w',encoding='utf-8-sig',newline='') as stream:
                csv.writer(stream).writerows([["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in row] for row in values])
            os.replace(temp,target)
        finally:
            temp.unlink(missing_ok=True)
    target = folder/'年度统计汇总.xlsx'
    temp = folder/('.'+uuid.uuid4().hex+'.xlsx')
    try:
        workbook.save(temp)
        os.replace(temp,target)
    finally:
        workbook.close()
        temp.unlink(missing_ok=True)


def rebuild_manual(profile: dict) -> dict:
    import qq_freight_parser as parser
    from freight_runtime import FreightDatabase
    state = read_year(profile['archive_database'])
    cycle = dict(state['current'])
    db = FreightDatabase(profile['database_file'])
    records = active_records(profile,db.fetch_records())
    records = [parser.normalize_record_cargo(r) for r in records if parser.validate_cargo_route(r)[0]]
    today = parser.trusted_today().isoformat()
    current = [r for r in records if r.get('日期','')<=today]
    if cycle['state']!='active':
        cycle['name']='暂停统计（无当前年度）'
    write_reports(Path(profile['output_dir'])/CURRENT_FOLDER,tables_for(current,cycle,context_now(),today))
    db.export_rejections_csv(str(Path(profile['output_dir'])/parser.REJECTED_CSV))
    return {'active_records':len(current),'historical_active_records':db.count_records(),
            'pending_records':len(records)-len(current),'rejected_records':len(db.fetch_rejections()),
            'excel_updated':True,'as_of_date':today,'manual_year':cycle['name']}


class ArchiveService:
    def __init__(self, profiles: dict, refresh=None) -> None:
        self.profiles = profiles
        self.path = store_path(profiles)
        self.refresh = refresh
        self._worker_lock = threading.Lock()
        self._worker: threading.Thread | None = None
        for profile in profiles.values():
            profile['archive_database'] = str(self.path)

    @contextmanager
    def session(self):
        if not self.path:
            raise ValueError('请先配置至少一个QQ群。')
        self.path.parent.mkdir(parents=True,exist_ok=True)
        with closing(sqlite3.connect(self.path,timeout=30)) as con, con:
            con.row_factory=sqlite3.Row
            con.executescript('''CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
                INSERT OR IGNORE INTO settings VALUES('version','0');
                CREATE TABLE IF NOT EXISTS cycles(id INTEGER PRIMARY KEY,name TEXT NOT NULL,start_day TEXT NOT NULL,
                  end_day TEXT NOT NULL DEFAULT '',state TEXT NOT NULL,opened_at TEXT NOT NULL,closed_at TEXT NOT NULL DEFAULT '',
                  context_json TEXT NOT NULL DEFAULT '{}',error TEXT NOT NULL DEFAULT '');
                CREATE TABLE IF NOT EXISTS exclusions(cycle_id INTEGER,group_id TEXT,record_key TEXT,PRIMARY KEY(cycle_id,group_id,record_key));
                CREATE TABLE IF NOT EXISTS snapshots(cycle_id INTEGER,group_id TEXT,record_key TEXT,record_json TEXT NOT NULL,PRIMARY KEY(cycle_id,group_id,record_key));
                CREATE TABLE IF NOT EXISTS cycle_groups(cycle_id INTEGER,group_id TEXT,name TEXT,output_dir TEXT,tables_json TEXT,PRIMARY KEY(cycle_id,group_id));''')
            yield con

    @contextmanager
    def source_locks(self):
        import qq_freight_parser as parser
        from freight_runtime import FreightDatabase
        with ExitStack() as stack:
            sources = {}
            for key,profile in sorted(self.profiles.items()):
                stack.enter_context(parser.group_output_lock(profile))
                con = stack.enter_context(FreightDatabase(profile['database_file']).session())
                con.execute('BEGIN IMMEDIATE')
                if con.execute("SELECT 1 FROM inbound_events WHERE status IN ('pending','processing') LIMIT 1").fetchone():
                    raise ValueError('有已收到的群消息正在排队处理，请稍后刷新并重新确认年度操作。')
                sources[key]=con
            yield sources

    def status(self) -> dict:
        import qq_freight_parser as parser
        state = read_year(self.path)
        today = parser.trusted_today().isoformat()
        for cycle in state['cycles']:
            cycle.pop('context_json',None)
        state['today']=today
        state['suggested_start']=date(parser.trusted_today().year,1,1).isoformat() if not state['current'] else max(today,(date.fromisoformat(state['current']['end_day'] or today)+timedelta(days=1)).isoformat())
        state['groups']=[{'id':k,'name':v['group_name']} for k,v in self.profiles.items()]
        return state

    def _version(self, con, body: dict) -> None:
        current=int(con.execute("SELECT value FROM settings WHERE key='version'").fetchone()[0])
        if type(body.get('version')) is not int or body['version']!=current:
            raise ValueError('年度状态已改变，请刷新页面后重新确认。')
        if body.get('confirm') is not True:
            raise ValueError('请先确认此年度操作。')

    @staticmethod
    def _bump(con) -> None:
        con.execute("UPDATE settings SET value=CAST(value AS INTEGER)+1 WHERE key='version'")

    def open(self, body: dict) -> dict:
        import qq_freight_parser as parser
        name=str(body.get('name','')).strip()
        if not name or len(name)>60 or any(ord(c)<32 for c in name):
            raise ValueError('年度名称需为1至60个普通字符。')
        start=parse_day(body.get('start_day'))
        today=parser.trusted_today().isoformat()
        if start>today:
            raise ValueError('开始日期不能晚于今天，请到希望开始的日期再开启。')
        with self.source_locks() as sources, self.session() as con:
            con.execute('BEGIN IMMEDIATE')
            self._version(con,body)
            previous=con.execute('SELECT * FROM cycles ORDER BY id DESC LIMIT 1').fetchone()
            if previous and previous['state']!='closed':
                raise ValueError('请先结束当前年度并完成归档。')
            if previous and (start<=previous['end_day'] or start<today):
                raise ValueError('新年度从开启当天开始，且必须晚于上年度截止日期；暂停期不补入。')
            if con.execute('SELECT 1 FROM cycles WHERE name=?',(name,)).fetchone():
                raise ValueError('年度名称已使用，请填写新的名称。')
            cur=con.execute("INSERT INTO cycles(name,start_day,state,opened_at) VALUES(?,?,'active',?)",(name,start,datetime.now().astimezone().isoformat()))
            if previous:
                for key,source in sources.items():
                    con.executemany('INSERT INTO exclusions VALUES(?,?,?)',((cur.lastrowid,key,r[0]) for r in source.execute('SELECT dedup_key FROM freight_records WHERE effective_date>=?',(start,))))
            self._bump(con)
        self._refresh_all()
        return {'message':('统计年度已开启：按所选起始日纳入已有记录。' if not previous else '新统计年度已开启：开启前已入库的记录不补入，后续新记录开始累计。')}

    def preview_close(self, body: dict) -> dict:
        import qq_freight_parser as parser
        state=read_year(self.path)
        cycle=state['current']
        end=parse_day(body.get('end_day'))
        if not cycle or cycle['state']!='active':
            raise ValueError('当前没有正在统计的年度。')
        if not cycle['start_day']<=end<=parser.trusted_today().isoformat():
            raise ValueError('截止日期必须介于本年度开始日和今天之间。')
        counts=[]
        with self.session() as con:
            for key,profile in self.profiles.items():
                records=[json.loads(r['record_json']) for r in selected_rows(profile,cycle,con)]
                records=[r for r in records if r['日期']<=end and parser.validate_cargo_route(r)[0]]
                counts.append({'group_id':key,'name':profile['group_name'],'count':len(records)})
        return {'version':state['version'],'name':cycle['name'],'start_day':cycle['start_day'],'end_day':end,'groups':counts,
                'message':'确认后停止本年度统计，保存关账时快照。预览后的新消息可能使最终数量变化；暂停期间不计入年度。'}

    def close(self, body: dict) -> dict:
        import qq_freight_parser as parser
        preview=self.preview_close(body)
        with self.source_locks() as sources, self.session() as con:
            con.execute('BEGIN IMMEDIATE')
            self._version(con,body)
            cycle=dict(con.execute('SELECT * FROM cycles ORDER BY id DESC LIMIT 1').fetchone())
            if cycle['state']!='active' or cycle['start_day']>preview['end_day']:
                raise ValueError('年度状态已改变，请重新预览。')
            context=context_now()
            cycle.update(end_day=preview['end_day'],state='closing')
            for key,profile in self.profiles.items():
                excluded={r[0] for r in con.execute('SELECT record_key FROM exclusions WHERE cycle_id=? AND group_id=?',(cycle['id'],key))}
                records=[]
                for row in sources[key].execute('SELECT dedup_key,record_json FROM freight_records WHERE effective_date>=? AND effective_date<=? ORDER BY effective_date,dedup_key',(cycle['start_day'],cycle['end_day'])):
                    record=parser.normalize_record_cargo(json.loads(row['record_json']))
                    if row['dedup_key'] in excluded or not parser.validate_cargo_route(record)[0]:
                        continue
                    records.append(record)
                    con.execute('INSERT INTO snapshots VALUES(?,?,?,?)',(cycle['id'],key,row['dedup_key'],json.dumps(record,ensure_ascii=False)))
                tables=tables_for(records,cycle,context,cycle['end_day'])
                con.execute('INSERT INTO cycle_groups VALUES(?,?,?,?,?)',(cycle['id'],key,profile['group_name'],str(Path(profile['output_dir']).resolve()),json.dumps(tables,ensure_ascii=False)))
            con.execute("UPDATE cycles SET end_day=?,state='closing',closed_at=?,context_json=? WHERE id=?",(cycle['end_day'],datetime.now().astimezone().isoformat(),json.dumps(context,ensure_ascii=False),cycle['id']))
            self._bump(con)
        self.resume()
        self._refresh_all()
        return {'message':'年度已停止统计，正在从固定快照生成归档；完成后可开启下一年度。'}

    def retry(self, body: dict) -> dict:
        with self.session() as con:
            con.execute('BEGIN IMMEDIATE')
            self._version(con,body)
            row=con.execute('SELECT * FROM cycles ORDER BY id DESC LIMIT 1').fetchone()
            if not row or row['state']!='error':
                raise ValueError('当前没有失败的归档任务。')
            con.execute("UPDATE cycles SET state='closing',error='' WHERE id=?",(row['id'],))
            self._bump(con)
        self.resume()
        return {'message':'正在使用原关账快照重新生成归档，不重新取数。'}

    def resume(self) -> None:
        state=read_year(self.path)
        if not state['current'] or state['current']['state']!='closing':
            return
        with self._worker_lock:
            if self._worker and self._worker.is_alive():
                return
            self._worker=threading.Thread(target=self._export,args=(state['current']['id'],),daemon=True,name='freight-year-export')
            self._worker.start()

    def _export(self, cycle_id: int) -> None:
        # The transaction is also a cross-process export lock. A restart rolls back
        # to closing; a second worker rechecks state before touching fixed reports.
        with self.session() as con:
            con.execute('BEGIN IMMEDIATE')
            cycle=con.execute('SELECT state FROM cycles WHERE id=?',(cycle_id,)).fetchone()
            if not cycle or cycle['state']!='closing':
                return
            try:
                groups=[dict(r) for r in con.execute('SELECT * FROM cycle_groups WHERE cycle_id=?',(cycle_id,))]
                for group in groups:
                    write_reports(Path(group['output_dir'])/ARCHIVE_FOLDER/str(cycle_id),json.loads(group['tables_json']))
            except Exception as exc:
                con.execute("UPDATE cycles SET state='error',error=? WHERE id=?",('归档文件写入失败，请关闭占用文件、检查磁盘空间后重试。'+type(exc).__name__,cycle_id))
            else:
                con.execute("UPDATE cycles SET state='closed',error='' WHERE id=? AND state='closing'",(cycle_id,))
            self._bump(con)

    def details(self, params: dict) -> dict:
        raw=str(params.get('cycle_id',''))
        if not raw.isdigit() or not self.path or not self.path.exists():
            raise ValueError('归档年度不存在。')
        with self.session() as con:
            cycle=con.execute('SELECT * FROM cycles WHERE id=?',(int(raw),)).fetchone()
            if not cycle or cycle['state']!='closed':
                raise ValueError('该年度尚未完成归档。')
            groups=[]
            for row in con.execute('SELECT * FROM cycle_groups WHERE cycle_id=?',(int(raw),)):
                tables=json.loads(row['tables_json'])
                groups.append({'id':row['group_id'],'name':row['name'],'count':len(tables['报价明细']),
                    'annual':tables['大类年度汇总'],'path':f'{ARCHIVE_FOLDER}/{raw}/年度统计汇总.xlsx',
                    'configured':row['group_id'] in self.profiles})
        return {'name':cycle['name'],'groups':groups}

    def _refresh_all(self) -> None:
        from freight_runtime import FreightDatabase
        for key,profile in self.profiles.items():
            with FreightDatabase(profile['database_file']).session() as con:
                con.execute("INSERT OR REPLACE INTO metadata VALUES('reports_dirty',?)",(uuid.uuid4().hex,))
            if self.refresh:
                self.refresh(key,force_full=True)
