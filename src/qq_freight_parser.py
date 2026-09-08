import re
import csv
import os
import sys
import time
import json
import math
import hashlib
import argparse
import calendar
import queue
import threading
from collections import defaultdict
from statistics import mean
from datetime import date, datetime, timedelta

import pandas as pd
import matplotlib

matplotlib.use("Agg", force=True)
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from lunardate import LunarDate

from freight_runtime import (
    FreightDatabase,
    FreightConfigurationManager,
    NetworkClock,
    RuntimeStatus,
    WindowsOcr,
    configure_logging,
    create_daily_backup,
    normalize_data_lifecycle_config,
    normalize_group_data_lifecycle_configs,
    normalize_excluded_origins,
    prune_daily_backups,
    start_status_dashboard,
    validate_dashboard_host,
    validate_tcp_port,
    validate_websocket_url,
)


QQ_OUTPUT_BUILD_LOCK = threading.RLock()
_GROUP_OUTPUT_LOCKS: dict[str, threading.RLock] = {}
_GROUP_OUTPUT_LOCKS_GUARD = threading.Lock()


def group_output_lock(profile: dict):
    key = os.path.abspath(profile["output_dir"])
    with _GROUP_OUTPUT_LOCKS_GUARD:
        return _GROUP_OUTPUT_LOCKS.setdefault(key, threading.RLock())

# =========================
# 1. 配置区
# =========================

# 贵港群中未写始发地时，默认按贵港统计
DEFAULT_ORIGIN = "贵港"

# 同一价格区内的地点统一归并为标准始发地
ORIGIN_GROUPS = {
    "贵港": ["贵港市", "贵港", "港南区", "港南", "港北区", "港北", "覃塘区", "覃塘", "东龙镇", "东龙"],
    "南宁": ["南宁市", "南宁"],
    "武鸣": ["武鸣区", "武鸣"]
}

# 计划采集的 10 个目的城市及其常见区县
DEST_GROUPS = {
    "佛山": ["佛山", "南海", "顺德", "三水", "高明"],
    "杭州": ["杭州", "余杭", "萧山", "富阳", "临安", "建德"],
    "成都": ["成都", "双流", "龙泉驿", "龙泉", "郫都", "新都", "温江", "青白江"],
    "重庆": ["重庆", "渝北", "江北", "九龙坡", "沙坪坝", "巴南", "涪陵", "永川", "合川", "江津"],
    "长沙": ["长沙", "长沙县", "浏阳", "宁乡", "望城"],
    "南昌": ["南昌", "新建", "进贤", "安义"],
    "武汉": ["武汉", "汉口", "武昌", "汉阳", "东西湖", "蔡甸", "江夏", "黄陂", "新洲"],
    "西安": ["西安", "长安", "临潼", "高陵", "鄠邑", "蓝田", "周至"],
    "福州": ["福州", "福清", "长乐", "闽侯", "连江", "罗源", "闽清", "永泰"],
    "临沂": ["临沂", "兰山", "罗庄", "河东", "郯城", "兰陵", "沂水", "沂南", "平邑", "费县", "蒙阴", "莒南", "临沭"]
}

# 默认货物规则。BOARD_TYPES 保留给旧配置、旧接口和外部脚本兼容使用。
DEFAULT_CARGO_SUBCATEGORY = "大板"
DEFAULT_CARGO_CATEGORY = "板材"
BOARD_TYPES = ["大板", "红板", "拼板"]
CARGO_TYPES = {
    name: {"aliases": [name], "category": DEFAULT_CARGO_CATEGORY}
    for name in BOARD_TYPES
}
# 线路白名单：小类规则优先于大类规则；未配置规则代表该货物不限制线路。
CARGO_ROUTE_SCOPES = {"subcategory": {}, "category": {}}

# 大于这个值，视为整车价，不进入每日平均运价
PRICE_THRESHOLD = 1000

# 所有最终价格统计统一保留三位小数
DECIMAL_PLACES = 3
CSV_FLOAT_FORMAT = f"%.{DECIMAL_PLACES}f"
EXCEL_DECIMAL_FORMAT = "0." + ("0" * DECIMAL_PLACES)
PRICE_OUTPUT_COLUMNS = {
    "最低报价", "最高报价", "平均报价",
    "平均运价", "最低运价", "最高运价",
    "7日均线", "价格指数", "数据覆盖率", "月份覆盖率"
}

# 无效消息关键词
INVALID_LINE_KEYWORDS = [
    "不要标价", "有车了", "求车", "找车", "联系", "有车联系"
]

# 常见前缀
LINE_PREFIX_WORDS = [
    "今明", "明天", "后天", "今天", "急", "现货", "计划"
]

HEADER_PATTERN = re.compile(
    r'^(?P<sender>.+?):\s*(?:(?P<year>\d{4})-)?(?P<month>\d{2})-(?P<day>\d{2})\s+(?P<time>\d{2}:\d{2}:\d{2})$'
)

# 支持：
# # 03-12
# 03-12
# 3/12
# 3月12日
DATE_MARKER_PATTERN = re.compile(
    r'^(?:#\s*)?(?:(?P<year>\d{4})[-/年])?(?P<month>\d{1,2})[-/月](?P<day>\d{1,2})(?:日)?$'
)

# 支持：
# 03-12 武鸣到佛山南海 大板 120
INLINE_DATE_PATTERN = re.compile(
    r'^(?:(?P<year>\d{4})[-/])?(?P<month>\d{1,2})[-/](?P<day>\d{1,2})\s+(?P<body>.+)$'
)

INPUT_FILE = "qq_chat.txt"
FORMATTED_INPUT_FILE = "格式化后_qq_chat.txt"
DETAIL_CSV = "报价明细.csv"
PENDING_CSV = "待生效运价.csv"
DAILY_CSV = "每日平均运价.csv"
HISTORY_DAILY_CSV = "历史每日平均运价.csv"
WEEKLY_CSV = "每周平均运价.csv"
MONTHLY_CSV = "每月平均运价.csv"
YEARLY_CSV = "每年平均运价.csv"
REPORT_XLSX = "物流报价汇总.xlsx"
PENDING_REPORT_XLSX = "物流报价汇总_待更新.xlsx"
CHART_FOLDER = "charts"
SMALL_CATEGORY_OUTPUT_FOLDER = "小类统计"
SMALL_CATEGORY_ARCHIVE_FOLDER = "年度归档"
SMALL_CATEGORY_ARCHIVE_STATE = ".归档状态.json"
SMALL_CATEGORY_ARCHIVE_SCHEMA_VERSION = 1
BUSINESS_ARCHIVE_FOLDER = "年度归档"
LARGE_CATEGORY_ARCHIVE_FOLDER = "大类统计"
LARGE_CATEGORY_ARCHIVE_STATE = ".归档状态.json"
LARGE_CATEGORY_ARCHIVE_SCHEMA_VERSION = 1
OBSERVATION_CONFIG_FILE = "运价观察周期.json"
QQ_LIVE_CONFIG_FILE = "qq_live_config.json"
QQ_LIVE_STATE_FILE = ".qq_live_state.json"
RULES_CONFIG_FILE = "freight_rules.json"
DATABASE_FILE = "运价数据.db"
REJECTED_CSV = "未识别消息.csv"
MAX_SEEN_QQ_MESSAGES = 10000
SUPPORTED_ONEBOT_MESSAGE_POST_TYPES = {"message", "message_sent"}
DETAIL_COLUMNS = [
    "日期", "发布人", "始发地", "目的地", "目的城市", "货物小类", "货物大类",
    "报价文本", "最低报价", "最高报价", "平均报价", "是否整车价", "原始消息",
]

WEEKLY_COLUMNS = [
    "周序号", "周期开始", "周期结束", "始发地", "目的城市", "货物",
    "平均运价", "最低运价", "最高运价", "数据天数", "观察天数",
    "数据覆盖率", "报价数量", "周期状态"
]
MONTHLY_COLUMNS = [
    "月份序号", "月份", "周期开始", "周期结束", "实际观察开始", "实际观察结束",
    "自然月天数", "观察天数", "始发地", "目的城市", "货物", "平均运价",
    "最低运价", "最高运价", "数据天数", "数据覆盖率", "报价数量", "周期状态"
]
YEARLY_COLUMNS = [
    "观察年序号", "观察年", "周期开始", "周期结束", "实际观察结束", "始发地",
    "目的城市", "货物", "平均运价", "最低运价", "最高运价", "数据月份数",
    "已过月份数", "月份覆盖率", "数据天数", "报价数量", "周期状态"
]

DEBUG = False
CLOCK = NetworkClock()
LOGGER = None


def application_dir() -> str:
    """返回源码运行或PyInstaller发行版的可写应用目录。"""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def trusted_today() -> date:
    return CLOCK.today()


# =========================
# 2. 调试输出
# =========================

def debug(msg):
    if DEBUG:
        print(msg)


# =========================
# 3. 文本处理
# =========================

def normalize_text(text: str) -> str:
    text = text.strip()
    text = text.replace("　", " ")
    text = text.translate(str.maketrans("０１２３４５６７８９／．－", "0123456789/.-"))
    text = text.replace("→", "到").replace("->", "到")
    text = re.sub(r"(?<=\d)\s*[～~—–]\s*(?=\d)", "-", text)
    text = re.sub(r"(?<=\d)\s*([/-])\s*(?=\d)", r"\1", text)
    text = re.sub(r"\s+", " ", text)
    return text


def preprocess_line(line: str) -> str:
    text = line.strip()
    for w in LINE_PREFIX_WORDS:
        if text.startswith(w):
            text = text[len(w):].strip()
    return text


# =========================
# 4. 文件处理
# =========================

def delete_generated_outputs():
    old_files = [
        DETAIL_CSV, PENDING_CSV, DAILY_CSV, HISTORY_DAILY_CSV,
        WEEKLY_CSV, MONTHLY_CSV, YEARLY_CSV,
    ]
    for file in old_files:
        if os.path.exists(file):
            try:
                os.remove(file)
                print(f"已删除旧文件: {file}")
            except Exception as e:
                print(f"删除失败 {file}: {e}")

    if os.path.exists(CHART_FOLDER):
        for name in os.listdir(CHART_FOLDER):
            path = os.path.join(CHART_FOLDER, name)
            if os.path.isfile(path):
                try:
                    os.remove(path)
                except Exception as e:
                    print(f"删除旧图片失败 {path}: {e}")
    else:
        os.makedirs(CHART_FOLDER, exist_ok=True)


# =========================
# 5. 消息头解析
# =========================

def parse_header(line: str):
    m = HEADER_PATTERN.match(line)
    if not m:
        return None
    year = int(m.group("year") or trusted_today().year)
    month = int(m.group("month"))
    day = int(m.group("day"))
    try:
        parsed_date = date(year, month, day)
    except ValueError:
        return None
    return {
        "sender": m.group("sender"),
        "date": parsed_date.isoformat(),
        "time": m.group("time")
    }


def parse_date_marker(line: str):
    m = DATE_MARKER_PATTERN.match(line.strip())
    if not m:
        return None
    year = int(m.group("year") or trusted_today().year)
    month = int(m.group("month"))
    day = int(m.group("day"))
    try:
        return date(year, month, day).isoformat()
    except ValueError:
        return None


def parse_inline_date_line(line: str):
    m = INLINE_DATE_PATTERN.match(line.strip())
    if not m:
        return None
    year = int(m.group("year") or trusted_today().year)
    month = int(m.group("month"))
    day = int(m.group("day"))
    body = m.group("body").strip()
    try:
        parsed_date = date(year, month, day).isoformat()
    except ValueError:
        return None
    return parsed_date, body


RELATIVE_DATE_PATTERNS = [
    (re.compile(r"明天\s*(?:和|及|与|、|/)?\s*后天|明后(?:两天|天)?"), (1, 2)),
    (re.compile(r"今明(?:两天)?"), (0, 1)),
    (re.compile(r"后天"), (2,)),
    (re.compile(r"明天|明日"), (1,)),
    (re.compile(r"今天|今日"), (0,)),
]

WEEKDAY_MAP = {"一": 0, "二": 1, "三": 2, "四": 3, "五": 4, "六": 5, "日": 6, "天": 6}
EXPLICIT_CHINESE_DATE_PATTERN = re.compile(
    r"(?:(?P<year>\d{4})年)?(?P<month>\d{1,2})月(?P<day>\d{1,2})日?"
)
EXPLICIT_SLASH_DATE_PATTERN = re.compile(
    r"(?<![\d./-])(?:(?P<year>\d{4})[-/])?"
    r"(?P<month>\d{1,2})[-/](?P<day>\d{1,2})(?![\d./-])"
)
WEEKDAY_PATTERN = re.compile(
    r"(?P<prefix>下周|下星期|下礼拜|本周|本星期|本礼拜|周|星期|礼拜)"
    r"(?P<weekday>[一二三四五六日天])"
)


def load_rules_config(config_path: str = RULES_CONFIG_FILE) -> dict:
    """加载可编辑业务规则；文件不存在时继续使用代码内置默认值。"""
    global DEFAULT_ORIGIN, ORIGIN_GROUPS, DEST_GROUPS, BOARD_TYPES
    global DEFAULT_CARGO_SUBCATEGORY, DEFAULT_CARGO_CATEGORY
    global CARGO_TYPES, CARGO_ROUTE_SCOPES
    global PRICE_THRESHOLD, INVALID_LINE_KEYWORDS, RELATIVE_DATE_PATTERNS

    if not os.path.exists(config_path):
        return {}
    with open(config_path, "r", encoding="utf-8") as stream:
        config = json.load(stream)

    if isinstance(config.get("origin_groups"), dict) and config["origin_groups"]:
        ORIGIN_GROUPS = {
            str(name): [str(alias) for alias in aliases]
            for name, aliases in config["origin_groups"].items()
        }
    if isinstance(config.get("destination_groups"), dict) and config["destination_groups"]:
        DEST_GROUPS = {
            str(name): [str(alias) for alias in aliases]
            for name, aliases in config["destination_groups"].items()
        }
    configured_board_types = config.get("board_types")
    if isinstance(configured_board_types, list) and configured_board_types:
        BOARD_TYPES = [
            str(value).strip()
            for value in configured_board_types
            if str(value).strip()
        ]

    default_category = str(
        config.get("default_cargo_category", "板材")
    ).strip() or "板材"
    configured_cargo_types = config.get("cargo_types")
    cargo_types = {}
    if isinstance(configured_cargo_types, list) and configured_cargo_types:
        for item in configured_cargo_types:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name", "")).strip()
            if not name:
                continue
            aliases = item.get("aliases", [])
            if not isinstance(aliases, list):
                aliases = []
            clean_aliases = [name]
            for alias in aliases:
                alias = str(alias).strip()
                if alias and alias not in clean_aliases:
                    clean_aliases.append(alias)
            cargo_types[name] = {
                "aliases": clean_aliases,
                "category": str(item.get("category", "")).strip() or default_category,
            }
    if not cargo_types:
        cargo_types = {
            name: {"aliases": [name], "category": default_category}
            for name in BOARD_TYPES
        }

    DEFAULT_CARGO_CATEGORY = default_category
    CARGO_TYPES = cargo_types
    BOARD_TYPES = list(cargo_types)
    configured_default_cargo = str(
        config.get("default_cargo_subcategory", "大板")
    ).strip()
    DEFAULT_CARGO_SUBCATEGORY = (
        configured_default_cargo
        if configured_default_cargo in CARGO_TYPES
        else next(iter(CARGO_TYPES))
    )

    scopes = {"subcategory": {}, "category": {}}
    configured_scopes = config.get("cargo_route_scopes", [])
    if isinstance(configured_scopes, list):
        for item in configured_scopes:
            if not isinstance(item, dict):
                continue
            level = str(item.get("level", "")).strip()
            cargo = str(item.get("cargo", "")).strip()
            if level not in scopes or not cargo:
                continue
            routes = set()
            for route in item.get("routes", []):
                if not isinstance(route, dict):
                    continue
                origin = str(route.get("origin", "")).strip()
                destination = str(route.get("destination", "")).strip()
                if origin and destination:
                    routes.add((origin, destination))
            if routes:
                scopes[level][cargo] = routes
    CARGO_ROUTE_SCOPES = scopes
    if isinstance(config.get("invalid_line_keywords"), list):
        INVALID_LINE_KEYWORDS = [str(value) for value in config["invalid_line_keywords"]]
    if config.get("price_threshold") is not None:
        PRICE_THRESHOLD = float(config["price_threshold"])
    configured_default = str(config.get("default_origin", DEFAULT_ORIGIN)).strip()
    if configured_default in ORIGIN_GROUPS:
        DEFAULT_ORIGIN = configured_default

    aliases = config.get("relative_date_aliases")
    if isinstance(aliases, dict):
        def compile_aliases(name: str, fallback: list[str]) -> re.Pattern:
            values = aliases.get(name, fallback)
            values = sorted({str(value) for value in values if str(value)}, key=len, reverse=True)
            return re.compile("|".join(re.escape(value) for value in values))

        RELATIVE_DATE_PATTERNS = [
            (compile_aliases("tomorrow_and_day_after", ["明后天"]), (1, 2)),
            (compile_aliases("today_and_tomorrow", ["今明"]), (0, 1)),
            (compile_aliases("day_after_tomorrow", ["后天"]), (2,)),
            (compile_aliases("tomorrow", ["明天", "明日"]), (1,)),
            (compile_aliases("today", ["今天", "今日"]), (0,)),
        ]
    return config


def iter_cargo_aliases():
    """按全局最长别名优先返回（别名、标准小类），避免短名称抢先匹配。"""
    matches = []
    for name, item in CARGO_TYPES.items():
        for alias in item.get("aliases", []):
            alias = str(alias).strip()
            if alias:
                matches.append((alias, name))
    return sorted(matches, key=lambda value: (-len(value[0]), value[0], value[1]))


def cargo_category(cargo_subcategory: str) -> str:
    item = CARGO_TYPES.get(str(cargo_subcategory or "").strip(), {})
    return str(item.get("category", "")).strip() or DEFAULT_CARGO_CATEGORY


def normalize_record_cargo(record: dict) -> dict:
    """按当前配置补齐/重算大类，让历史数据在重建时也能采用最新归类。"""
    normalized = dict(record)
    subcategory = str(normalized.get("货物小类", "")).strip()
    if subcategory:
        normalized["货物大类"] = cargo_category(subcategory)
    return normalized


def validate_cargo_route(record: dict) -> tuple[bool, str]:
    """校验货物线路白名单；小类有专属规则时覆盖其大类规则。"""
    subcategory = str(record.get("货物小类", "")).strip()
    category = cargo_category(subcategory)
    route = (
        str(record.get("始发地", "")).strip(),
        str(record.get("目的城市", "")).strip(),
    )
    scoped_level = (
        "subcategory"
        if subcategory in CARGO_ROUTE_SCOPES["subcategory"]
        else "category"
    )
    scoped_cargo = subcategory if scoped_level == "subcategory" else category
    allowed_routes = CARGO_ROUTE_SCOPES[scoped_level].get(scoped_cargo)
    if not allowed_routes or route in allowed_routes:
        return True, ""
    level_name = "小类" if scoped_level == "subcategory" else "大类"
    return (
        False,
        f"货物{level_name}“{scoped_cargo}”未开放线路：{route[0]}到{route[1]}",
    )


def expand_relative_freight_dates(line: str, base_date_text: str) -> list[tuple[str, str]]:
    """把相对日期报价展开为一个或多个生效日期，并移除日期提示词。"""
    try:
        base_date = datetime.strptime(base_date_text, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return []

    normalized_line = normalize_text(line)
    for pattern, offsets in RELATIVE_DATE_PATTERNS:
        match = pattern.search(normalized_line)
        if not match:
            continue
        freight_text = normalize_text(
            normalized_line[:match.start()] + " " + normalized_line[match.end():]
        )
        return [
            ((base_date + timedelta(days=offset)).isoformat(), freight_text)
            for offset in offsets
        ]

    for pattern in (EXPLICIT_CHINESE_DATE_PATTERN, EXPLICIT_SLASH_DATE_PATTERN):
        matches = list(pattern.finditer(normalized_line))
        if pattern is EXPLICIT_SLASH_DATE_PATTERN:
            # 裸月/日仅在行首或明确装货日期上下文识别；不能把车长或报价当日期。
            matches = [m for m in matches if m.group("year") or (
                m.start() == 0
                or re.search(r"(?:日期|装货|发货)[:： ]*$", normalized_line[:m.start()])
                or re.match(r"\s*(?:日|号|装货|发货|装车)", normalized_line[m.end():])
            )]
        if not matches:
            continue
        match = matches[0]
        try:
            explicit_date = date(
                int(match.group("year") or base_date.year),
                int(match.group("month")),
                int(match.group("day")),
            )
        except ValueError:
            return []
        freight_text = normalize_text(
            normalized_line[:match.start()] + " " + normalized_line[match.end():]
        )
        return [(explicit_date.isoformat(), freight_text)]

    # 只有“22号”时无法可靠确定月份，不要默默记到消息当天。
    if re.search(r"(?<!\d)\d{1,2}号", normalized_line):
        return []

    weekday_match = WEEKDAY_PATTERN.search(normalized_line)
    if weekday_match:
        target_weekday = WEEKDAY_MAP[weekday_match.group("weekday")]
        prefix = weekday_match.group("prefix")
        week_start = base_date - timedelta(days=base_date.weekday())
        if prefix.startswith("下"):
            effective_date = week_start + timedelta(days=7 + target_weekday)
        elif prefix.startswith("本"):
            effective_date = week_start + timedelta(days=target_weekday)
        else:
            effective_date = base_date + timedelta(
                days=(target_weekday - base_date.weekday()) % 7
            )
        freight_text = normalize_text(
            normalized_line[:weekday_match.start()] + " " + normalized_line[weekday_match.end():]
        )
        return [(effective_date.isoformat(), freight_text)]

    return [(base_date.isoformat(), normalized_line)]


# =========================
# 6. 筛选规则
# =========================

def is_invalid_line(line: str) -> bool:
    return any(k in line for k in INVALID_LINE_KEYWORDS)


def normalize_origin(origin: str):
    """把贵港下辖片区等别名统一为标准始发地。"""
    text = re.sub(r"\s+", "", origin or "")
    for city, aliases in ORIGIN_GROUPS.items():
        for alias in sorted(aliases, key=len, reverse=True):
            if alias in text:
                return city
    return None


def extract_origin_prefix(text: str):
    """识别未使用“到”字时写在行首的始发地。"""
    compact = text.lstrip()
    for city, aliases in ORIGIN_GROUPS.items():
        for alias in sorted(aliases, key=len, reverse=True):
            if compact.startswith(alias):
                remaining = compact[len(alias):].lstrip(" 市区县镇-—至到:：,，")
                return city, remaining
    return None, text


def normalize_destination_city(dest: str):
    for city, keywords in DEST_GROUPS.items():
        for keyword in sorted(keywords, key=len, reverse=True):
            if keyword in dest:
                return city
    return None


# =========================
# 7. 路线提取
# =========================

def extract_route(line: str, default_origin: str = DEFAULT_ORIGIN):
    """
    支持：
    武鸣到佛山南海 大板 120
    武鸣到佛山南海九江大板125/130
    南宁到成都双流大板210
    """
    text = preprocess_line(line)

    if "到" in text:
        left, right = text.split("到", 1)
        # Only an absent origin may inherit the group's default. An explicitly
        # supplied out-of-scope origin must still be rejected, never reclassified.
        origin = normalize_origin(left) if left.strip() else normalize_origin(default_origin)
        if not origin:
            return None, f"始发地不在范围内: {left.strip()}"
        rest_all = right.strip()
    else:
        origin, rest_all = extract_origin_prefix(text)
        if not origin:
            origin = default_origin
            rest_all = text

    # 选择最靠前的目的地关键词；位置相同时优先匹配更长的区县名。
    destination_matches = []
    for city, keywords in DEST_GROUPS.items():
        for keyword in keywords:
            idx = rest_all.find(keyword)
            if idx != -1:
                destination_matches.append((idx, -len(keyword), city, keyword))

    if not destination_matches:
        return None, f"目的地不在范围内: {rest_all}"

    dest_start, _, matched_city, matched_keyword = min(destination_matches)

    # 目的地原文截止到货物名称或第一个价格数字，保留“佛山南海九江”等明细地点。
    end_candidates = []
    for cargo_alias, _cargo_name in iter_cargo_aliases():
        cargo_pos = rest_all.find(cargo_alias, dest_start + len(matched_keyword))
        if cargo_pos != -1:
            end_candidates.append(cargo_pos)

    price_match = re.search(r"\d", rest_all[dest_start + len(matched_keyword):])
    if price_match:
        end_candidates.append(dest_start + len(matched_keyword) + price_match.start())

    dest_end = min(end_candidates) if end_candidates else len(rest_all)
    matched_dest_raw = rest_all[dest_start:dest_end].strip(" -—:：,，、;；")
    matched_dest_raw = re.sub(r"\s+", "", matched_dest_raw) or matched_keyword
    rest = rest_all[dest_end:].strip()

    return (origin, matched_dest_raw, matched_city, rest), None


# =========================
# 8. 货物提取
# =========================

def extract_cargo(rest: str):
    for alias, cargo in iter_cargo_aliases():
        if alias in rest:
            return cargo
    return None


# =========================
# 9. 报价提取
# =========================

def is_vehicle_size_number(num_str: str) -> bool:
    try:
        val = float(num_str)
    except ValueError:
        return False

    vehicle_sizes = [9.6, 13, 13.5, 13.75, 14, 16, 17.5]
    return any(abs(val - x) < 0.01 for x in vehicle_sizes)


def quote_rejection_reason(text: str) -> str:
    """Check the complete line before route/price extraction can discard context."""
    text = normalize_text(text)
    volume_unit = r"(?:立方米?|方|m\s*(?:\^?\s*3|³)|㎥|cbm)"
    volume_pricing = (
        rf"(?:/\s*{volume_unit}|每\s*{volume_unit}"
        rf"|(?:元|块)\s*(?:一\s*)?{volume_unit}"
        rf"|一\s*{volume_unit}\s*[:：]?\s*\d"
        rf"|\d\s*一\s*{volume_unit}|按\s*{volume_unit}"
        r"|(?<!木)(?:立方|方)价)"
    )
    if re.search(volume_pricing, text, flags=re.IGNORECASE):
        return "按方/立方米计价的运价不采集"
    # A hyphen between positive numbers is a range, not a unary minus.
    # normalize_text has already joined valid ranges such as "120 - 125".
    if re.search(r"(?<![\d.])[-−﹣—–]\s*\d|负\s*\d", text):
        return "报价含负数或疑似负价，拒绝采集，请核实后重新发布"
    return ""


def parse_price_text(price_text: str):
    price_text = normalize_text(price_text)
    # Never recover unsigned fragments from malformed/negative price tokens.
    if not re.fullmatch(r"\d+(?:\.\d+)?(?:[/-]\d+(?:\.\d+)?)*", price_text):
        return None
    nums = re.findall(r"\d+(?:\.\d+)?", price_text)
    values = [float(x) for x in nums]
    if not all(math.isfinite(value) and value > 0 for value in values):
        return None
    return {
        "price_text": price_text,
        "price_min": round(min(values), DECIMAL_PLACES),
        "price_max": round(max(values), DECIMAL_PLACES),
        "price_avg": round(sum(values) / len(values), DECIMAL_PLACES)
    }


def extract_price_with_reason(rest: str):
    """
    支持：
    118
    120/125
    160-170
    3400/3450
    过滤 13.75 / 17.5 等车型长度
    """
    text = normalize_text(rest)
    rejection = quote_rejection_reason(text)
    if rejection:
        return None, rejection
    text = re.sub(r"(?<!\d)\d{7,}(?!\d)", " ", text)
    if re.search(r"\d\s*(?:加|另加|另收|另付)\s*\d", text):
        return None, "价格含加价说明，需人工确认最终运价"

    candidates = []
    unexplained_numbers = []
    pattern = r"\d+(?:\.\d+)?(?:[/-]\d+(?:\.\d+)?)*"
    for m in re.finditer(pattern, text):
        token = m.group()
        nums = re.findall(r"\d+(?:\.\d+)?", token)

        # “元/吨”“一吨”是价格单位；紧跟数字的“吨、米、装、卸”等是货量/车长。
        suffix = text[m.end():].lstrip()
        prefix = text[:m.start()].rstrip()
        if re.match(r"(?:吨(?!位|包)|公斤|千克|公里|千米|米|方|立方|装|卸|车(?!长|型)|趟|排|件|包|张|号|天|小时)", suffix):
            continue
        if re.search(r"(?:(?<![大有])吨位|重量|车长|距离|装卸费|信息费|定金)\s*[:：]?\s*$", prefix):
            continue

        if (nums and all(is_vehicle_size_number(x) for x in nums)
                and not re.match(r"(?:元|块|一吨|/\s*吨)", suffix)):
            continue

        parsed = parse_price_text(token)
        if parsed and parsed["price_avg"] >= 20:
            candidates.append(parsed)
        else:
            unexplained_numbers.append(token)

    if not candidates:
        return None, "未识别到有效价格"
    if unexplained_numbers:
        return None, "存在未说明用途的数字或疑似拆分价格，需人工确认"
    if len({item["price_text"] for item in candidates}) > 1:
        return None, "存在多个未标明关系的价格数字，需人工确认"
    return candidates[0], ""


def extract_price(rest: str):
    return extract_price_with_reason(rest)[0]


# =========================
# 10. 自动整理输入
# =========================

def format_freight_line_with_reason(
    line: str,
    default_origin: str = DEFAULT_ORIGIN,
) -> tuple[str | None, str]:
    """整理单行运价，并在不改变业务规则的前提下返回明确拒绝原因。"""
    line = normalize_text(line)
    line = preprocess_line(line)

    if not line:
        return None, "消息行为空"

    rejection = quote_rejection_reason(line)
    if rejection:
        return None, rejection

    invalid_keyword = next(
        (keyword for keyword in INVALID_LINE_KEYWORDS if keyword in line),
        "",
    )
    if invalid_keyword:
        return None, f"命中无效关键词: {invalid_keyword}"

    route, err = extract_route(line, default_origin)
    if not route:
        return None, err or "未识别到有效线路"

    origin, dest_raw, _dest_city, rest = route
    price_info, price_reason = extract_price_with_reason(rest)
    if not price_info:
        return None, price_reason

    cargo = extract_cargo(rest)
    if cargo is None and "到" in line:
        cargo = extract_cargo(line.split("到", 1)[0])
    if cargo is None:
        cargo = DEFAULT_CARGO_SUBCATEGORY

    return f"{origin}到{dest_raw} {cargo} {price_info['price_text']}", ""


def auto_format_freight_line(line: str, default_origin: str = DEFAULT_ORIGIN):
    """
    把原始货运文本整理成程序容易识别的格式：
    始发地到目的地 货物 价格

    规则：
    - 用现有 extract_route / extract_price 识别
    - 若没写货物种类，默认按 大板 处理
    """
    formatted, _reason = format_freight_line_with_reason(line, default_origin)
    return formatted


def auto_format_input_file(
    raw_file_path: str,
    formatted_output_file: str = FORMATTED_INPUT_FILE,
    default_origin: str = DEFAULT_ORIGIN
):
    """
    把原始 qq_chat.txt 自动整理成程序可识别格式，并输出到 FORMATTED_INPUT_FILE
    """
    if not os.path.exists(raw_file_path):
        print(f"未找到输入文件: {raw_file_path}")
        return raw_file_path

    output_lines = []
    current_date = None
    current_sender = "系统整理"
    current_time = "00:00:00"

    def append_freight_line(raw_freight_line: str, base_date_text: str) -> None:
        for effective_date, freight_text in expand_relative_freight_dates(
            raw_freight_line,
            base_date_text
        ):
            formatted = auto_format_freight_line(freight_text, default_origin)
            if not formatted:
                continue
            output_lines.append(f"{current_sender}: {effective_date} {current_time}")
            output_lines.append(formatted)

    with open(raw_file_path, "r", encoding="utf-8") as f:
        for raw_line in f:
            line = normalize_text(raw_line)
            if not line:
                continue

            # 1. 原始 QQ 头
            header = parse_header(line)
            if header:
                current_date = header["date"]
                current_sender = header["sender"]
                current_time = header["time"]
                continue

            # 2. 日期块
            marker_date = parse_date_marker(line)
            if marker_date:
                current_date = marker_date
                current_time = "00:00:00"
                continue

            # 3. 行内日期
            inline_date = parse_inline_date_line(line)
            if inline_date:
                current_date, body = inline_date
                current_time = "00:00:00"
                append_freight_line(body, current_date)
                continue

            # 4. 普通货运行
            if current_date:
                append_freight_line(line, current_date)

    with open(formatted_output_file, "w", encoding="utf-8") as f:
        for line in output_lines:
            f.write(line + "\n")

    print(f"已自动整理输入文件：{formatted_output_file}")
    return formatted_output_file


# =========================
# 11. 单条消息解析
# =========================

def parse_freight_line(
    line: str,
    current_date: str,
    current_sender: str,
    default_origin: str = DEFAULT_ORIGIN
):
    if quote_rejection_reason(line):
        return None
    if is_invalid_line(line):
        return None

    route, err = extract_route(line, default_origin)
    if not route:
        debug(f"跳过 | {err} | {line}")
        return None

    origin, dest_raw, dest_city, rest = route

    price_info = extract_price(rest)
    if not price_info:
        return None

    cargo = extract_cargo(rest)
    if cargo is None:
        return None

    is_full_truck_price = price_info["price_avg"] > PRICE_THRESHOLD

    return {
        "日期": current_date,
        "发布人": current_sender,
        "始发地": origin,
        "目的地": dest_raw,
        "目的城市": dest_city,
        "货物小类": cargo,
        "货物大类": cargo_category(cargo),
        "报价文本": price_info["price_text"],
        "最低报价": price_info["price_min"],
        "最高报价": price_info["price_max"],
        "平均报价": price_info["price_avg"],
        "是否整车价": "是" if is_full_truck_price else "否",
        "原始消息": line
    }


# =========================
# 12. 读取聊天文件
# =========================

def parse_chat_file(file_path: str, default_origin: str = DEFAULT_ORIGIN):
    if not os.path.exists(file_path):
        print(f"未找到输入文件: {file_path}")
        return []

    records = []
    current_sender = None
    current_date = None

    with open(file_path, "r", encoding="utf-8") as f:
        for raw_line in f:
            line = normalize_text(raw_line)
            if not line:
                continue

            header = parse_header(line)
            if header:
                current_sender = header["sender"]
                current_date = header["date"]
                continue

            if not current_date:
                continue

            rec = parse_freight_line(line, current_date, current_sender, default_origin)
            if rec:
                allowed, _reason = validate_cargo_route(rec)
                if allowed:
                    records.append(rec)

    return records


# =========================
# 13. 同日同线路报价去重
# =========================


def normalize_location_for_dedup(value: str) -> str:
    """统一地点写法，消除空格及“市/区/县/镇”等行政区后缀差异。"""
    text = re.sub(r"\s+", "", str(value or ""))
    for suffix in ["省", "市", "区", "县", "镇"]:
        text = text.replace(suffix, "")
    return text


def build_destination_dedup_key(record: dict) -> tuple[str, str]:
    """把“佛山南海”和“南海”等同一目的地写法归为同一个键。"""
    city = normalize_location_for_dedup(record.get("目的城市", ""))
    destination = normalize_location_for_dedup(record.get("目的地", ""))
    if city and destination.startswith(city):
        destination = destination[len(city):]
    return city, destination


def build_price_dedup_key(record: dict) -> tuple[float, ...]:
    """按完整报价数字查重，并兼容 120/125、120-125 等不同分隔写法。"""
    numbers = re.findall(r"\d+(?:\.\d+)?", str(record.get("报价文本", "")))
    if numbers:
        return tuple(sorted(round(float(number), DECIMAL_PLACES) for number in numbers))

    # 兼容缺少原始报价文本的旧记录。
    fallback = [
        record.get("最低报价"),
        record.get("最高报价"),
        record.get("平均报价")
    ]
    return tuple(round(float(value), DECIMAL_PLACES) for value in fallback if value is not None)


def build_record_content_dedup_key(record: dict) -> str:
    """正式数据和未来缓存共用同一内容查重键，发布人不参与。"""
    normalized_origin = normalize_origin(record.get("始发地", ""))
    if not normalized_origin:
        normalized_origin = normalize_location_for_dedup(record.get("始发地", ""))
    key_payload = {
        "date": str(record.get("日期", "")),
        "origin": normalized_origin,
        "destination": build_destination_dedup_key(record),
        "cargo": str(record.get("货物小类", "")).strip(),
        "prices": build_price_dedup_key(record),
    }
    serialized = json.dumps(key_payload, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


def deduplicate_records(records: list[dict]) -> list[dict]:
    """
    去重规则：
    日期 + 标准化始发地 + 具体目的地 + 货物小类 + 完整报价数字

    发布人不参与查重；同一天不同人转发相同货运信息时，只保留最先出现的一条。
    """
    unique_map = {}
    for r in records:
        key = build_record_content_dedup_key(r)
        if key not in unique_map:
            unique_map[key] = r

    return list(unique_map.values())


# =========================
# 14. 当次每日平均运价
# =========================

def summarize_daily_quote(records):
    """
    小类按配置归入大类，再按 日期 + 始发地 + 目的城市 + 货物大类 分组。
    整车价 (>1000) 不参与统计
    """
    groups = defaultdict(list)

    for r in records:
        if r["平均报价"] > PRICE_THRESHOLD:
            continue
        category = cargo_category(r.get("货物小类", ""))
        key = (r["日期"], r["始发地"], r["目的城市"], category)
        groups[key].append(r["平均报价"])

    result = []
    for (date, origin, city, category), prices in groups.items():
        result.append({
            "日期": date,
            "始发地": origin,
            "目的城市": city,
            "货物": category,
            "平均运价": round(mean(prices), DECIMAL_PLACES),
            "最低运价": round(min(prices), DECIMAL_PLACES),
            "最高运价": round(max(prices), DECIMAL_PLACES),
            "报价数量": len(prices)
        })

    result.sort(key=lambda x: (
        datetime.strptime(x["日期"], "%Y-%m-%d"),
        x["始发地"],
        x["目的城市"],
        x["货物"],
    ))
    return result


def summarize_daily_quote_by_subcategory(records):
    """按货物小类生成日统计；现有大类统计保持不变。"""
    groups = defaultdict(list)

    for record in records:
        if record["平均报价"] > PRICE_THRESHOLD:
            continue
        subcategory = str(
            record.get("货物小类", DEFAULT_CARGO_SUBCATEGORY)
        ).strip() or DEFAULT_CARGO_SUBCATEGORY
        key = (
            record["日期"],
            record["始发地"],
            record["目的城市"],
            subcategory,
        )
        groups[key].append(record["平均报价"])

    result = []
    for (record_date, origin, city, subcategory), prices in groups.items():
        result.append({
            "日期": record_date,
            "始发地": origin,
            "目的城市": city,
            "货物": subcategory,
            "平均运价": round(mean(prices), DECIMAL_PLACES),
            "最低运价": round(min(prices), DECIMAL_PLACES),
            "最高运价": round(max(prices), DECIMAL_PLACES),
            "报价数量": len(prices),
        })

    result.sort(key=lambda item: (
        datetime.strptime(item["日期"], "%Y-%m-%d"),
        item["始发地"],
        item["目的城市"],
        item["货物"],
    ))
    return result


# =========================
# 15. 增加历史指标
# =========================

def enrich_history_indicators(history_df):
    if history_df is None or history_df.empty:
        return history_df

    df = history_df.copy()
    df["日期排序"] = pd.to_datetime(df["日期"], format="%Y-%m-%d", errors="coerce")

    result_list = []

    if "始发地" not in df.columns:
        df["始发地"] = "历史未区分"

    if "货物" not in df.columns:
        df["货物"] = DEFAULT_CARGO_CATEGORY
    df["货物"] = df["货物"].fillna(DEFAULT_CARGO_CATEGORY)

    for (origin, city, cargo), sub in df.groupby(["始发地", "目的城市", "货物"]):
        sub = sub.sort_values("日期排序").copy()
        sub["7日均线"] = sub["平均运价"].rolling(window=7, min_periods=1).mean().round(DECIMAL_PLACES)

        first_price = sub["平均运价"].iloc[0]
        if first_price and first_price != 0:
            sub["价格指数"] = (sub["平均运价"] / first_price * 100).round(DECIMAL_PLACES)
        else:
            sub["价格指数"] = 100.00

        result_list.append(sub)

    result = pd.concat(result_list, ignore_index=True)
    result = result.drop(columns=["日期排序"])
    return result


# =========================
# 16. 历史日均价内存计算
# =========================

def merge_daily_history(new_summary, minimum_date: date | None = None):
    """生成供Excel和周期统计使用的日均历史，不再落地独立CSV。"""
    new_df = pd.DataFrame(new_summary)
    base_columns = [
        "日期", "始发地", "目的城市", "货物", "平均运价", "最低运价", "最高运价", "报价数量"
    ]

    if minimum_date is not None and not new_df.empty:
        new_dates = pd.to_datetime(new_df["日期"], errors="coerce").dt.date
        new_df = new_df[new_dates >= minimum_date].copy()

    if new_df.empty:
        return enrich_history_indicators(pd.DataFrame(columns=base_columns))

    for column in base_columns:
        if column not in new_df.columns:
            new_df[column] = None
    history_df = new_df[base_columns].drop_duplicates(
        subset=["日期", "始发地", "目的城市", "货物"],
        keep="last",
    )
    history_df["日期排序"] = pd.to_datetime(
        history_df["日期"], format="%Y-%m-%d", errors="coerce"
    )
    history_df = history_df.sort_values(
        ["日期排序", "始发地", "目的城市", "货物"]
    ).drop(columns=["日期排序"])
    return enrich_history_indicators(history_df)


# =========================
# 17. 部署日起的周、月、年统计
# =========================

def add_months(month_date: date, months: int) -> date:
    """在自然月层面移动月份，返回目标月 1 日。"""
    month_index = month_date.year * 12 + (month_date.month - 1) + months
    year, zero_based_month = divmod(month_index, 12)
    return date(year, zero_based_month + 1, 1)


def get_month_end(month_start: date) -> date:
    last_day = calendar.monthrange(month_start.year, month_start.month)[1]
    return date(month_start.year, month_start.month, last_day)


def get_or_create_observation_start(
    config_file: str = OBSERVATION_CONFIG_FILE,
    default_date: date | None = None
) -> date:
    """首次部署时固定观察起点，后续重建不随运行日期漂移。"""
    if os.path.exists(config_file):
        try:
            with open(config_file, "r", encoding="utf-8") as f:
                config = json.load(f)
            return datetime.strptime(
                str(config["observation_start_date"]), "%Y-%m-%d"
            ).date()
        except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
            raise ValueError(
                f"观察周期配置无效: {config_file}；"
                "请确认 observation_start_date 使用 YYYY-MM-DD 格式。"
            ) from exc

    observation_start = default_date or trusted_today()
    config_dir = os.path.dirname(os.path.abspath(config_file))
    os.makedirs(config_dir, exist_ok=True)
    temp_file = config_file + ".tmp"
    config = {
        "observation_start_date": observation_start.isoformat(),
        "week_rule": "从观察起点每7个自然日为一周",
        "month_rule": "按自然月，首月不足整月也计入",
        "year_rule": "从部署所在月起连续12个观察月为一年"
    }
    with open(temp_file, "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)
    os.replace(temp_file, config_file)
    print(f"已建立运价观察周期，起始日期: {observation_start}")
    return observation_start


def prepare_observation_daily_data(
    history_df: pd.DataFrame,
    observation_start: date,
    as_of_date: date
) -> pd.DataFrame:
    if history_df is None or history_df.empty:
        return pd.DataFrame()

    df = history_df.copy()
    if "始发地" not in df.columns:
        df["始发地"] = "历史未区分"
    if "报价数量" not in df.columns:
        df["报价数量"] = 0
    if "货物" not in df.columns:
        df["货物"] = DEFAULT_CARGO_CATEGORY
    df["货物"] = df["货物"].fillna(DEFAULT_CARGO_CATEGORY)

    df["日期排序"] = pd.to_datetime(df["日期"], errors="coerce")
    df = df.dropna(subset=["日期排序", "平均运价"])
    start_timestamp = pd.Timestamp(observation_start)
    end_timestamp = pd.Timestamp(as_of_date)
    df = df[(df["日期排序"] >= start_timestamp) & (df["日期排序"] <= end_timestamp)]

    for column in ["平均运价", "最低运价", "最高运价", "报价数量"]:
        df[column] = pd.to_numeric(df[column], errors="coerce")
    df = df.dropna(subset=["平均运价", "最低运价", "最高运价"])
    df["报价数量"] = df["报价数量"].fillna(0)
    return df


def aggregate_weekly_statistics(
    daily_df: pd.DataFrame,
    observation_start: date,
    as_of_date: date
) -> pd.DataFrame:
    """按部署日起每 7 个自然日聚合，只输出已走满 7 天的周。"""
    if daily_df.empty:
        return pd.DataFrame(columns=WEEKLY_COLUMNS)

    df = daily_df.copy()
    start_timestamp = pd.Timestamp(observation_start)
    df["周序号"] = ((df["日期排序"] - start_timestamp).dt.days // 7 + 1).astype(int)

    rows = []
    group_columns = ["周序号", "始发地", "目的城市", "货物"]
    for (week_number, origin, city, cargo), sub in df.groupby(group_columns):
        period_start = observation_start + timedelta(days=(int(week_number) - 1) * 7)
        period_end = period_start + timedelta(days=6)
        if period_end > as_of_date:
            continue

        data_days = int(sub["日期排序"].dt.date.nunique())
        rows.append({
            "周序号": int(week_number),
            "周期开始": period_start.isoformat(),
            "周期结束": period_end.isoformat(),
            "始发地": origin,
            "目的城市": city,
            "货物": cargo,
            "平均运价": round(float(sub["平均运价"].mean()), DECIMAL_PLACES),
            "最低运价": round(float(sub["最低运价"].min()), DECIMAL_PLACES),
            "最高运价": round(float(sub["最高运价"].max()), DECIMAL_PLACES),
            "数据天数": data_days,
            "观察天数": 7,
            "数据覆盖率": round(data_days / 7 * 100, DECIMAL_PLACES),
            "报价数量": int(sub["报价数量"].sum()),
            "周期状态": "完整"
        })

    return pd.DataFrame(rows, columns=WEEKLY_COLUMNS).sort_values(
        ["周序号", "始发地", "目的城市", "货物"], ignore_index=True
    ) if rows else pd.DataFrame(columns=WEEKLY_COLUMNS)


def aggregate_monthly_statistics(
    daily_df: pd.DataFrame,
    observation_start: date,
    as_of_date: date
) -> pd.DataFrame:
    """按真实自然月聚合，首月和当前月不足整月也保留。"""
    if daily_df.empty:
        return pd.DataFrame(columns=MONTHLY_COLUMNS)

    df = daily_df.copy()
    start_month = date(observation_start.year, observation_start.month, 1)
    df["月份序号"] = (
        (df["日期排序"].dt.year - start_month.year) * 12
        + (df["日期排序"].dt.month - start_month.month)
        + 1
    ).astype(int)

    rows = []
    group_columns = ["月份序号", "始发地", "目的城市", "货物"]
    for (month_number, origin, city, cargo), sub in df.groupby(group_columns):
        month_start = add_months(start_month, int(month_number) - 1)
        month_end = get_month_end(month_start)
        actual_start = max(observation_start, month_start)
        actual_end = min(as_of_date, month_end)
        observation_days = (actual_end - actual_start).days + 1
        data_days = int(sub["日期排序"].dt.date.nunique())
        rows.append({
            "月份序号": int(month_number),
            "月份": month_start.strftime("%Y-%m"),
            "周期开始": month_start.isoformat(),
            "周期结束": month_end.isoformat(),
            "实际观察开始": actual_start.isoformat(),
            "实际观察结束": actual_end.isoformat(),
            "自然月天数": calendar.monthrange(month_start.year, month_start.month)[1],
            "观察天数": observation_days,
            "始发地": origin,
            "目的城市": city,
            "货物": cargo,
            "平均运价": round(float(sub["平均运价"].mean()), DECIMAL_PLACES),
            "最低运价": round(float(sub["最低运价"].min()), DECIMAL_PLACES),
            "最高运价": round(float(sub["最高运价"].max()), DECIMAL_PLACES),
            "数据天数": data_days,
            "数据覆盖率": round(data_days / observation_days * 100, DECIMAL_PLACES),
            "报价数量": int(sub["报价数量"].sum()),
            "周期状态": "完整" if as_of_date >= month_end else "进行中"
        })

    return pd.DataFrame(rows, columns=MONTHLY_COLUMNS).sort_values(
        ["月份序号", "始发地", "目的城市", "货物"], ignore_index=True
    ) if rows else pd.DataFrame(columns=MONTHLY_COLUMNS)


def aggregate_yearly_statistics(
    monthly_df: pd.DataFrame,
    observation_start: date,
    as_of_date: date
) -> pd.DataFrame:
    """每连续 12 个观察月形成一个观察年，年均价等权平均各月均价。"""
    if monthly_df.empty:
        return pd.DataFrame(columns=YEARLY_COLUMNS)

    df = monthly_df.copy()
    df["观察年序号"] = ((df["月份序号"] - 1) // 12 + 1).astype(int)
    start_month = date(observation_start.year, observation_start.month, 1)
    rows = []

    group_columns = ["观察年序号", "始发地", "目的城市", "货物"]
    for (year_number, origin, city, cargo), sub in df.groupby(group_columns):
        year_start_month = add_months(start_month, (int(year_number) - 1) * 12)
        year_end_month = add_months(year_start_month, 11)
        period_start = observation_start if int(year_number) == 1 else year_start_month
        period_end = get_month_end(year_end_month)
        actual_end = min(as_of_date, period_end)
        elapsed_months = min(
            12,
            (actual_end.year - year_start_month.year) * 12
            + actual_end.month - year_start_month.month + 1
        )
        data_months = int(sub["月份"].nunique())
        year_label = (
            f"第{int(year_number)}观察年 "
            f"({year_start_month:%Y-%m}至{year_end_month:%Y-%m})"
        )

        rows.append({
            "观察年序号": int(year_number),
            "观察年": year_label,
            "周期开始": period_start.isoformat(),
            "周期结束": period_end.isoformat(),
            "实际观察结束": actual_end.isoformat(),
            "始发地": origin,
            "目的城市": city,
            "货物": cargo,
            "平均运价": round(float(sub["平均运价"].mean()), DECIMAL_PLACES),
            "最低运价": round(float(sub["最低运价"].min()), DECIMAL_PLACES),
            "最高运价": round(float(sub["最高运价"].max()), DECIMAL_PLACES),
            "数据月份数": data_months,
            "已过月份数": elapsed_months,
            "月份覆盖率": round(data_months / elapsed_months * 100, DECIMAL_PLACES),
            "数据天数": int(sub["数据天数"].sum()),
            "报价数量": int(sub["报价数量"].sum()),
            "周期状态": "完整" if as_of_date >= period_end else "进行中"
        })

    return pd.DataFrame(rows, columns=YEARLY_COLUMNS).sort_values(
        ["观察年序号", "始发地", "目的城市", "货物"], ignore_index=True
    ) if rows else pd.DataFrame(columns=YEARLY_COLUMNS)


def build_period_statistics(
    history_df: pd.DataFrame,
    observation_start: date,
    as_of_date: date | None = None
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    current_date = as_of_date or trusted_today()
    daily_df = prepare_observation_daily_data(history_df, observation_start, current_date)
    weekly_df = aggregate_weekly_statistics(daily_df, observation_start, current_date)
    monthly_df = aggregate_monthly_statistics(daily_df, observation_start, current_date)
    yearly_df = aggregate_yearly_statistics(monthly_df, observation_start, current_date)
    return weekly_df, monthly_df, yearly_df


def replace_generated_output_resilient(temp_file: str, filename: str) -> bool:
    """原子替换生成文件；目标被占用时保存待更新副本并保持采集运行。"""
    absolute_filename = os.path.abspath(filename)
    directory = os.path.dirname(absolute_filename)
    stem, extension = os.path.splitext(os.path.basename(absolute_filename))
    pending_file = os.path.join(directory, f"{stem}_待更新{extension}")
    try:
        os.replace(temp_file, absolute_filename)
        if os.path.exists(pending_file):
            try:
                os.remove(pending_file)
            except OSError:
                pass
        return True
    except PermissionError:
        os.replace(temp_file, pending_file)
        message = f"{filename} 正在被打开，最新数据已保存到 {pending_file}"
        print(f"警告：{message}")
        if LOGGER:
            LOGGER.warning(message)
        return False


def generated_output_temp_path(filename: str) -> str:
    absolute_filename = os.path.abspath(filename)
    directory = os.path.dirname(absolute_filename)
    stem, extension = os.path.splitext(os.path.basename(absolute_filename))
    os.makedirs(directory, exist_ok=True)
    return os.path.join(directory, f".{stem}_new{extension}")


def save_period_csv(period_df: pd.DataFrame, filename: str, columns: list[str]) -> bool:
    output_df = period_df.copy() if period_df is not None else pd.DataFrame(columns=columns)
    for column in columns:
        if column not in output_df.columns:
            output_df[column] = None
    output_df = output_df[columns]
    temp_file = generated_output_temp_path(filename)
    output_df.to_csv(
        temp_file,
        index=False,
        encoding="utf-8-sig",
        float_format=CSV_FLOAT_FORMAT
    )
    return replace_generated_output_resilient(temp_file, filename)


# =========================
# 18. 保存 CSV
# =========================

def format_csv_row(row: dict) -> dict:
    """让价格结果在 CSV 中固定显示三位小数。"""
    formatted = row.copy()
    for column in PRICE_OUTPUT_COLUMNS:
        value = formatted.get(column)
        if value is None or value == "":
            continue
        try:
            formatted[column] = f"{float(value):.{DECIMAL_PLACES}f}"
        except (TypeError, ValueError):
            continue
    return formatted


def save_csv(data, filename) -> bool:
    temp_file = generated_output_temp_path(filename)
    if not data:
        if filename in {DETAIL_CSV, PENDING_CSV}:
            headers = [
                "日期", "发布人", "始发地", "目的地", "目的城市", "货物小类", "货物大类",
                "报价文本", "最低报价", "最高报价", "平均报价", "是否整车价", "原始消息"
            ]
        else:
            headers = ["日期", "始发地", "目的城市", "货物", "平均运价", "最低运价", "最高运价", "报价数量"]

        with open(temp_file, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
        return replace_generated_output_resilient(temp_file, filename)

    fieldnames = list(data[0].keys())
    with open(temp_file, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(format_csv_row(row) for row in data)
    return replace_generated_output_resilient(temp_file, filename)


# =========================
# 18. 保存 Excel
# =========================

def auto_adjust_excel_width_and_date(filename):
    wb = load_workbook(filename)

    for ws in wb.worksheets:
        if ws["A1"].value == "日期":
            ws.column_dimensions["A"].width = 15

        header_columns = {
            cell.value: cell.column
            for cell in ws[1]
            if cell.value is not None
        }
        for header in PRICE_OUTPUT_COLUMNS:
            column_index = header_columns.get(header)
            if column_index is None:
                continue
            for row_index in range(2, ws.max_row + 1):
                ws.cell(row=row_index, column=column_index).number_format = EXCEL_DECIMAL_FORMAT

        for col in ws.columns:
            col_letter = col[0].column_letter
            max_length = 0
            for cell in col:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            current_width = ws.column_dimensions[col_letter].width or 0
            ws.column_dimensions[col_letter].width = max(current_width, max_length + 2, 12)

    wb.save(filename)


def save_excel(
    detail_data,
    history_daily_data,
    weekly_data,
    monthly_data,
    yearly_data,
    filename,
    pending_data=None
):
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        detail_df = pd.DataFrame(detail_data)
        history_df = pd.DataFrame(history_daily_data)
        weekly_df = pd.DataFrame(weekly_data)
        monthly_df = pd.DataFrame(monthly_data)
        yearly_df = pd.DataFrame(yearly_data)
        pending_df = pd.DataFrame(pending_data or [])

        if detail_df.empty:
            detail_df = pd.DataFrame(columns=[
                "日期", "发布人", "始发地", "目的地", "目的城市",
                "货物小类", "货物大类", "报价文本", "最低报价", "最高报价",
                "平均报价", "是否整车价", "原始消息"
            ])

        if history_df.empty:
            history_df = pd.DataFrame(columns=[
                "日期", "始发地", "目的城市", "货物", "平均运价", "最低运价", "最高运价", "报价数量", "7日均线", "价格指数"
            ])

        if pending_df.empty:
            pending_df = pd.DataFrame(columns=[
                "日期", "发布人", "始发地", "目的地", "目的城市",
                "货物小类", "货物大类", "报价文本", "最低报价", "最高报价",
                "平均报价", "是否整车价", "原始消息"
            ])

        detail_columns = [
            "日期", "发布人", "始发地", "目的地", "目的城市",
            "货物小类", "货物大类", "报价文本", "最低报价", "最高报价",
            "平均报价", "是否整车价", "原始消息"
        ]
        daily_columns = [
            "日期", "始发地", "目的城市", "货物", "平均运价", "最低运价", "最高运价", "报价数量", "7日均线", "价格指数"
        ]

        detail_df = detail_df[detail_columns]
        pending_df = pending_df[detail_columns]
        history_df = history_df[daily_columns]
        weekly_df = weekly_df.reindex(columns=WEEKLY_COLUMNS)
        monthly_df = monthly_df.reindex(columns=MONTHLY_COLUMNS)
        yearly_df = yearly_df.reindex(columns=YEARLY_COLUMNS)

        detail_df.to_excel(writer, index=False, sheet_name="报价明细")
        pending_df.to_excel(writer, index=False, sheet_name="待生效运价")
        history_df.to_excel(writer, index=False, sheet_name="历史每日平均运价")
        weekly_df.to_excel(writer, index=False, sheet_name="每周平均运价")
        monthly_df.to_excel(writer, index=False, sheet_name="每月平均运价")
        yearly_df.to_excel(writer, index=False, sheet_name="每年平均运价")

    auto_adjust_excel_width_and_date(filename)


def save_excel_resilient(
    detail_data,
    history_daily_data,
    weekly_data,
    monthly_data,
    yearly_data,
    filename,
    pending_data=None
) -> bool:
    """先生成临时Excel；原文件被打开时保存为待更新副本，不中断采集。"""
    absolute_filename = os.path.abspath(filename)
    directory = os.path.dirname(absolute_filename)
    base_name = os.path.splitext(os.path.basename(absolute_filename))[0]
    temp_file = os.path.join(directory, f".{base_name}_new.xlsx")
    pending_file = os.path.join(directory, PENDING_REPORT_XLSX)
    save_excel(
        detail_data,
        history_daily_data,
        weekly_data,
        monthly_data,
        yearly_data,
        temp_file,
        pending_data,
    )
    try:
        os.replace(temp_file, absolute_filename)
        if os.path.exists(pending_file):
            try:
                os.remove(pending_file)
            except OSError:
                pass
        return True
    except PermissionError:
        os.replace(temp_file, pending_file)
        print(
            f"警告：{filename} 正在被打开，已保存最新副本: {PENDING_REPORT_XLSX}；"
            "关闭原Excel后，下次更新会自动替换。"
        )
        if LOGGER:
            LOGGER.warning("Excel被占用，最新副本已保存到 %s", pending_file)
        return False


def observation_year_window(
    observation_start: date,
    as_of_date: date,
) -> tuple[int, date, date]:
    """返回当前观察年序号及边界，与现有连续12个观察月口径一致。"""
    start_month = date(observation_start.year, observation_start.month, 1)
    month_offset = max(
        0,
        (as_of_date.year - start_month.year) * 12
        + as_of_date.month - start_month.month,
    )
    year_number = month_offset // 12 + 1
    year_start_month = add_months(start_month, (year_number - 1) * 12)
    period_start = observation_start if year_number == 1 else year_start_month
    period_end = get_month_end(add_months(year_start_month, 11))
    return year_number, period_start, period_end


def current_observation_year_daily_data(
    history_df: pd.DataFrame,
    observation_start: date,
    as_of_date: date,
) -> tuple[pd.DataFrame, int, date, date]:
    """筛出当前观察年的每日小类数据，年度切换后旧年度不再写入Excel。"""
    year_number, period_start, period_end = observation_year_window(
        observation_start,
        as_of_date,
    )
    if history_df is None or history_df.empty:
        return pd.DataFrame(), year_number, period_start, period_end

    result = history_df.copy()
    result["日期排序"] = pd.to_datetime(result["日期"], errors="coerce")
    result = result.dropna(subset=["日期排序"])
    actual_end = min(as_of_date, period_end)
    result = result[
        (result["日期排序"] >= pd.Timestamp(period_start))
        & (result["日期排序"] <= pd.Timestamp(actual_end))
    ].copy()
    result = result.sort_values(
        ["日期排序", "始发地", "目的城市", "货物"],
        ignore_index=True,
    ).drop(columns=["日期排序"])
    return result, year_number, period_start, period_end


def spring_festival_date(year: int) -> date:
    """返回公历年份对应的春节（农历正月初一）日期。"""
    try:
        return LunarDate(int(year), 1, 1).toSolarDate()
    except (TypeError, ValueError, RuntimeError) as exc:
        raise ValueError(
            f"无法计算 {year} 年春节；当前农历库支持1900至2099年。"
        ) from exc


def spring_festival_year_window(spring_year: int) -> tuple[date, date]:
    """返回一个春节年度的起止日期，结束日为下一年春节前一天。"""
    period_start = spring_festival_date(spring_year)
    period_end = spring_festival_date(spring_year + 1) - timedelta(days=1)
    return period_start, period_end


def spring_festival_year_for_date(value: date) -> int:
    """按北京时间日期判断其所属的春节年度。"""
    this_year_festival = spring_festival_date(value.year)
    return value.year if value >= this_year_festival else value.year - 1


def spring_festival_year_daily_data(
    history_df: pd.DataFrame,
    spring_year: int,
    as_of_date: date,
) -> tuple[pd.DataFrame, date, date]:
    """筛出指定春节年度内截至给定日期的每日小类数据。"""
    period_start, period_end = spring_festival_year_window(spring_year)
    if history_df is None or history_df.empty:
        return pd.DataFrame(), period_start, period_end

    result = history_df.copy()
    result["日期排序"] = pd.to_datetime(result["日期"], errors="coerce")
    result = result.dropna(subset=["日期排序"])
    actual_end = min(as_of_date, period_end)
    result = result[
        (result["日期排序"] >= pd.Timestamp(period_start))
        & (result["日期排序"] <= pd.Timestamp(actual_end))
    ].copy()
    result = result.sort_values(
        ["日期排序", "始发地", "目的城市", "货物"],
        ignore_index=True,
    ).drop(columns=["日期排序"])
    # 年度切换后，7日均线和价格指数也必须从新春节年度重新起算。
    result = enrich_history_indicators(result)
    return result, period_start, period_end


def aggregate_spring_festival_year_statistics(
    monthly_df: pd.DataFrame,
    spring_year: int,
    period_start: date,
    period_end: date,
    as_of_date: date,
) -> pd.DataFrame:
    """按春节年度汇总月度数据，年度边界不沿用原大类的连续12个月口径。"""
    if monthly_df is None or monthly_df.empty:
        return pd.DataFrame(columns=YEARLY_COLUMNS)

    actual_end = min(as_of_date, period_end)
    elapsed_months = (
        (actual_end.year - period_start.year) * 12
        + actual_end.month - period_start.month
        + 1
    )
    rows = []
    for (origin, city, cargo), sub in monthly_df.groupby(
        ["始发地", "目的城市", "货物"]
    ):
        data_months = int(sub["月份"].nunique())
        rows.append({
            "观察年序号": int(spring_year),
            "观察年": f"{spring_year}春节年度",
            "周期开始": period_start.isoformat(),
            "周期结束": period_end.isoformat(),
            "实际观察结束": actual_end.isoformat(),
            "始发地": origin,
            "目的城市": city,
            "货物": cargo,
            "平均运价": round(float(sub["平均运价"].mean()), DECIMAL_PLACES),
            "最低运价": round(float(sub["最低运价"].min()), DECIMAL_PLACES),
            "最高运价": round(float(sub["最高运价"].max()), DECIMAL_PLACES),
            "数据月份数": data_months,
            "已过月份数": elapsed_months,
            "月份覆盖率": round(
                data_months / elapsed_months * 100,
                DECIMAL_PLACES,
            ),
            "数据天数": int(sub["数据天数"].sum()),
            "报价数量": int(sub["报价数量"].sum()),
            "周期状态": "完整" if as_of_date >= period_end else "进行中",
        })

    return pd.DataFrame(rows, columns=YEARLY_COLUMNS).sort_values(
        ["观察年序号", "始发地", "目的城市", "货物"],
        ignore_index=True,
    )


def build_spring_festival_period_statistics(
    history_df: pd.DataFrame,
    spring_year: int,
    as_of_date: date,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """生成一个春节年度内的小类周、月、年统计。"""
    period_start, period_end = spring_festival_year_window(spring_year)
    actual_end = min(as_of_date, period_end)
    daily_df = prepare_observation_daily_data(
        history_df,
        period_start,
        actual_end,
    )
    weekly_df = aggregate_weekly_statistics(
        daily_df,
        period_start,
        actual_end,
    )
    monthly_df = aggregate_monthly_statistics(
        daily_df,
        period_start,
        actual_end,
    )
    yearly_df = aggregate_spring_festival_year_statistics(
        monthly_df,
        spring_year,
        period_start,
        period_end,
        actual_end,
    )
    return weekly_df, monthly_df, yearly_df


def safe_statistics_component(value: str, fallback: str) -> str:
    """生成单级Windows目录/文件名，防止配置名称逃逸统计根目录。"""
    safe_value = re.sub(r'[\\/:*?"<>|]', "_", str(value or ""))
    safe_value = safe_value.strip().rstrip(".")
    if safe_value in {"", ".", ".."}:
        return fallback
    return safe_value


def resolve_small_category_output_dir(category: str, subcategory: str) -> str:
    """建立“小类统计/大类/小类”目录，并校验最终路径仍在统计根目录内。"""
    root = os.path.abspath(SMALL_CATEGORY_OUTPUT_FOLDER)
    category_name = safe_statistics_component(category, DEFAULT_CARGO_CATEGORY)
    subcategory_name = safe_statistics_component(
        subcategory,
        DEFAULT_CARGO_SUBCATEGORY,
    )
    output_dir = os.path.abspath(os.path.join(root, category_name, subcategory_name))
    if os.path.commonpath([root, output_dir]) != root:
        raise ValueError("小类统计输出目录超出允许范围。")
    os.makedirs(output_dir, exist_ok=True)
    return output_dir


def save_small_category_daily_excel(
    daily_df: pd.DataFrame,
    subcategory: str,
    category: str,
    spring_year: int,
    period_start: date,
    period_end: date,
    filename: str,
) -> None:
    """保存一个小类在指定春节年度内的每日统计。"""
    columns = [
        "春节年度", "年度开始", "年度结束", "日期", "始发地", "目的城市",
        "货物大类", "货物小类", "平均运价", "最低运价", "最高运价",
        "报价数量", "7日均线", "价格指数",
    ]
    output_df = daily_df.copy() if daily_df is not None else pd.DataFrame()
    if output_df.empty:
        output_df = pd.DataFrame(columns=columns)
    else:
        output_df["春节年度"] = f"{spring_year}春节年度"
        output_df["年度开始"] = period_start.isoformat()
        output_df["年度结束"] = period_end.isoformat()
        output_df["货物大类"] = category
        output_df["货物小类"] = subcategory
        output_df = output_df.reindex(columns=columns)

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        output_df.to_excel(writer, index=False, sheet_name="每日统计")
    auto_adjust_excel_width_and_date(filename)


def save_small_category_daily_excel_resilient(
    daily_df: pd.DataFrame,
    subcategory: str,
    category: str,
    spring_year: int,
    period_start: date,
    period_end: date,
    filename: str,
) -> bool:
    """原子更新小类每日Excel；文件被占用时保存待更新副本。"""
    absolute_filename = os.path.abspath(filename)
    output_dir = os.path.dirname(absolute_filename)
    os.makedirs(output_dir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(absolute_filename))[0]
    temp_file = os.path.join(output_dir, f".{stem}_new.xlsx")
    pending_file = os.path.join(output_dir, f"{stem}_待更新.xlsx")
    save_small_category_daily_excel(
        daily_df,
        subcategory,
        category,
        spring_year,
        period_start,
        period_end,
        temp_file,
    )
    try:
        os.replace(temp_file, absolute_filename)
        if os.path.exists(pending_file):
            try:
                os.remove(pending_file)
            except OSError:
                pass
        return True
    except PermissionError:
        os.replace(temp_file, pending_file)
        if LOGGER:
            LOGGER.warning("小类每日Excel被占用，最新副本已保存到 %s", pending_file)
        return False


# =========================
# 19. 按始发地与目的城市画历史折线图
# =========================

def draw_destination_daily_charts(history_df):
    if history_df is None or history_df.empty:
        print("历史日均价为空，无法画图。")
        return

    plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False

    history_df = history_df.copy()
    history_df["日期排序"] = pd.to_datetime(history_df["日期"], format="%Y-%m-%d", errors="coerce")

    if "始发地" not in history_df.columns:
        history_df["始发地"] = "历史未区分"
    if "货物" not in history_df.columns:
        history_df["货物"] = DEFAULT_CARGO_CATEGORY

    for (origin, city, cargo), sub in history_df.groupby(
        ["始发地", "目的城市", "货物"]
    ):
        sub = sub.copy()
        sub = sub.sort_values("日期排序")
        route_name = f"{origin}到{city}"
        safe_route_name = re.sub(r'[\\/:*?"<>|]', "_", route_name)
        safe_cargo = re.sub(r'[\\/:*?"<>|]', "_", str(cargo))

        plt.figure(figsize=(10, 5))
        plt.plot(sub["日期"], sub["平均运价"], marker="o", label="日均运价")
        if "7日均线" in sub.columns:
            plt.plot(sub["日期"], sub["7日均线"], marker="s", label="7日均线")
        plt.title(f"{route_name}-{cargo}日均运价走势图")
        plt.xlabel("日期")
        plt.ylabel("平均运价")
        plt.legend()
        plt.xticks(rotation=45)
        plt.grid(True)
        plt.tight_layout()
        out_path = os.path.join(
            CHART_FOLDER,
            f"{safe_route_name}_{safe_cargo}_日均运价走势图.png",
        )
        plt.savefig(out_path, dpi=150)
        plt.close()

        if "价格指数" in sub.columns:
            plt.figure(figsize=(10, 5))
            plt.plot(sub["日期"], sub["价格指数"], marker="o", label="价格指数")
            plt.title(f"{route_name}-{cargo}价格指数走势图")
            plt.xlabel("日期")
            plt.ylabel("指数（首日=100）")
            plt.legend()
            plt.xticks(rotation=45)
            plt.grid(True)
            plt.tight_layout()
            out_path = os.path.join(
                CHART_FOLDER,
                f"{safe_route_name}_{safe_cargo}_价格指数走势图.png",
            )
            plt.savefig(out_path, dpi=150)
            plt.close()

        print(f"已更新 {route_name} 的运价图和指数图")


def draw_route_period_chart(
    period_df: pd.DataFrame,
    x_column: str,
    label_column: str,
    period_name: str,
    filename_suffix: str,
    output_dir: str = CHART_FOLDER,
) -> None:
    if period_df is None or period_df.empty:
        print(f"{period_name}统计为空，暂不生成{period_name}走势图。")
        return

    for (origin, city, cargo), sub in period_df.groupby(
        ["始发地", "目的城市", "货物"]
    ):
        sub = sub.sort_values(x_column).copy()
        route_name = f"{origin}到{city}"
        safe_route_name = re.sub(r'[\\/:*?"<>|]', "_", route_name)
        safe_cargo = re.sub(r'[\\/:*?"<>|]', "_", str(cargo))
        if period_name == "周":
            labels = sub[label_column].map(lambda value: f"第{int(value)}周")
        elif period_name == "年":
            if label_column == "观察年":
                labels = sub[label_column].astype(str)
            else:
                labels = sub[label_column].map(lambda value: f"第{int(value)}观察年")
        else:
            labels = sub[label_column].astype(str)
        if "周期状态" in sub.columns:
            labels = labels + sub["周期状态"].map(
                lambda status: "（进行中）" if status == "进行中" else ""
            )

        plt.figure(figsize=(11, 5.5))
        plt.plot(labels, sub["平均运价"], marker="o", label=f"{period_name}平均运价")
        plt.title(f"{route_name}-{cargo}{period_name}价格走势图")
        plt.xlabel(period_name)
        plt.ylabel("平均运价")
        plt.legend()
        plt.xticks(rotation=45, ha="right")
        plt.grid(True)
        plt.tight_layout()
        os.makedirs(output_dir, exist_ok=True)
        out_path = os.path.join(
            output_dir,
            f"{safe_route_name}_{safe_cargo}_{filename_suffix}.png"
        )
        plt.savefig(out_path, dpi=150)
        plt.close()
        print(f"已更新 {route_name} 的{period_name}价格走势图")


def draw_observation_year_monthly_charts(monthly_df: pd.DataFrame) -> None:
    """每个观察年以月份为横轴，并用虚线标出该观察年的月均价格。"""
    if monthly_df is None or monthly_df.empty:
        return

    df = monthly_df.copy()
    df["观察年序号"] = ((df["月份序号"] - 1) // 12 + 1).astype(int)
    for (origin, city, cargo, year_number), sub in df.groupby(
        ["始发地", "目的城市", "货物", "观察年序号"]
    ):
        sub = sub.sort_values("月份序号")
        route_name = f"{origin}到{city}"
        safe_route_name = re.sub(r'[\\/:*?"<>|]', "_", route_name)
        safe_cargo = re.sub(r'[\\/:*?"<>|]', "_", str(cargo))
        year_average = round(float(sub["平均运价"].mean()), DECIMAL_PLACES)

        plt.figure(figsize=(11, 5.5))
        plt.plot(sub["月份"], sub["平均运价"], marker="o", label="月平均运价")
        plt.axhline(
            year_average,
            color="orange",
            linestyle="--",
            label=f"观察年平均 {year_average:.{DECIMAL_PLACES}f}"
        )
        plt.title(f"{route_name}-{cargo}-第{int(year_number)}观察年月度价格走势")
        plt.xlabel("月份")
        plt.ylabel("平均运价")
        plt.legend()
        plt.xticks(rotation=45, ha="right")
        plt.grid(True)
        plt.tight_layout()
        out_path = os.path.join(
            CHART_FOLDER,
            f"{safe_route_name}_{safe_cargo}_第{int(year_number)}观察年_月度价格走势图.png"
        )
        plt.savefig(out_path, dpi=150)
        plt.close()


def draw_period_price_charts(
    weekly_df: pd.DataFrame,
    monthly_df: pd.DataFrame,
    yearly_df: pd.DataFrame
) -> None:
    plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False

    draw_route_period_chart(weekly_df, "周序号", "周序号", "周", "周价格走势图")
    draw_route_period_chart(monthly_df, "月份序号", "月份", "月", "月价格走势图")
    draw_route_period_chart(
        yearly_df,
        "观察年序号",
        "观察年序号",
        "年",
        "年价格走势图"
    )
    draw_observation_year_monthly_charts(monthly_df)


def draw_spring_festival_period_price_charts(
    weekly_df: pd.DataFrame,
    monthly_df: pd.DataFrame,
    yearly_df: pd.DataFrame,
) -> None:
    """为春节年度的大类结果生成周、月、年图，不改变旧观察年函数。"""
    plt.rcParams["font.sans-serif"] = [
        "SimHei", "Microsoft YaHei", "Arial Unicode MS", "DejaVu Sans"
    ]
    plt.rcParams["axes.unicode_minus"] = False
    draw_route_period_chart(
        weekly_df, "周序号", "周序号", "周", "周价格走势图"
    )
    draw_route_period_chart(
        monthly_df, "月份序号", "月份", "月", "月价格走势图"
    )
    draw_route_period_chart(
        yearly_df, "观察年序号", "观察年", "年", "年价格走势图"
    )


def small_category_archive_fingerprint(
    daily_df: pd.DataFrame,
    subcategory: str,
    category: str,
    spring_year: int,
) -> str:
    """生成稳定摘要；旧年度数据未变化时不重复生成归档。"""
    canonical = daily_df.copy() if daily_df is not None else pd.DataFrame()
    if not canonical.empty:
        canonical = canonical.reindex(sorted(canonical.columns), axis=1)
        sort_columns = [
            column for column in ("日期", "始发地", "目的城市", "货物")
            if column in canonical.columns
        ]
        if sort_columns:
            canonical = canonical.sort_values(sort_columns, ignore_index=True)
    payload = "|".join([
        str(SMALL_CATEGORY_ARCHIVE_SCHEMA_VERSION),
        str(spring_year),
        category,
        subcategory,
        canonical.to_csv(index=False, float_format=CSV_FLOAT_FORMAT),
    ])
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def read_json_object(path: str) -> dict:
    try:
        with open(path, "r", encoding="utf-8") as stream:
            value = json.load(stream)
        return value if isinstance(value, dict) else {}
    except (OSError, TypeError, ValueError, json.JSONDecodeError):
        return {}


def write_json_object_atomic(path: str, value: dict) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    temp_path = path + ".tmp"
    with open(temp_path, "w", encoding="utf-8") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)
    os.replace(temp_path, path)


def filter_cargo_period(dataframe: pd.DataFrame, subcategory: str) -> pd.DataFrame:
    if dataframe is None or dataframe.empty or "货物" not in dataframe.columns:
        return pd.DataFrame()
    return dataframe[
        dataframe["货物"].astype(str) == subcategory
    ].copy()


def draw_small_category_period_charts(
    subcategory: str,
    weekly_df: pd.DataFrame,
    monthly_df: pd.DataFrame,
    yearly_df: pd.DataFrame,
    output_dir: str,
    suffix_label: str = "",
) -> int:
    before_files = {
        name for name in os.listdir(output_dir) if name.lower().endswith(".png")
    }
    chart_specs = (
        (weekly_df, "周序号", "周序号", "周", f"{suffix_label}周价格走势图"),
        (monthly_df, "月份序号", "月份", "月", f"{suffix_label}月价格走势图"),
        (yearly_df, "观察年序号", "观察年", "年", f"{suffix_label}年价格走势图"),
    )
    for dataframe, x_column, label_column, period_name, suffix in chart_specs:
        cargo_period = filter_cargo_period(dataframe, subcategory)
        if cargo_period.empty:
            continue
        draw_route_period_chart(
            cargo_period,
            x_column,
            label_column,
            period_name,
            suffix,
            output_dir,
        )
    after_files = {
        name for name in os.listdir(output_dir) if name.lower().endswith(".png")
    }
    return len(after_files - before_files)


def write_small_category_outputs(
    history_df: pd.DataFrame,
    as_of_date: date,
) -> dict:
    """先补齐旧春节年度归档，再生成当前春节年度的小类业务结果。"""
    cargo_names = set()
    cargo_names.update(
        str(name).strip()
        for name in CARGO_TYPES
        if str(name).strip()
    )
    if history_df is not None and not history_df.empty and "货物" in history_df.columns:
        cargo_names.update(
            str(value).strip()
            for value in history_df["货物"].dropna().tolist()
            if str(value).strip()
        )

    current_spring_year = spring_festival_year_for_date(as_of_date)
    history_years = set()
    # Existing generated archives must also be regenerated when their last quote is deleted.
    for subcategory in cargo_names:
        archive_root = os.path.join(resolve_small_category_output_dir(cargo_category(subcategory), subcategory),
                                    SMALL_CATEGORY_ARCHIVE_FOLDER)
        if os.path.isdir(archive_root):
            for name in os.listdir(archive_root):
                match = re.fullmatch(r"(\d{4})春节年度", name)
                if match and os.path.isfile(os.path.join(archive_root, name, SMALL_CATEGORY_ARCHIVE_STATE)):
                    history_years.add(int(match.group(1)))
    if history_df is not None and not history_df.empty:
        for value in pd.to_datetime(history_df["日期"], errors="coerce").dropna():
            history_years.add(spring_festival_year_for_date(value.date()))
    completed_years = sorted(
        year for year in history_years if year < current_spring_year
    )
    statistics_by_year = {}
    for spring_year in [*completed_years, current_spring_year]:
        statistics_by_year[spring_year] = build_spring_festival_period_statistics(
            history_df,
            spring_year,
            as_of_date,
        )

    yearly_frames = [
        statistics[2]
        for statistics in statistics_by_year.values()
        if statistics[2] is not None and not statistics[2].empty
    ]
    all_yearly_df = (
        pd.concat(yearly_frames, ignore_index=True)
        if yearly_frames else pd.DataFrame(columns=YEARLY_COLUMNS)
    )

    output_summary = {
        "subcategory_count": 0,
        "daily_excel_count": 0,
        "daily_excel_pending_count": 0,
        "chart_count": 0,
        "archived_year_count": 0,
        "archive_skipped_count": 0,
        "archive_pending_count": 0,
        "current_spring_year": current_spring_year,
    }

    # 归档阶段必须先于当前年度输出，确保春节期间停机后首次启动先补旧账。
    for spring_year in completed_years:
        weekly_df, monthly_df, yearly_df = statistics_by_year[spring_year]
        period_start, period_end = spring_festival_year_window(spring_year)
        year_label = f"{spring_year}春节年度"
        for subcategory in sorted(cargo_names):
            cargo_history = filter_cargo_period(history_df, subcategory)
            year_daily, _start, _end = spring_festival_year_daily_data(
                cargo_history,
                spring_year,
                period_end,
            )
            category = cargo_category(subcategory)
            base_dir = resolve_small_category_output_dir(category, subcategory)
            archive_dir = os.path.join(
                base_dir,
                SMALL_CATEGORY_ARCHIVE_FOLDER,
                year_label,
            )
            if year_daily.empty and not os.path.isdir(archive_dir):
                continue
            os.makedirs(archive_dir, exist_ok=True)
            safe_subcategory = safe_statistics_component(
                subcategory,
                DEFAULT_CARGO_SUBCATEGORY,
            )
            excel_name = f"{safe_subcategory}_{year_label}_每日统计.xlsx"
            excel_path = os.path.join(archive_dir, excel_name)
            state_path = os.path.join(
                archive_dir,
                SMALL_CATEGORY_ARCHIVE_STATE,
            )
            fingerprint = small_category_archive_fingerprint(
                year_daily,
                subcategory,
                category,
                spring_year,
            )
            state = read_json_object(state_path)
            expected_files = state.get("files", [])
            archive_is_current = (
                state.get("schema_version") == SMALL_CATEGORY_ARCHIVE_SCHEMA_VERSION
                and state.get("fingerprint") == fingerprint
                and isinstance(expected_files, list)
                and expected_files
                and all(
                    os.path.isfile(os.path.join(archive_dir, str(name)))
                    for name in expected_files
                )
            )
            if archive_is_current:
                output_summary["archive_skipped_count"] += 1
                continue

            for name in os.listdir(archive_dir):
                path = os.path.join(archive_dir, name)
                if os.path.isfile(path) and name.lower().endswith(".png"):
                    try:
                        os.remove(path)
                    except OSError:
                        if LOGGER:
                            LOGGER.warning("无法清理旧年度小类图表: %s", path)
            excel_updated = save_small_category_daily_excel_resilient(
                year_daily,
                subcategory,
                category,
                spring_year,
                period_start,
                period_end,
                excel_path,
            )
            chart_count = draw_small_category_period_charts(
                subcategory,
                weekly_df,
                monthly_df,
                yearly_df,
                archive_dir,
                f"{year_label}_",
            )
            output_summary["chart_count"] += chart_count
            if not excel_updated:
                output_summary["archive_pending_count"] += 1
                continue
            generated_files = sorted(
                name for name in os.listdir(archive_dir)
                if os.path.isfile(os.path.join(archive_dir, name))
                and (name.lower().endswith(".png") or name == excel_name)
            )
            write_json_object_atomic(state_path, {
                "schema_version": SMALL_CATEGORY_ARCHIVE_SCHEMA_VERSION,
                "spring_year": spring_year,
                "period_start": period_start.isoformat(),
                "period_end": period_end.isoformat(),
                "fingerprint": fingerprint,
                "files": generated_files,
            })
            output_summary["archived_year_count"] += 1

    current_weekly, current_monthly, _current_yearly = statistics_by_year[
        current_spring_year
    ]
    for subcategory in sorted(cargo_names):
        category = cargo_category(subcategory)
        output_dir = resolve_small_category_output_dir(category, subcategory)
        for name in os.listdir(output_dir):
            path = os.path.join(output_dir, name)
            if os.path.isfile(path) and name.lower().endswith(".png"):
                try:
                    os.remove(path)
                except OSError:
                    if LOGGER:
                        LOGGER.warning("无法清理旧小类图表: %s", path)

        cargo_history = (
            history_df[history_df["货物"].astype(str) == subcategory].copy()
            if history_df is not None and not history_df.empty
            else pd.DataFrame()
        )
        current_daily, period_start, period_end = (
            spring_festival_year_daily_data(
                cargo_history,
                current_spring_year,
                as_of_date,
            )
        )
        safe_subcategory = safe_statistics_component(
            subcategory,
            DEFAULT_CARGO_SUBCATEGORY,
        )
        excel_path = os.path.join(
            output_dir,
            f"{safe_subcategory}_每日统计.xlsx",
        )
        excel_updated = save_small_category_daily_excel_resilient(
            current_daily,
            subcategory,
            category,
            current_spring_year,
            period_start,
            period_end,
            excel_path,
        )
        output_summary["daily_excel_count"] += int(excel_updated)
        output_summary["daily_excel_pending_count"] += int(not excel_updated)

        output_summary["chart_count"] += draw_small_category_period_charts(
            subcategory,
            current_weekly,
            current_monthly,
            all_yearly_df,
            output_dir,
        )

        output_summary["subcategory_count"] += 1

    return output_summary


def record_spring_festival_year(record: dict) -> int | None:
    try:
        record_date = datetime.strptime(str(record["日期"]), "%Y-%m-%d").date()
    except (KeyError, TypeError, ValueError):
        return None
    return spring_festival_year_for_date(record_date)


def large_category_archive_fingerprint(
    records: list[dict],
    spring_year: int,
) -> str:
    canonical_records = []
    for record in records:
        canonical_records.append({
            str(key): record.get(key)
            for key in sorted(record)
        })
    canonical_records.sort(key=lambda item: (
        str(item.get("日期", "")),
        str(item.get("始发地", "")),
        str(item.get("目的城市", "")),
        str(item.get("货物小类", "")),
        str(item.get("原始消息", "")),
    ))
    payload = json.dumps(
        {
            "schema_version": LARGE_CATEGORY_ARCHIVE_SCHEMA_VERSION,
            "spring_year": spring_year,
            "records": canonical_records,
        },
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def generated_business_files(output_dir: str) -> list[str]:
    result = []
    for root, _directories, files in os.walk(output_dir):
        for name in files:
            if name in {LARGE_CATEGORY_ARCHIVE_STATE, PENDING_REPORT_XLSX}:
                continue
            relative = os.path.relpath(os.path.join(root, name), output_dir)
            if (
                name.lower().endswith((".csv", ".xlsx", ".png"))
                and not name.startswith(".")
            ):
                result.append(relative)
    return sorted(result)


def write_large_category_period_outputs(
    records: list[dict],
    pending_records: list[dict],
    spring_year: int,
    as_of_date: date,
    output_dir: str,
) -> dict:
    """在指定目录生成一个春节年度的大类业务文件。"""
    period_start, period_end = spring_festival_year_window(spring_year)
    actual_end = min(as_of_date, period_end)
    daily_summary = summarize_daily_quote(records)
    history_df = merge_daily_history(daily_summary)
    weekly_df, monthly_df, yearly_df = build_spring_festival_period_statistics(
        history_df,
        spring_year,
        actual_end,
    )

    os.makedirs(output_dir, exist_ok=True)
    previous_directory = os.getcwd()
    try:
        os.chdir(output_dir)
        os.makedirs(CHART_FOLDER, exist_ok=True)
        save_csv(records, DETAIL_CSV)
        save_csv(pending_records, PENDING_CSV)
        save_csv(daily_summary, DAILY_CSV)
        save_period_csv(weekly_df, WEEKLY_CSV, WEEKLY_COLUMNS)
        save_period_csv(monthly_df, MONTHLY_CSV, MONTHLY_COLUMNS)
        save_period_csv(yearly_df, YEARLY_CSV, YEARLY_COLUMNS)
        excel_updated = save_excel_resilient(
            records,
            history_df.to_dict(orient="records"),
            weekly_df.to_dict(orient="records"),
            monthly_df.to_dict(orient="records"),
            yearly_df.to_dict(orient="records"),
            REPORT_XLSX,
            pending_records,
        )
        draw_spring_festival_period_price_charts(
            weekly_df,
            monthly_df,
            yearly_df,
        )
    finally:
        os.chdir(previous_directory)

    return {
        "daily_summary": daily_summary,
        "history_df": history_df,
        "weekly_df": weekly_df,
        "monthly_df": monthly_df,
        "yearly_df": yearly_df,
        "excel_updated": excel_updated,
        "period_start": period_start,
        "period_end": period_end,
    }


def ensure_large_category_archives(
    records: list[dict],
    current_spring_year: int,
) -> dict:
    """幂等补齐所有已结束春节年度的大类归档。"""
    records_by_year = defaultdict(list)
    for record in records:
        spring_year = record_spring_festival_year(record)
        if spring_year is not None and spring_year < current_spring_year:
            records_by_year[spring_year].append(record)

    if os.path.isdir(BUSINESS_ARCHIVE_FOLDER):
        for name in os.listdir(BUSINESS_ARCHIVE_FOLDER):
            match = re.fullmatch(r"(\d{4})春节年度", name)
            if match and int(match.group(1)) < current_spring_year and os.path.isfile(os.path.join(
                BUSINESS_ARCHIVE_FOLDER, name, LARGE_CATEGORY_ARCHIVE_FOLDER, LARGE_CATEGORY_ARCHIVE_STATE
            )):
                records_by_year.setdefault(int(match.group(1)), [])

    summary = {
        "archived_years": [],
        "skipped_years": [],
        "pending_years": [],
    }
    for spring_year in sorted(records_by_year):
        year_label = f"{spring_year}春节年度"
        archive_dir = os.path.abspath(os.path.join(
            BUSINESS_ARCHIVE_FOLDER,
            year_label,
            LARGE_CATEGORY_ARCHIVE_FOLDER,
        ))
        os.makedirs(archive_dir, exist_ok=True)
        state_path = os.path.join(archive_dir, LARGE_CATEGORY_ARCHIVE_STATE)
        fingerprint = large_category_archive_fingerprint(
            records_by_year[spring_year],
            spring_year,
        )
        state = read_json_object(state_path)
        expected_files = state.get("files", [])
        if (
            state.get("schema_version") == LARGE_CATEGORY_ARCHIVE_SCHEMA_VERSION
            and state.get("fingerprint") == fingerprint
            and isinstance(expected_files, list)
            and expected_files
            and all(
                os.path.isfile(os.path.join(archive_dir, str(name)))
                for name in expected_files
            )
        ):
            summary["skipped_years"].append(spring_year)
            continue

        chart_dir = os.path.join(archive_dir, CHART_FOLDER)
        if os.path.isdir(chart_dir):
            for name in os.listdir(chart_dir):
                path = os.path.join(chart_dir, name)
                if os.path.isfile(path) and name.lower().endswith(".png"):
                    try:
                        os.remove(path)
                    except OSError:
                        if LOGGER:
                            LOGGER.warning("无法清理旧年度大类图表: %s", path)
        period_start, period_end = spring_festival_year_window(spring_year)
        result = write_large_category_period_outputs(
            records_by_year[spring_year],
            [],
            spring_year,
            period_end,
            archive_dir,
        )
        if not result["excel_updated"]:
            summary["pending_years"].append(spring_year)
            continue
        files = generated_business_files(archive_dir)
        write_json_object_atomic(state_path, {
            "schema_version": LARGE_CATEGORY_ARCHIVE_SCHEMA_VERSION,
            "spring_year": spring_year,
            "period_start": period_start.isoformat(),
            "period_end": period_end.isoformat(),
            "fingerprint": fingerprint,
            "files": files,
        })
        summary["archived_years"].append(spring_year)
    return summary


# =========================
# 20. 一次完整重建
# =========================

def rebuild_all(
    default_origin: str = DEFAULT_ORIGIN,
    as_of_date: date | None = None,
    database_file: str | None = None,
    history_retention_days: int = 0,
) -> dict:
    print("\n检测到文件变化，开始重新统计...")
    effective_as_of_date = as_of_date or trusted_today()

    if database_file:
        database = FreightDatabase(database_file)
        all_records = database.fetch_records()
        database.export_rejections_csv(REJECTED_CSV)
    else:
        formatted_file = auto_format_input_file(
            INPUT_FILE,
            FORMATTED_INPUT_FILE,
            default_origin
        )
        all_records = deduplicate_records(parse_chat_file(formatted_file, default_origin))
    active_history_records = []
    pending_records = []
    for raw_record in all_records:
        record = normalize_record_cargo(raw_record)
        allowed, _reason = validate_cargo_route(record)
        if not allowed:
            continue
        try:
            record_date = datetime.strptime(record["日期"], "%Y-%m-%d").date()
        except (KeyError, TypeError, ValueError):
            continue
        if record_date <= effective_as_of_date:
            active_history_records.append(record)
        else:
            pending_records.append(record)

    current_spring_year = spring_festival_year_for_date(effective_as_of_date)

    # 先归档全部已结束春节年度，再触碰当前年度文件。
    large_archive_summary = ensure_large_category_archives(
        active_history_records,
        current_spring_year,
    )
    current_records = [
        record for record in active_history_records
        if record_spring_festival_year(record) == current_spring_year
    ]

    delete_generated_outputs()
    current_large_output = write_large_category_period_outputs(
        current_records,
        pending_records,
        current_spring_year,
        effective_as_of_date,
        os.getcwd(),
    )
    excel_updated = current_large_output["excel_updated"]

    small_daily_summary = summarize_daily_quote_by_subcategory(
        active_history_records
    )
    history_minimum_date = None
    if int(history_retention_days) > 0:
        history_minimum_date = effective_as_of_date - timedelta(
            days=int(history_retention_days)
        )
    small_history_df = merge_daily_history(
        small_daily_summary,
        history_minimum_date,
    )
    small_output_summary = write_small_category_outputs(
        small_history_df,
        effective_as_of_date,
    )

    print(f"当前春节年度明细记录: {len(current_records)}")
    print(f"历史有效记录总数: {len(active_history_records)}")
    print(f"待生效缓存记录: {len(pending_records)}")
    print("已生成文件：")
    print(f"1. {DETAIL_CSV}")
    print(f"2. {PENDING_CSV}")
    print(f"3. {DAILY_CSV}")
    print(f"4. {WEEKLY_CSV}")
    print(f"5. {MONTHLY_CSV}")
    print(f"6. {YEARLY_CSV}")
    print(f"7. {REPORT_XLSX}")
    print(f"8. {CHART_FOLDER} 文件夹中的周、月、年走势图")
    if database_file:
        print(f"9. {REJECTED_CSV}")
        print(f"10. {DATABASE_FILE}")
    print(
        f"11. {SMALL_CATEGORY_OUTPUT_FOLDER} 文件夹中的小类春节年度每日Excel及周、月、年走势图"
    )
    print(
        f"12. {BUSINESS_ARCHIVE_FOLDER} 文件夹中的旧春节年度大类业务归档"
    )

    rejected_count = 0
    if database_file:
        rejected_count = len(FreightDatabase(database_file).fetch_rejections())
    return {
        "active_records": len(current_records),
        "historical_active_records": len(active_history_records),
        "pending_records": len(pending_records),
        "rejected_records": rejected_count,
        "excel_updated": excel_updated,
        "current_spring_year": current_spring_year,
        "large_category_archives": large_archive_summary,
        "small_category_outputs": small_output_summary,
        "as_of_date": effective_as_of_date.isoformat(),
    }


# =========================
# 21. QQ / OneBot 实时接入
# =========================

def load_qq_live_config(config_path: str, require_group_ids: bool = True) -> dict:
    """读取并校验 NapCatQQ / OneBot 11 的本地连接配置。"""
    if not os.path.exists(config_path):
        raise FileNotFoundError(
            f"未找到 QQ 实时配置文件: {config_path}\n"
            "请复制 qq_live_config.example.json 为 qq_live_config.json，并填写群号。"
        )

    with open(config_path, "r", encoding="utf-8") as f:
        config = json.load(f)

    config_dir = os.path.dirname(os.path.abspath(config_path))

    def resolve_config_path(value: str, fallback: str = "") -> str:
        path = str(value or fallback).strip()
        if path and not os.path.isabs(path):
            path = os.path.join(config_dir, path)
        return os.path.abspath(path) if path else ""

    rules_file = resolve_config_path(config.get("rules_file"), RULES_CONFIG_FILE)
    load_rules_config(rules_file)

    ws_url = validate_websocket_url(config.get("ws_url"))

    raw_group_ids = config.get("group_ids", [])
    if not isinstance(raw_group_ids, list):
        raise ValueError("group_ids 必须是 QQ 群号列表。")

    group_ids = {str(group_id).strip() for group_id in raw_group_ids}
    group_ids.discard("")
    if any(not group_id.isdigit() for group_id in group_ids):
        raise ValueError("group_ids 中的 QQ 群号必须是纯数字。")
    if require_group_ids and not group_ids:
        raise ValueError("请在 group_ids 中填写至少一个 QQ 群号。")

    raw_group_names = config.get("group_names", {})
    if not isinstance(raw_group_names, dict):
        raise ValueError("group_names 必须是以群号为键、群名为值的对象。")
    group_names = {
        str(group_id).strip(): str(group_name).strip()
        for group_id, group_name in raw_group_names.items()
        if str(group_id).strip() in group_ids
    }

    raw_default_origins = config.get("group_default_origins", {})
    if not isinstance(raw_default_origins, dict):
        raise ValueError("group_default_origins 必须是以群号为键、始发地为值的对象。")
    group_default_origins = {}
    for group_id in group_ids:
        raw_origin = str(raw_default_origins.get(group_id, DEFAULT_ORIGIN)).strip()
        normalized_origin = normalize_origin(raw_origin)
        if not normalized_origin:
            raise ValueError(f"群 {group_id} 的默认始发地不在支持范围内: {raw_origin}")
        group_default_origins[group_id] = normalized_origin

    raw_exclusions = config.get("group_excluded_origins", {})
    if not isinstance(raw_exclusions, dict):
        raise ValueError("group_excluded_origins 必须是以群号为键、标准始发地列表为值的对象。")
    if set(raw_exclusions) - group_ids:
        raise ValueError("group_excluded_origins 引用了未配置的QQ群。")
    group_excluded_origins = {}
    for group_id, values in raw_exclusions.items():
        blocked = normalize_excluded_origins(
            values, set(ORIGIN_GROUPS), group_default_origins[group_id], group_id,
        )
        if blocked:
            group_excluded_origins[group_id] = blocked

    output_root = str(config.get("output_root", "QQ实时统计结果")).strip()
    if not output_root:
        raise ValueError("output_root 不能为空。")
    if not os.path.isabs(output_root):
        output_root = os.path.join(config_dir, output_root)
    output_root = os.path.abspath(output_root)

    try:
        reconnect_seconds = float(config.get("reconnect_seconds", 5))
    except (TypeError, ValueError) as exc:
        raise ValueError("reconnect_seconds 必须是数字。") from exc
    if not 1 <= reconnect_seconds <= 300:
        raise ValueError("reconnect_seconds 必须在1到300秒之间。")
    try:
        heartbeat_seconds = float(config.get("heartbeat_seconds", 15))
    except (TypeError, ValueError) as exc:
        raise ValueError("heartbeat_seconds 必须是数字。") from exc
    if not 5 <= heartbeat_seconds <= 300:
        raise ValueError("heartbeat_seconds 必须在5到300秒之间。")
    try:
        excel_batch_seconds = float(config.get("excel_batch_seconds", 0.5))
        excel_full_refresh_seconds = float(
            config.get("excel_full_refresh_seconds", 60)
        )
    except (TypeError, ValueError) as exc:
        raise ValueError("Excel更新间隔必须是数字。") from exc
    if not 0 <= excel_batch_seconds <= 10:
        raise ValueError("excel_batch_seconds 必须在0到10秒之间。")
    if not 10 <= excel_full_refresh_seconds <= 3600:
        raise ValueError("excel_full_refresh_seconds 必须在10到3600秒之间。")

    data_lifecycle = normalize_data_lifecycle_config(
        config.get("data_lifecycle"),
        config.get("backup_retention_days", 30),
    )
    group_data_lifecycle = normalize_group_data_lifecycle_configs(
        config.get("group_data_lifecycle"),
        data_lifecycle,
        group_ids,
    )

    dashboard = config.get("status_dashboard", {})
    if not isinstance(dashboard, dict):
        raise ValueError("status_dashboard 必须是对象。")
    time_check_urls = config.get("time_check_urls", [])
    if not isinstance(time_check_urls, list):
        raise ValueError("time_check_urls 必须是网址列表。")

    return {
        "config_path": os.path.abspath(config_path),
        "ws_url": ws_url,
        "access_token": str(config.get("access_token", "")).strip(),
        "group_ids": group_ids,
        "group_names": group_names,
        "group_default_origins": group_default_origins,
        "group_excluded_origins": group_excluded_origins,
        "output_root": output_root,
        "rules_file": rules_file,
        "accept_self_messages": bool(config.get("accept_self_messages", False)),
        "ocr_enabled": bool(config.get("ocr_enabled", False)),
        "ocr_script": resolve_config_path(config.get("ocr_script")),
        "dashboard_enabled": bool(dashboard.get("enabled", True)),
        "dashboard_host": validate_dashboard_host(
            dashboard.get("host", "127.0.0.1")
        ),
        "dashboard_port": validate_tcp_port(
            dashboard.get("port", 8765),
            "管理页面端口",
        ),
        "backup_retention_days": data_lifecycle["backup_retention_days"],
        "data_lifecycle": data_lifecycle,
        "group_data_lifecycle": group_data_lifecycle,
        "time_check_urls": [str(value) for value in time_check_urls],
        "reconnect_seconds": reconnect_seconds,
        "heartbeat_seconds": heartbeat_seconds,
        "excel_batch_seconds": excel_batch_seconds,
        "excel_full_refresh_seconds": excel_full_refresh_seconds,
        "rebuild_on_start": bool(config.get("rebuild_on_start", True))
    }


def safe_folder_component(value: str) -> str:
    """生成可读且适用于 Windows 的单级目录名。"""
    safe_value = re.sub(r'[\\/:*?"<>|]', "_", str(value or ""))
    safe_value = safe_value.strip().rstrip(".")
    return safe_value or "QQ群"


def resolve_qq_group_output_dir(
    output_root: str,
    group_id: str,
    group_name: str,
) -> tuple[str, str]:
    """使用群号稳定定位输出目录，并兼容迁移旧的“群名_群号”目录。"""
    stable_dir = os.path.join(output_root, f"QQ群_{group_id}")
    if os.path.isdir(stable_dir):
        return stable_dir, ""

    suffix = f"_{group_id}"
    expected_legacy = os.path.join(
        output_root,
        f"{safe_folder_component(group_name)}_{group_id}",
    )
    candidates = []
    if os.path.isdir(output_root):
        for entry in os.scandir(output_root):
            if (
                entry.is_dir()
                and entry.path != stable_dir
                and entry.name.endswith(suffix)
            ):
                candidates.append(entry.path)

    if not candidates:
        return stable_dir, ""

    if expected_legacy in candidates:
        source_dir = expected_legacy
    else:
        def candidate_score(path: str) -> tuple[int, float]:
            database_path = os.path.join(path, DATABASE_FILE)
            scored_path = database_path if os.path.exists(database_path) else path
            try:
                modified_at = os.path.getmtime(scored_path)
            except OSError:
                modified_at = 0.0
            return int(os.path.exists(database_path)), modified_at

        source_dir = max(candidates, key=candidate_score)

    if len(candidates) > 1:
        return source_dir, (
            f"检测到群 {group_id} 的多个历史目录，暂时继续使用 "
            f"{source_dir}；请合并其他同群目录后再迁移。"
        )

    try:
        os.replace(source_dir, stable_dir)
        return stable_dir, f"已把历史群目录迁移为稳定目录: {stable_dir}"
    except OSError as exc:
        return source_dir, (
            f"群 {group_id} 的历史目录暂未迁移（{exc}），"
            f"本次继续使用原目录: {source_dir}"
        )


def build_qq_group_profiles(config: dict) -> dict[str, dict]:
    """为每个允许群建立独立输入、状态和统计输出目录。"""
    profiles = {}
    for group_id in sorted(config["group_ids"]):
        group_name = config["group_names"].get(group_id) or f"QQ群_{group_id}"
        output_dir, output_dir_note = resolve_qq_group_output_dir(
            config["output_root"],
            group_id,
            group_name,
        )
        profiles[group_id] = {
            "group_id": group_id,
            "group_name": group_name,
            "default_origin": config["group_default_origins"][group_id],
            "excluded_origins": list(config.get("group_excluded_origins", {}).get(group_id, [])),
            "output_dir": output_dir,
            "output_dir_note": output_dir_note,
            "input_file": os.path.join(output_dir, INPUT_FILE),
            "state_file": os.path.join(output_dir, QQ_LIVE_STATE_FILE),
            "database_file": os.path.join(output_dir, DATABASE_FILE),
            "archive_database": os.path.abspath(os.path.join(config['output_root'], '统计年度管理.db')),
            "rejected_file": os.path.join(output_dir, REJECTED_CSV),
            "backup_retention_days": int(config.get("backup_retention_days", 30)),
            "backup_enabled": "backup_retention_days" in config,
            "data_lifecycle": dict(
                config.get("group_data_lifecycle", {}).get(
                    group_id,
                    config.get("data_lifecycle", {}),
                )
            ),
            "lifecycle_policy_source": (
                "group_override"
                if group_id in config.get("group_data_lifecycle", {})
                else "default_template"
            ),
            "rules_file": config.get("rules_file", ""),
            "config_path": config.get("config_path", "")
        }
    return profiles


def initialize_group_database(profile: dict) -> FreightDatabase:
    """首次升级时把已有TXT解析结果迁移进SQLite；之后只做增量写入。"""
    database = FreightDatabase(profile["database_file"])
    with database.session() as connection:
        initialized = connection.execute("SELECT value FROM metadata WHERE key='txt_migration_done'").fetchone()
    if initialized:
        return database
    if database.count_records() > 0:
        with database.session() as connection:
            connection.execute("INSERT OR REPLACE INTO metadata VALUES ('txt_migration_done','1')")
        return database
    if not os.path.exists(profile["input_file"]) or os.path.getsize(profile["input_file"]) == 0:
        with database.session() as connection:
            connection.execute("INSERT OR REPLACE INTO metadata VALUES ('txt_migration_done','1')")
        return database

    previous_directory = os.getcwd()
    try:
        os.chdir(profile["output_dir"])
        formatted_file = auto_format_input_file(
            INPUT_FILE,
            FORMATTED_INPUT_FILE,
            profile["default_origin"],
        )
        records = deduplicate_records(
            parse_chat_file(formatted_file, profile["default_origin"])
        )
        keyed_records = [
            (build_record_content_dedup_key(record), record)
            for record in records
        ]
        database.insert_records(keyed_records, source="txt_migration")
        with database.session() as connection:
            connection.execute("INSERT OR REPLACE INTO metadata VALUES ('txt_migration_done','1')")
    finally:
        os.chdir(previous_directory)
    return database


def _rebuild_qq_group_output_unlocked(profile: dict) -> dict:
    """在目标群目录中运行原有统计流程，并保证结束后恢复工作目录。"""
    output_dir = profile["output_dir"]
    os.makedirs(output_dir, exist_ok=True)
    database = initialize_group_database(profile)
    with database.session() as connection:
        dirty = connection.execute("SELECT value FROM metadata WHERE key='reports_dirty'").fetchone()
    previous_directory = os.getcwd()
    try:
        os.chdir(output_dir)
        from freight_archive import read_year, rebuild_manual
        if read_year(profile.get('archive_database'))['mode']=='manual':
            summary = rebuild_manual(profile)
        else:
            summary = rebuild_all(
                profile["default_origin"],
                database_file=profile["database_file"],
                history_retention_days=profile.get("data_lifecycle", {}).get(
                    "freight_record_retention_days", 0
                ),
            )
    finally:
        os.chdir(previous_directory)
    complete = (summary.get("excel_updated", True)
                and not summary.get("large_category_archives", {}).get("pending_years")
                and not summary.get("small_category_outputs", {}).get("daily_excel_pending_count")
                and not summary.get("small_category_outputs", {}).get("archive_pending_count"))
    if dirty and complete:
        check_dir = os.path.join(output_dir,'手动年度统计','当前') if summary.get('manual_year') else output_dir
        for root, directories, files in os.walk(check_dir):
            directories[:] = [name for name in directories if name not in {"备份", "logs"}]
            if any("待更新" in name and name.lower().endswith((".csv", ".xlsx")) for name in files):
                complete = False
                break
    if dirty and complete:
        with database.session() as connection:
            connection.execute("DELETE FROM metadata WHERE key='reports_dirty' AND value=?", (dirty["value"],))
    summary["report_pending"] = bool(dirty and not complete)
    if profile.get("backup_enabled"):
        create_daily_backup(
            output_dir,
            profile["database_file"],
            extra_files=[
                profile["input_file"],
                os.path.join(output_dir, REPORT_XLSX),
                profile.get("rules_file", ""),
                profile.get("config_path", ""),
            ],
            retention_days=profile.get("backup_retention_days", 30),
            backup_date=trusted_today(),
            prune_expired=profile.get("data_lifecycle", {}).get("enabled", True),
        )
    return summary


def rebuild_qq_group_output(profile: dict) -> dict:
    """全量构建会改变进程工作目录，因此跨群也必须受全局锁保护。"""
    with group_output_lock(profile), QQ_OUTPUT_BUILD_LOCK:
        return _rebuild_qq_group_output_unlocked(profile)


def _append_detail_csv(path: str, records: list[dict]) -> bool:
    if not records:
        return True
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    write_header = not os.path.exists(path) or os.path.getsize(path) == 0
    try:
        with open(path, "a", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(
                stream,
                fieldnames=DETAIL_COLUMNS,
                extrasaction="ignore",
            )
            if write_header:
                writer.writeheader()
            writer.writerows(format_csv_row(record) for record in records)
        return True
    except PermissionError:
        if LOGGER:
            LOGGER.warning(
                "CSV被占用，已跳过本次增量写入并等待全量刷新: %s",
                path,
            )
        return False


def append_qq_group_excel_records(profile: dict, records: list[dict]) -> dict:
    # A deleted quote can still be present in an older queued append batch.
    with group_output_lock(profile):
        from freight_archive import read_year
        if read_year(profile.get('archive_database'))['mode']=='manual':
            # Year boundaries cannot use the legacy Spring-Festival append routing.
            with QQ_OUTPUT_BUILD_LOCK:
                return _rebuild_qq_group_output_unlocked(profile)
        database = FreightDatabase(profile["database_file"])
        with database.session() as connection:
            deleted_keys = {row[0] for row in connection.execute(
                "SELECT record_id FROM deleted_data WHERE kind='freight'")}
        records = [record for record in records
                   if build_record_content_dedup_key(record) not in deleted_keys]
        return _append_qq_group_excel_records_unlocked(profile, records)


def _append_qq_group_excel_records_unlocked(profile: dict, records: list[dict]) -> dict:
    """把新入库记录追加到明细/待生效表；统计汇总由低频全量刷新完成。"""
    allowed_records = []
    for record in records:
        normalized = normalize_record_cargo(record)
        if validate_cargo_route(normalized)[0]:
            allowed_records.append(normalized)
    records = allowed_records
    output_dir = profile["output_dir"]
    os.makedirs(output_dir, exist_ok=True)
    report_path = os.path.join(output_dir, REPORT_XLSX)
    pending_report_path = os.path.join(output_dir, PENDING_REPORT_XLSX)
    database = FreightDatabase(profile["database_file"])
    today = trusted_today()
    current_spring_year = spring_festival_year_for_date(today)
    current_period_start, _current_period_end = spring_festival_year_window(
        current_spring_year
    )

    def current_counts() -> tuple[int, int]:
        active_count = 0
        pending_count = 0
        for stored_record in database.fetch_records():
            stored_record = normalize_record_cargo(stored_record)
            if not validate_cargo_route(stored_record)[0]:
                continue
            try:
                stored_date = datetime.strptime(
                    stored_record["日期"], "%Y-%m-%d"
                ).date()
            except (KeyError, TypeError, ValueError):
                continue
            if stored_date > today:
                pending_count += 1
            elif stored_date >= current_period_start:
                active_count += 1
        return active_count, pending_count

    if not os.path.exists(report_path) and not os.path.exists(pending_report_path):
        summary = rebuild_qq_group_output(profile)
        return {
            **summary,
            "excel_update_mode": "full_initialize",
            "excel_append_count": len(records),
            "aggregate_state": "ready",
        }

    if not records:
        database.export_rejections_csv(os.path.join(output_dir, REJECTED_CSV))
        active_count, pending_count = current_counts()
        return {
            "active_records": active_count,
            "pending_records": pending_count,
            "rejected_records": len(database.fetch_rejections()),
            "excel_updated": True,
            "as_of_date": today.isoformat(),
            "excel_update_mode": "metadata_only",
            "excel_append_count": 0,
        }

    source_path = (
        pending_report_path
        if os.path.exists(pending_report_path)
        else report_path
    )
    workbook = load_workbook(source_path)
    required_sheets = {"报价明细", "待生效运价"}
    if not required_sheets.issubset(workbook.sheetnames):
        workbook.close()
        summary = rebuild_qq_group_output(profile)
        return {
            **summary,
            "excel_update_mode": "full_repair",
            "excel_append_count": len(records),
            "aggregate_state": "ready",
        }

    excel_dedup_columns = [
        "日期", "发布人", "始发地", "目的地", "货物小类", "报价文本", "原始消息",
    ]
    existing_excel_keys = set()
    for sheet_name in required_sheets:
        worksheet = workbook[sheet_name]
        header_map = {
            str(cell.value): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        if not all(column in header_map for column in excel_dedup_columns):
            continue
        for row_index in range(2, worksheet.max_row + 1):
            existing_excel_keys.add(tuple(
                str(worksheet.cell(row=row_index, column=header_map[column]).value or "")
                for column in excel_dedup_columns
            ))

    active_records = []
    pending_records = []
    archived_period_records = []
    for record in records:
        excel_key = tuple(str(record.get(column, "") or "") for column in excel_dedup_columns)
        if excel_key in existing_excel_keys:
            continue
        existing_excel_keys.add(excel_key)
        try:
            effective_date = datetime.strptime(record["日期"], "%Y-%m-%d").date()
        except (KeyError, TypeError, ValueError):
            effective_date = today
        if effective_date > today:
            target = pending_records
        elif effective_date >= current_period_start:
            target = active_records
        else:
            target = archived_period_records
        target.append(record)

    for sheet_name, target_records in (
        ("报价明细", active_records),
        ("待生效运价", pending_records),
    ):
        worksheet = workbook[sheet_name]
        header_map = {
            str(cell.value): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        for record in target_records:
            row_values = [record.get(column, "") for column in DETAIL_COLUMNS]
            worksheet.append(row_values)
            row_index = worksheet.max_row
            for price_column in PRICE_OUTPUT_COLUMNS:
                column_index = header_map.get(price_column)
                if column_index:
                    worksheet.cell(row=row_index, column=column_index).number_format = (
                        EXCEL_DECIMAL_FORMAT
                    )

    temp_path = os.path.join(output_dir, ".物流报价汇总_append.xlsx")
    try:
        workbook.save(temp_path)
    finally:
        workbook.close()

    excel_updated = True
    try:
        os.replace(temp_path, report_path)
        if os.path.exists(pending_report_path):
            try:
                os.remove(pending_report_path)
            except OSError:
                pass
    except PermissionError:
        os.replace(temp_path, pending_report_path)
        excel_updated = False
        if LOGGER:
            LOGGER.warning(
                "Excel被占用，增量内容已保存到 %s",
                pending_report_path,
            )

    _append_detail_csv(os.path.join(output_dir, DETAIL_CSV), active_records)
    _append_detail_csv(os.path.join(output_dir, PENDING_CSV), pending_records)
    database.export_rejections_csv(os.path.join(output_dir, REJECTED_CSV))
    active_count, pending_count = current_counts()
    return {
        "active_records": active_count,
        "pending_records": pending_count,
        "rejected_records": len(database.fetch_rejections()),
        "excel_updated": excel_updated,
        "as_of_date": today.isoformat(),
        "excel_update_mode": "append",
        "excel_append_count": len(active_records) + len(pending_records),
        "archive_refresh_count": len(archived_period_records),
        "aggregate_state": "pending_refresh",
    }


def extract_onebot_plain_text(event: dict) -> str:
    """从 OneBot 字符串或消息段数组中提取纯文本，不把图片等 CQ 码当作运价。"""
    message = event.get("message")
    text_parts = []

    if isinstance(message, list):
        for segment in message:
            if not isinstance(segment, dict) or segment.get("type") != "text":
                continue
            data = segment.get("data") or {}
            text_parts.append(str(data.get("text", "")))
    elif isinstance(message, str):
        text_parts.append(message)

    if not text_parts:
        text_parts.append(str(event.get("raw_message", "")))

    text = "".join(text_parts)
    text = re.sub(r"\[CQ:[^\]]+\]", " ", text)
    return text.strip()


def extract_onebot_image_sources(event: dict) -> list[str]:
    message = event.get("message")
    if not isinstance(message, list):
        return []
    sources = []
    for segment in message:
        if not isinstance(segment, dict) or segment.get("type") != "image":
            continue
        data = segment.get("data") or {}
        source = data.get("url") or data.get("file") or data.get("path")
        if source and str(source) not in sources:
            sources.append(str(source))
    return sources


def build_onebot_message_key(event: dict, message_text: str) -> str:
    """优先使用 NapCat 消息 ID；缺失时使用稳定摘要。"""
    group_id = str(event.get("group_id", ""))
    message_id = str(event.get("message_id", "")).strip()
    if message_id:
        return f"{group_id}:{message_id}"

    fallback = {
        "group_id": group_id,
        "user_id": str(event.get("user_id", "")),
        "time": str(event.get("time", "")),
        "text": message_text
    }
    payload = json.dumps(fallback, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def build_inbound_event_key(event: dict) -> str:
    """为持久接收队列生成稳定键，优先使用群号和NapCat消息ID。"""
    group_id = str(event.get("group_id", "")).strip()
    message_id = str(event.get("message_id", "")).strip()
    post_type = str(event.get("post_type", "")).strip()
    if group_id and message_id:
        return f"{post_type}:{group_id}:{message_id}"
    payload = json.dumps(event, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def classify_onebot_group_event(event: object) -> tuple[str, str]:
    """兼容 OneBot 普通群消息和 NapCat 自身消息事件。"""
    if not isinstance(event, dict):
        return "", "事件不是JSON对象"

    post_type = str(event.get("post_type", "")).strip().lower()
    if post_type not in SUPPORTED_ONEBOT_MESSAGE_POST_TYPES:
        return post_type, "不是支持的消息事件"

    message_type = str(event.get("message_type", "")).strip().lower()
    if message_type and message_type != "group":
        return post_type, "不是群消息"

    if not str(event.get("group_id", "")).strip():
        return post_type, "群消息缺少群号"

    return post_type, ""


def load_seen_qq_message_keys(state_file: str) -> list[str]:
    if not os.path.exists(state_file):
        return []

    try:
        with open(state_file, "r", encoding="utf-8") as f:
            data = json.load(f)
        keys = data.get("seen_message_keys", [])
        if isinstance(keys, list):
            return [str(key) for key in keys[-MAX_SEEN_QQ_MESSAGES:]]
    except (OSError, json.JSONDecodeError, TypeError):
        print(f"警告：状态文件损坏，将重新建立: {state_file}")
    return []


def save_seen_qq_message_keys(state_file: str, keys: list[str]) -> None:
    """原子更新状态文件，避免程序中断后留下半个 JSON。"""
    state_dir = os.path.dirname(os.path.abspath(state_file))
    os.makedirs(state_dir, exist_ok=True)
    temp_file = state_file + ".tmp"
    data = {"seen_message_keys": keys[-MAX_SEEN_QQ_MESSAGES:]}
    with open(temp_file, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    os.replace(temp_file, state_file)


class OneBotFreightIngestor:
    """把允许群消息增量写入SQLite；内容查重同时覆盖正式和未来记录。"""

    def __init__(
        self,
        group_profiles: dict[str, dict],
        accept_self_messages: bool = False,
        ocr: WindowsOcr | None = None,
        logger=None,
    ) -> None:
        self.group_profiles = group_profiles
        self.accept_self_messages = accept_self_messages
        self.ocr = ocr
        self.logger = logger
        self.databases = {}
        for group_id, profile in group_profiles.items():
            os.makedirs(profile["output_dir"], exist_ok=True)
            self.databases[group_id] = initialize_group_database(profile)

    def _record_rejection(
        self,
        database: FreightDatabase,
        message_key: str,
        group_id: str,
        sender: str,
        message_time: datetime,
        reason: str,
        message_text: str,
        media_source: str = "",
    ) -> None:
        payload = f"{message_key}|{reason}|{message_text}|{media_source}"
        rejection_key = hashlib.sha256(payload.encode("utf-8")).hexdigest()
        database.add_rejection(
            rejection_key,
            message_key,
            group_id,
            sender,
            message_time.isoformat(timespec="seconds"),
            reason,
            message_text,
            media_source,
        )

    def ingest(self, event: dict) -> dict:
        event_type, event_error = classify_onebot_group_event(event)
        if event_error:
            return {
                "status": "ignored",
                "reason": event_error,
                "event_type": event_type,
            }

        group_id = str(event.get("group_id", ""))
        if group_id not in self.group_profiles:
            return {
                "status": "ignored",
                "reason": "群号不在允许列表",
                "event_type": event_type,
                "group_id": group_id,
            }
        profile = self.group_profiles[group_id]
        database = self.databases[group_id]

        if not self.accept_self_messages:
            self_id = str(event.get("self_id", "")).strip()
            user_id = str(event.get("user_id", "")).strip()
            if self_id and user_id and self_id == user_id:
                return {
                    "status": "ignored",
                    "reason": "登录账号自己的消息已关闭",
                    "event_type": event_type,
                    "group_id": group_id,
                }

        event_time = event.get("time", time.time())
        try:
            message_time = datetime.fromtimestamp(float(event_time))
        except (TypeError, ValueError, OSError):
            message_time = CLOCK.now().replace(tzinfo=None)

        sender_info = event.get("sender") or {}
        sender = (
            sender_info.get("card")
            or sender_info.get("nickname")
            or event.get("user_id")
            or "QQ用户"
        )
        sender = re.sub(r"[:\r\n]", "_", str(sender)).strip() or "QQ用户"

        plain_text = extract_onebot_plain_text(event)
        image_sources = extract_onebot_image_sources(event)
        ocr_texts = []
        ocr_errors = []
        ocr_active = bool(self.ocr and self.ocr.enabled)
        if image_sources and ocr_active:
            for source in image_sources:
                text, error = self.ocr.recognize(source)
                if text:
                    ocr_texts.append(text)
                elif error:
                    ocr_errors.append((source, error))
        combined_text = "\n".join(value for value in [plain_text, *ocr_texts] if value).strip()
        key_material = combined_text or "|".join(image_sources)
        message_key = build_onebot_message_key(event, key_material)
        if database.has_processed_message(message_key):
            return {
                "status": "duplicate",
                "event_type": event_type,
                "group_id": group_id,
            }

        if image_sources and not combined_text and not ocr_active:
            database.mark_processed_message(message_key)
            return {"status": "ignored", "reason": "仅采集文字，图片消息已跳过",
                    "event_type": event_type, "group_id": group_id,
                    "group_name": profile["group_name"], "sender": sender}

        for source, error in ocr_errors:
            self._record_rejection(
                database, message_key, group_id, sender, message_time,
                f"图片OCR失败: {error}", plain_text, source,
            )

        if not combined_text:
            self._record_rejection(
                database, message_key, group_id, sender, message_time,
                "没有可识别的文字或图片OCR文本", "", "|".join(image_sources),
            )
            database.mark_processed_message(message_key)
            return {
                "status": "rejected",
                "reason": "没有可识别文字",
                "event_type": event_type,
                "group_id": group_id,
                "group_name": profile["group_name"],
                "sender": sender,
            }

        parsed_records = []
        raw_valid_lines = []
        rejected_lines = []
        received_line_count = 0
        for raw_line in combined_text.splitlines():
            line = normalize_text(raw_line)
            if not line:
                continue
            received_line_count += 1
            line_records = []
            line_reasons = []
            expanded_dates = expand_relative_freight_dates(
                line,
                message_time.date().isoformat(),
            )
            if not expanded_dates:
                line_reasons.append("日期无效或不明确，请使用完整年月日或明天/后天")
            for effective_date, freight_text in expanded_dates:
                formatted, format_reason = format_freight_line_with_reason(
                    freight_text,
                    profile["default_origin"],
                )
                if not formatted:
                    if format_reason:
                        line_reasons.append(format_reason)
                    continue
                record = parse_freight_line(
                    formatted,
                    effective_date,
                    sender,
                    profile["default_origin"],
                )
                if record:
                    record = normalize_record_cargo(record)
                    # Check the parsed canonical origin, after aliases/default fallback.
                    # Ingestion only: changing this setting never deletes/relabels history.
                    if record["始发地"] in profile.get("excluded_origins", []):
                        line_reasons.append(f"本群不统计始发地为{record['始发地']}的运价（群始发地过滤）")
                        continue
                    allowed, scope_reason = validate_cargo_route(record)
                    if allowed:
                        line_records.append(record)
                    else:
                        line_reasons.append(scope_reason)
                else:
                    line_reasons.append("格式化后未生成统计记录")
            if line_records:
                raw_valid_lines.append(line)
                parsed_records.extend(line_records)
            else:
                reason = next(iter(dict.fromkeys(line_reasons)), "不是可识别的运价信息")
                rejected_lines.append((line, reason))

        for line, reason in rejected_lines:
            self._record_rejection(
                database, message_key, group_id, sender, message_time,
                reason, line, "|".join(image_sources),
            )

        unique_records = {}
        for record in parsed_records:
            unique_records.setdefault(build_record_content_dedup_key(record), record)
        keyed_records = list(unique_records.items())
        existing_keys = database.existing_record_keys(list(unique_records))
        new_records = [
            record
            for dedup_key, record in keyed_records
            if dedup_key not in existing_keys
        ]
        pending_inserted = sum(
            1
            for dedup_key, record in keyed_records
            if dedup_key not in existing_keys
            and datetime.strptime(record["日期"], "%Y-%m-%d").date() > trusted_today()
        )
        inserted, duplicate_content = database.insert_records(
            keyed_records,
            message_key=message_key,
            source="qq_ocr" if ocr_texts else "qq_text",
        )
        database.mark_processed_message(message_key)

        if inserted > 0:
            input_file = profile["input_file"]
            os.makedirs(os.path.dirname(os.path.abspath(input_file)), exist_ok=True)
            needs_separator = os.path.exists(input_file) and os.path.getsize(input_file) > 0
            with open(input_file, "a", encoding="utf-8") as stream:
                if needs_separator:
                    stream.write("\n")
                stream.write(f"{sender}: {message_time:%Y-%m-%d %H:%M:%S}\n")
                for line in raw_valid_lines:
                    stream.write(line + "\n")

        if inserted == 0 and duplicate_content > 0:
            status = "duplicate_content"
        elif inserted > 0:
            status = "added"
        else:
            status = "rejected"

        rejected_reasons = list(dict.fromkeys(
            reason for _line, reason in rejected_lines
        ))
        if status == "added":
            result_reason = f"新增 {inserted} 条运价"
            if rejected_lines:
                result_reason += (
                    f"；另有 {len(rejected_lines)} 行未识别: "
                    f"{rejected_reasons[0]}"
                )
        elif status == "duplicate_content":
            result_reason = f"{duplicate_content} 条内容与已有记录重复"
        else:
            result_reason = "；".join(rejected_reasons[:3]) or "不是可识别的运价信息"

        return {
            "status": status,
            "reason": result_reason,
            "event_type": event_type,
            "group_id": group_id,
            "group_name": profile["group_name"],
            "default_origin": profile["default_origin"],
            "output_dir": profile["output_dir"],
            "sender": sender,
            "received_line_count": received_line_count,
            "line_count": len(parsed_records),
            "rejected_line_count": len(rejected_lines),
            "rejected_reasons": rejected_reasons,
            "inserted_count": inserted,
            "duplicate_content_count": duplicate_content,
            "pending_count": pending_inserted,
            "ocr_used": bool(ocr_texts),
            "_inserted_records": new_records[:inserted],
        }


def update_runtime_event_status(
    runtime_status: RuntimeStatus,
    group_profiles: dict[str, dict],
    event: object,
    result: dict,
) -> None:
    """先记录事件处理结果，再执行较慢的Excel重建，保证页面及时反馈。"""
    event_data = event if isinstance(event, dict) else {}
    event_at = CLOCK.now().isoformat(timespec="seconds")
    group_id = str(result.get("group_id") or event_data.get("group_id") or "").strip()
    event_type = str(
        result.get("event_type") or event_data.get("post_type") or "unknown"
    ).strip()
    event_status = str(result.get("status") or "unknown").strip()
    event_reason = str(result.get("reason") or "").strip()
    if not event_reason:
        event_reason = {
            "duplicate": "消息已处理",
            "ignored": "事件已忽略",
        }.get(event_status, "-")
    event_reason = event_reason[:300]
    message_id = str(event_data.get("message_id", "")).strip()
    runtime_status.update(
        last_event_at=event_at,
        last_event_type=event_type,
        last_event_status=event_status,
        last_event_reason=event_reason,
        last_group_id=group_id,
    )
    if group_id not in group_profiles:
        return
    runtime_status.update_group(
        group_id,
        last_message=event_at,
        last_event_at=event_at,
        last_event_type=event_type,
        last_event_status=event_status,
        last_event_reason=event_reason,
        last_message_id=message_id,
    )
    if event_status in {"added", "rejected"}:
        inserted_count = int(result.get("inserted_count", 0) or 0)
        pending_count = int(result.get("pending_count", 0) or 0)
        rejected_count = int(result.get("rejected_line_count", 0) or 0)
        runtime_status.increment_group(
            group_id,
            active_records=max(0, inserted_count - pending_count),
            pending_records=pending_count,
            rejected_records=rejected_count,
        )


class ExcelRebuildCoordinator:
    """按群隔离Excel工作线程；同群严格串行，群间可并行追加。"""

    def __init__(
        self,
        group_profiles: dict[str, dict],
        runtime_status: RuntimeStatus,
        rebuild_callback=rebuild_qq_group_output,
        incremental_callback=append_qq_group_excel_records,
        logger=None,
        batch_delay_seconds: float = 0.5,
        full_refresh_seconds: float = 60.0,
    ) -> None:
        self.group_profiles = group_profiles
        self.runtime_status = runtime_status
        self.rebuild_callback = rebuild_callback
        self.incremental_callback = incremental_callback
        self.logger = logger
        self.batch_delay_seconds = max(0.0, float(batch_delay_seconds))
        self.full_refresh_seconds = max(1.0, float(full_refresh_seconds))
        self._queues = {
            group_id: queue.Queue(maxsize=1)
            for group_id in group_profiles
        }
        self._lock = threading.Lock()
        self._pending = set()
        self._active = set()
        self._dirty = set()
        self._records = defaultdict(list)
        self._force_full = set()
        self._last_full = {}
        self._full_timers = {}
        self._stopped = False
        self._accepting = True
        now = time.monotonic()
        for group_id, profile in group_profiles.items():
            if os.path.exists(os.path.join(profile["output_dir"], REPORT_XLSX)):
                self._last_full[group_id] = now
        self._threads = []
        for group_id in group_profiles:
            thread = threading.Thread(
                target=self._run_group,
                args=(group_id,),
                name=f"freight-excel-{group_id}",
                daemon=True,
            )
            thread.start()
            self._threads.append(thread)

    def _wake_group(self, group_id: str) -> None:
        try:
            self._queues[group_id].put_nowait(True)
        except queue.Full:
            pass

    def request(
        self,
        group_id: str,
        records: list[dict] | None = None,
        force_full: bool = False,
    ) -> bool:
        group_id = str(group_id)
        if group_id not in self.group_profiles:
            return False
        with self._lock:
            if self._stopped or not self._accepting:
                return False
            if records:
                self._records[group_id].extend(records)
            if force_full:
                self._force_full.add(group_id)
            if group_id in self._pending:
                self._dirty.add(group_id)
            else:
                self._pending.add(group_id)
                self._wake_group(group_id)
            queue_depth = 1 if group_id in self._pending else 0
        self.runtime_status.update_group(
            group_id,
            excel_state="queued",
            excel_update_mode="full" if force_full else "append",
            excel_queue_depth=queue_depth,
        )
        return True

    def _schedule_full_refresh(self, group_id: str) -> None:
        def trigger() -> None:
            self.request(group_id, force_full=True)

        with self._lock:
            if self._stopped or not self._accepting:
                return
            old_timer = self._full_timers.pop(group_id, None)
            if old_timer:
                old_timer.cancel()
            timer = threading.Timer(self.full_refresh_seconds, trigger)
            timer.daemon = True
            self._full_timers[group_id] = timer
            timer.start()

    def _run_group(self, group_id: str) -> None:
        target_queue = self._queues[group_id]
        while True:
            signal = target_queue.get()
            if signal is None:
                target_queue.task_done()
                return
            if self.batch_delay_seconds:
                time.sleep(self.batch_delay_seconds)
            with self._lock:
                self._active.add(group_id)
                records = self._records.pop(group_id, [])
                force_full = group_id in self._force_full
                self._force_full.discard(group_id)
                self._dirty.discard(group_id)
                last_full = self._last_full.get(group_id)
            report_exists = os.path.exists(os.path.join(
                self.group_profiles[group_id]["output_dir"],
                REPORT_XLSX,
            ))
            full_due = (
                force_full
                or not report_exists
                or last_full is None
                or time.monotonic() - last_full >= self.full_refresh_seconds
            )
            update_mode = "full" if full_due else "append"
            started = time.monotonic()
            self.runtime_status.update_group(
                group_id,
                excel_state="updating",
                excel_update_mode=update_mode,
                excel_queue_depth=1,
                excel_error="",
            )
            failed = False
            try:
                if full_due:
                    summary = self.rebuild_callback(self.group_profiles[group_id])
                    summary = {
                        **summary,
                        "excel_update_mode": "full",
                        "aggregate_state": "ready",
                    }
                    with self._lock:
                        self._last_full[group_id] = time.monotonic()
                        timer = self._full_timers.pop(group_id, None)
                        if timer:
                            timer.cancel()
                    if summary.get("report_pending"):
                        self._schedule_full_refresh(group_id)
                else:
                    summary = self.incremental_callback(
                        self.group_profiles[group_id],
                        records,
                    )
                    if records:
                        self._schedule_full_refresh(group_id)
                self.runtime_status.update_group(
                    group_id,
                    name=self.group_profiles[group_id]["group_name"],
                    excel_state="ready",
                    excel_last_updated_at=CLOCK.now().isoformat(timespec="seconds"),
                    excel_processing_ms=round((time.monotonic() - started) * 1000, 1),
                    **summary,
                )
            except Exception as exc:
                failed = True
                error_text = f"群 {group_id} Excel更新失败: {exc}"
                with self._lock:
                    if records:
                        self._records[group_id][0:0] = records
                    self._force_full.add(group_id)
                    self._dirty.add(group_id)
                self.runtime_status.update_group(
                    group_id,
                    excel_state="error",
                    excel_error=str(exc)[:300],
                )
                self.runtime_status.add_error(error_text)
                if self.logger:
                    self.logger.exception(error_text)
                time.sleep(1.0)
            finally:
                with self._lock:
                    self._active.discard(group_id)
                    has_more = (
                        group_id in self._dirty
                        or bool(self._records.get(group_id))
                        or group_id in self._force_full
                    )
                    if has_more and not self._stopped:
                        self._wake_group(group_id)
                    else:
                        self._pending.discard(group_id)
                    queue_depth = 1 if group_id in self._pending else 0
                self.runtime_status.update_group(
                    group_id,
                    excel_queue_depth=queue_depth,
                )
                target_queue.task_done()

    def wait_idle(self, timeout: float = 10.0) -> bool:
        deadline = time.monotonic() + max(0.0, timeout)
        while time.monotonic() <= deadline:
            with self._lock:
                idle = (
                    not self._pending
                    and not self._active
                    and all(value.unfinished_tasks == 0 for value in self._queues.values())
                )
            if idle:
                return True
            time.sleep(0.01)
        return False

    def stop(self, drain: bool = True, timeout: float = 10.0) -> bool:
        with self._lock:
            if self._stopped:
                return all(not thread.is_alive() for thread in self._threads)
            self._accepting = False
        if drain and not self.wait_idle(timeout):
            with self._lock:
                pending_groups = list(self._pending)
            for group_id in pending_groups:
                self.runtime_status.update_group(
                    group_id,
                    excel_state="drain_timeout",
                )
            return False
        with self._lock:
            self._stopped = True
            timers = list(self._full_timers.values())
            self._full_timers.clear()
        for timer in timers:
            timer.cancel()
        for target_queue in self._queues.values():
            target_queue.put(None)
        deadline = time.monotonic() + max(0.1, timeout)
        for thread in self._threads:
            thread.join(timeout=max(0.0, deadline - time.monotonic()))
        return all(not thread.is_alive() for thread in self._threads)


class DataLifecycleManager:
    """按群独立执行生命周期维护；群策略缺失时继承默认模板。"""

    def __init__(
        self,
        group_profiles: dict[str, dict],
        runtime_status: RuntimeStatus,
        policy: dict,
        refresh_callback=None,
        logger=None,
        interval_seconds: float | None = None,
    ) -> None:
        self.group_profiles = group_profiles
        self.runtime_status = runtime_status
        self.default_policy = normalize_data_lifecycle_config(
            policy,
            (policy or {}).get("backup_retention_days", 30)
            if isinstance(policy, dict) else 30,
        )
        self.policies = {
            group_id: normalize_data_lifecycle_config(
                profile.get("data_lifecycle", self.default_policy),
                self.default_policy["backup_retention_days"],
            )
            for group_id, profile in group_profiles.items()
        }
        self.refresh_callback = refresh_callback
        self.logger = logger
        self.interval_seconds = (
            max(0.01, float(interval_seconds)) if interval_seconds is not None else None
        )
        self._stop_event = threading.Event()
        self._deleted_totals = defaultdict(int)
        self._threads = []
        for group_id, profile in group_profiles.items():
            group_policy = self.policies[group_id]
            policy_source = profile.get(
                "lifecycle_policy_source", "default_template"
            )
            if not group_policy["enabled"]:
                self.runtime_status.update_group(
                    group_id,
                    lifecycle_state="disabled",
                    lifecycle_policy=dict(group_policy),
                    lifecycle_policy_source=policy_source,
                )
                continue
            self.runtime_status.update_group(
                group_id,
                lifecycle_state="starting",
                lifecycle_policy=dict(group_policy),
                lifecycle_policy_source=policy_source,
            )
            thread = threading.Thread(
                target=self._run_group,
                args=(group_id,),
                name=f"freight-data-lifecycle-{group_id}",
                daemon=True,
            )
            thread.start()
            self._threads.append(thread)

    def _group_interval_seconds(self, group_id: str) -> float:
        if self.interval_seconds is not None:
            return self.interval_seconds
        return float(self.policies[group_id]["maintenance_interval_minutes"] * 60)

    def _run_group(self, group_id: str) -> None:
        while not self._stop_event.is_set():
            self._run_group_once(group_id)
            if self._stop_event.wait(self._group_interval_seconds(group_id)):
                break

    def run_once(self) -> dict[str, dict]:
        results = {}
        for group_id in self.group_profiles:
            if self._stop_event.is_set():
                break
            if self.policies[group_id]["enabled"]:
                result = self._run_group_once(group_id)
                if result is not None:
                    results[group_id] = result
        return results

    def _run_group_once(self, group_id: str) -> dict | None:
        profile = self.group_profiles[group_id]
        group_policy = self.policies[group_id]
        started_at = CLOCK.now()
        self.runtime_status.update_group(
            group_id,
            lifecycle_state="running",
            lifecycle_started_at=started_at.isoformat(timespec="seconds"),
        )
        try:
            result = FreightDatabase(profile["database_file"]).cleanup_lifecycle(
                reference_time=started_at,
                processed_message_retention_days=group_policy[
                    "processed_message_retention_days"
                ],
                rejected_message_retention_days=group_policy[
                    "rejected_message_retention_days"
                ],
                dead_letter_retention_days=group_policy[
                    "dead_letter_retention_days"
                ],
                freight_record_retention_days=group_policy[
                    "freight_record_retention_days"
                ],
            )
            backup_deleted = prune_daily_backups(
                profile["output_dir"],
                group_policy["backup_retention_days"],
                started_at.date(),
            )
            result["deleted"]["daily_backups"] = backup_deleted
            result["deleted_total"] += backup_deleted
            self._deleted_totals[group_id] += result["deleted_total"]
            report_changed = bool(
                result["deleted"]["freight_records"]
                or result["deleted"]["rejected_messages"]
            )
            if report_changed and self.refresh_callback:
                self.refresh_callback(group_id, force_full=True)
            completed_at = CLOCK.now()
            self.runtime_status.update_group(
                group_id,
                lifecycle_state="ready",
                lifecycle_last_run_at=completed_at.isoformat(timespec="seconds"),
                lifecycle_next_run_at=(
                    completed_at
                    + timedelta(seconds=self._group_interval_seconds(group_id))
                ).isoformat(timespec="seconds"),
                lifecycle_last_deleted_total=result["deleted_total"],
                lifecycle_deleted_total=self._deleted_totals[group_id],
                lifecycle_deleted=result["deleted"],
                lifecycle_remaining=result["remaining"],
                lifecycle_database_bytes=result["database_bytes"],
                lifecycle_last_error="",
            )
            if result["deleted_total"] and self.logger:
                self.logger.info(
                    "群 %s 生命周期清理完成(%s): %s",
                    group_id,
                    profile.get("lifecycle_policy_source", "default_template"),
                    result["deleted"],
                )
            return result
        except Exception as exc:
            error_text = f"群 {group_id} 生命周期维护失败: {exc}"
            self.runtime_status.update_group(
                group_id,
                lifecycle_state="error",
                lifecycle_last_run_at=CLOCK.now().isoformat(timespec="seconds"),
                lifecycle_last_error=str(exc)[:300],
            )
            self.runtime_status.add_error(error_text)
            if self.logger:
                self.logger.exception(error_text)
            return None

    def stop(self, timeout: float = 30.0) -> bool:
        self._stop_event.set()
        deadline = time.monotonic() + max(0.0, timeout)
        for thread in self._threads:
            thread.join(timeout=max(0.0, deadline - time.monotonic()))
        return all(not thread.is_alive() for thread in self._threads)


class GroupMessageProcessor:
    """每群使用SQLite持久队列；内存队列只负责有界唤醒。"""

    def __init__(
        self,
        group_profiles: dict[str, dict],
        ingestor: OneBotFreightIngestor,
        runtime_status: RuntimeStatus,
        result_callback=None,
        logger=None,
        collection_control=None,
    ) -> None:
        self.group_profiles = group_profiles
        self.ingestor = ingestor
        self.runtime_status = runtime_status
        self.result_callback = result_callback
        self.logger = logger
        self.collection_control = collection_control
        self._queues = {
            group_id: queue.Queue(maxsize=1)
            for group_id in group_profiles
        }
        self._lock = threading.Lock()
        self._received = defaultdict(int)
        self._processed = defaultdict(int)
        self._active = set()
        self._accepting = True
        self._stop_event = threading.Event()
        self._threads = []
        ingestor_databases = getattr(ingestor, "databases", {})
        self._databases = {}
        for group_id in group_profiles:
            database = ingestor_databases.get(group_id)
            if database is None:
                database_path = group_profiles[group_id].get("database_file")
                if not database_path:
                    raise ValueError(f"群 {group_id} 缺少持久队列数据库路径。")
                database = FreightDatabase(database_path)
            database.reset_processing_inbound_events()
            self._databases[group_id] = database
            thread = threading.Thread(
                target=self._run_group,
                args=(group_id,),
                name=f"freight-group-{group_id}",
                daemon=True,
            )
            thread.start()
            self._threads.append(thread)
            pending_count = database.count_pending_inbound_events()
            if pending_count:
                self._wake_group(group_id)
            self.runtime_status.update_group(
                group_id,
                queue_state="queued" if pending_count else "idle",
                queue_depth=pending_count,
                queue_durable=True,
                queue_dead_letters=database.count_dead_letter_inbound_events(),
                received_events=0,
                processed_events=0,
            )

    def _wake_group(self, group_id: str) -> None:
        try:
            self._queues[group_id].put_nowait(True)
        except queue.Full:
            pass

    def submit(self, event: object) -> dict:
        if self.collection_control is not None:
            return self.collection_control.submit(event, lambda: self._submit_admitted(event))
        return self._submit_admitted(event)

    def _submit_admitted(self, event: object) -> dict:
        event_data = event if isinstance(event, dict) else {}
        group_id = str(event_data.get("group_id", "")).strip()
        target_queue = self._queues.get(group_id)
        if target_queue is None:
            result = self.ingestor.ingest(event)
            update_runtime_event_status(
                self.runtime_status,
                self.group_profiles,
                event,
                result,
            )
            return result

        if not isinstance(event, dict):
            return {
                "status": "ignored",
                "reason": "事件不是JSON对象",
                "group_id": group_id,
            }
        encoded_size = len(json.dumps(event, ensure_ascii=False).encode("utf-8"))
        if encoded_size > 2_000_000:
            self.runtime_status.add_error(f"群 {group_id} 事件超过2MB，已拒绝。")
            return {
                "status": "rejected",
                "reason": "事件超过2MB安全上限",
                "group_id": group_id,
            }
        with self._lock:
            if not self._accepting:
                return {
                    "status": "ignored",
                    "reason": "消息处理器已停止",
                    "group_id": group_id,
                }
            self._received[group_id] += 1
            received_count = self._received[group_id]
        event_key = build_inbound_event_key(event)
        persisted = self._databases[group_id].enqueue_inbound_event(event_key, event)
        queue_depth = self._databases[group_id].count_pending_inbound_events()
        self._wake_group(group_id)
        event_at = CLOCK.now().isoformat(timespec="seconds")
        self.runtime_status.update(
            last_received_at=event_at,
            last_received_group_id=group_id,
        )
        self.runtime_status.update_group(
            group_id,
            queue_state="queued",
            queue_depth=queue_depth,
            queue_durable=True,
            received_events=received_count,
            last_received_at=event_at,
        )
        return {
            "status": "queued" if persisted else "duplicate_queue_event",
            "group_id": group_id,
            "queue_depth": queue_depth,
            "durable": True,
        }

    def _run_group(self, group_id: str) -> None:
        target_queue = self._queues[group_id]
        database = self._databases[group_id]
        while not self._stop_event.is_set():
            try:
                target_queue.get(timeout=0.5)
                target_queue.task_done()
            except queue.Empty:
                pass

            while not self._stop_event.is_set():
                claimed = database.claim_next_inbound_event()
                if claimed is None:
                    break
                event = claimed["event"]
                event_id = claimed["id"]
                started = time.monotonic()
                with self._lock:
                    self._active.add(group_id)
                remaining = database.count_pending_inbound_events()
                self.runtime_status.update_group(
                    group_id,
                    queue_state="processing",
                    queue_depth=remaining,
                )
                try:
                    result = self.ingestor.ingest(event)
                    update_runtime_event_status(
                        self.runtime_status,
                        self.group_profiles,
                        event,
                        result,
                    )
                    if self.result_callback:
                        self.result_callback(event, result)
                    database.complete_inbound_event(event_id)
                    with self._lock:
                        self._processed[group_id] += 1
                        processed_count = self._processed[group_id]
                except Exception as exc:
                    retry_state = database.fail_inbound_event(event_id, str(exc))
                    processed_count = self._processed[group_id]
                    error_text = f"群 {group_id} 消息处理失败({retry_state}): {exc}"
                    self.runtime_status.update_group(
                        group_id,
                        last_event_status="error",
                        last_event_reason=str(exc)[:300],
                    )
                    self.runtime_status.add_error(error_text)
                    if self.logger:
                        self.logger.exception(error_text)
                finally:
                    with self._lock:
                        self._active.discard(group_id)
                    remaining = database.count_pending_inbound_events()
                    self.runtime_status.update_group(
                        group_id,
                        queue_state="queued" if remaining else "idle",
                        queue_depth=remaining,
                        queue_dead_letters=database.count_dead_letter_inbound_events(),
                        processed_events=processed_count,
                        last_processing_ms=round(
                            (time.monotonic() - started) * 1000,
                            1,
                        ),
                    )
    def wait_idle(self, timeout: float = 10.0) -> bool:
        deadline = time.monotonic() + max(0.0, timeout)
        while time.monotonic() <= deadline:
            with self._lock:
                active = bool(self._active)
            if not active and all(
                database.count_pending_inbound_events() == 0
                for database in self._databases.values()
            ):
                return True
            time.sleep(0.01)
        return False

    def stop(self, drain: bool = True, timeout: float = 10.0) -> bool:
        with self._lock:
            self._accepting = False
        if drain and not self.wait_idle(timeout):
            for group_id, database in self._databases.items():
                self.runtime_status.update_group(
                    group_id,
                    queue_state="drain_timeout",
                    queue_depth=database.count_pending_inbound_events(),
                )
            return False
        self._stop_event.set()
        for group_id in self._queues:
            self._wake_group(group_id)
        deadline = time.monotonic() + max(0.1, timeout)
        for thread in self._threads:
            thread.join(timeout=max(0.0, deadline - time.monotonic()))
        stopped = all(not thread.is_alive() for thread in self._threads)
        if stopped:
            for database in self._databases.values():
                database.close_inbound_event_queue()
        return stopped


def build_recent_data_snapshot(
    group_profiles: dict[str, dict],
    limit: int = 80,
) -> dict:
    """汇总各群最近的正式、待生效和不合格消息，供只读管理页展示。"""
    safe_limit = max(1, min(int(limit), 200))
    today = trusted_today()
    items = []
    for group_id, profile in group_profiles.items():
        database = FreightDatabase(profile["database_file"])
        for row in database.fetch_recent_records(safe_limit):
            record = normalize_record_cargo(row["record"])
            if not validate_cargo_route(record)[0]:
                continue
            effective_date_text = str(record.get("日期", ""))
            try:
                effective_date = datetime.strptime(
                    effective_date_text,
                    "%Y-%m-%d",
                ).date()
            except ValueError:
                effective_date = today
            kind = "待生效" if effective_date > today else "正式识别"
            items.append({
                "kind": kind,
                "group_id": group_id,
                "group_name": profile["group_name"],
                "received_at": row["created_at"],
                "effective_date": effective_date_text,
                "sender": record.get("发布人", ""),
                "route": f"{record.get('始发地', '')}到{record.get('目的地', '')}",
                "cargo": (
                    f"{record.get('货物小类', '')} / {record.get('货物大类', '')}"
                ).strip(" /"),
                "price": record.get("报价文本", ""),
                "detail": record.get("原始消息", ""),
            })
        for rejection in database.fetch_rejections(safe_limit):
            reason = str(rejection.get("原因", ""))
            message_text = str(rejection.get("消息文本", ""))
            detail = f"{reason}｜{message_text}" if message_text else reason
            items.append({
                "kind": "不合格",
                "group_id": group_id,
                "group_name": profile["group_name"],
                "received_at": rejection.get("记录时间", ""),
                "effective_date": "",
                "sender": rejection.get("发布人", ""),
                "route": "",
                "cargo": "",
                "price": "",
                "detail": detail,
            })
    items.sort(key=lambda item: str(item.get("received_at", "")), reverse=True)
    selected_items = items[:safe_limit]
    return {
        "updated_at": CLOCK.now().isoformat(timespec="seconds"),
        "limit": safe_limit,
        "items": selected_items,
        "counts": {
            kind: sum(1 for item in selected_items if item["kind"] == kind)
            for kind in ("正式识别", "待生效", "不合格")
        },
    }


def validate_qq_live_config_without_applying(config_path: str) -> None:
    """校验保存后的配置，同时恢复当前进程正在使用的业务规则。"""
    global DEFAULT_ORIGIN, ORIGIN_GROUPS, DEST_GROUPS, BOARD_TYPES
    global DEFAULT_CARGO_SUBCATEGORY, DEFAULT_CARGO_CATEGORY
    global CARGO_TYPES, CARGO_ROUTE_SCOPES
    global PRICE_THRESHOLD, INVALID_LINE_KEYWORDS, RELATIVE_DATE_PATTERNS

    previous_rules = (
        DEFAULT_ORIGIN,
        ORIGIN_GROUPS,
        DEST_GROUPS,
        BOARD_TYPES,
        DEFAULT_CARGO_SUBCATEGORY,
        DEFAULT_CARGO_CATEGORY,
        CARGO_TYPES,
        CARGO_ROUTE_SCOPES,
        PRICE_THRESHOLD,
        INVALID_LINE_KEYWORDS,
        RELATIVE_DATE_PATTERNS,
    )
    try:
        load_qq_live_config(config_path)
    finally:
        (
            DEFAULT_ORIGIN,
            ORIGIN_GROUPS,
            DEST_GROUPS,
            BOARD_TYPES,
            DEFAULT_CARGO_SUBCATEGORY,
            DEFAULT_CARGO_CATEGORY,
            CARGO_TYPES,
            CARGO_ROUTE_SCOPES,
            PRICE_THRESHOLD,
            INVALID_LINE_KEYWORDS,
            RELATIVE_DATE_PATTERNS,
        ) = previous_rules


def _run_qq_live_session(config_path: str) -> bool:
    """连接 NapCat 正向 WebSocket，实时接收指定群的运价信息。"""
    global LOGGER
    try:
        import websocket
    except ImportError as exc:
        raise RuntimeError(
            "缺少 websocket-client，请按项目文档安装 requirements.txt。"
        ) from exc

    config = load_qq_live_config(config_path, require_group_ids=False)
    group_profiles = build_qq_group_profiles(config)
    LOGGER = configure_logging(os.path.join(config["output_root"], "logs"), "freight-live")
    clock_status = CLOCK.sync(config.get("time_check_urls"))
    LOGGER.info("启动联网校时: %s", clock_status)
    runtime_status = RuntimeStatus(config["output_root"])
    runtime_status.update(clock=clock_status, connection="starting")
    restart_requested = threading.Event()
    websocket_holder = {"socket": None}

    def request_configuration_reload() -> None:
        runtime_status.update(
            connection="restarting",
            restart_stage="stopping_receive",
        )
        restart_requested.set()
        current_socket = websocket_holder.get("socket")
        if current_socket is not None:
            try:
                current_socket.close()
            except Exception:
                pass

    dashboard_server = None
    from freight_data import FreightDataService
    data_service = FreightDataService(group_profiles, runtime_status)
    if config["dashboard_enabled"]:
        configuration_manager = FreightConfigurationManager(
            config["config_path"],
            config["rules_file"],
            validate_callback=lambda: validate_qq_live_config_without_applying(
                config_path
            ),
            restart_callback=request_configuration_reload,
            logger=LOGGER,
        )
        dashboard_server = start_status_dashboard(
            runtime_status,
            config["dashboard_host"],
            config["dashboard_port"],
            LOGGER,
            configuration_manager,
            recent_data_provider=lambda limit: build_recent_data_snapshot(
                group_profiles,
                limit,
            ),
            data_service=data_service,
        )
    if not config["group_ids"]:
        runtime_status.update(connection="setup_required")
        LOGGER.info("尚未配置QQ群，等待用户在管理页面完成首次设置。")
        print("尚未配置QQ群，请打开管理页面完成首次设置。")
        while not restart_requested.is_set():
            try:
                restart_requested.wait(1)
            except KeyboardInterrupt:
                runtime_status.update(connection="stopped")
                if dashboard_server:
                    dashboard_server.shutdown()
                    dashboard_server.server_close()
                runtime_status.close()
                return False
        runtime_status.update(connection="restarting", restart_stage="reloading")
        if dashboard_server:
            dashboard_server.shutdown()
            dashboard_server.server_close()
        runtime_status.close()
        return True
    ocr = WindowsOcr(config["ocr_script"], config["ocr_enabled"])
    ingestor = OneBotFreightIngestor(
        group_profiles,
        accept_self_messages=config["accept_self_messages"],
        ocr=ocr,
        logger=LOGGER,
    )
    excel_coordinator = ExcelRebuildCoordinator(
        group_profiles,
        runtime_status,
        logger=LOGGER,
        batch_delay_seconds=config["excel_batch_seconds"],
        full_refresh_seconds=config["excel_full_refresh_seconds"],
    )
    data_service.refresh_callback = excel_coordinator.request
    data_service.archive.resume()

    def handle_processed_event(_event: object, result: dict) -> None:
        if result["status"] in {"ignored", "duplicate"}:
            return
        log_result = {
            key: value
            for key, value in result.items()
            if not key.startswith("_")
        }
        if result["status"] == "added":
            print(
                f"收到群 {result['group_name']} ({result['group_id']}) 的运价信息，"
                f"发布人: {result['sender']}，新增: {result['inserted_count']}，"
                f"缓存新增: {result['pending_count']}"
            )
            LOGGER.info("运价消息: %s", log_result)
        elif result["status"] == "duplicate_content":
            LOGGER.info("内容查重已拦截: %s", log_result)
        else:
            LOGGER.info("消息进入未识别队列: %s", log_result)
        excel_coordinator.request(
            result.get("group_id", ""),
            records=result.get("_inserted_records", []),
        )

    message_processor = GroupMessageProcessor(
        group_profiles,
        ingestor,
        runtime_status,
        result_callback=handle_processed_event,
        logger=LOGGER,
        collection_control=data_service.collection,
    )
    data_service.collection.pending_count = lambda: sum(
        database.count_pending_inbound_events() for database in message_processor._databases.values()
    )

    print("QQ 实时提取已启动")
    print(f"监听群数量: {len(config['group_ids'])}")
    print(f"OneBot 地址: {config['ws_url']}")
    print(f"独立输出根目录: {config['output_root']}")
    print(
        f"程序可信时间: {CLOCK.now():%Y-%m-%d %H:%M:%S} "
        f"({clock_status['method']})"
    )
    if dashboard_server:
        print(
            f"状态面板: http://{config['dashboard_host']}:"
            f"{config['dashboard_port']}"
        )
    for profile in group_profiles.values():
        if profile.get("output_dir_note"):
            LOGGER.warning(profile["output_dir_note"])
        print(
            f"- {profile['group_name']} ({profile['group_id']}) -> "
            f"{profile['output_dir']}，默认始发地: {profile['default_origin']}"
        )
    print(("支持文字和Windows图片OCR" if ocr.enabled else "仅采集文字，图片已跳过")
          + "；未识别文字会单独保存；按 Ctrl+C 停止。")

    if config["rebuild_on_start"]:
        for profile in group_profiles.values():
            os.makedirs(profile["output_dir"], exist_ok=True)
            if not os.path.exists(profile["input_file"]):
                with open(profile["input_file"], "a", encoding="utf-8"):
                    pass
            summary = rebuild_qq_group_output(profile)
            runtime_status.update_group(
                profile["group_id"],
                name=profile["group_name"],
                output_dir=profile["output_dir"],
                **summary,
            )

    # Resume unfinished report updates after a deletion even when startup rebuild is disabled.
    for group_id in group_profiles:
        if data_service.report_state(group_id)["pending"]:
            excel_coordinator.request(group_id, force_full=True)

    lifecycle_manager = DataLifecycleManager(
        group_profiles,
        runtime_status,
        config["data_lifecycle"],
        refresh_callback=excel_coordinator.request,
        logger=LOGGER,
    )

    last_calendar_date = trusted_today()
    reconnect_count = 0

    stopped_by_user = False
    while not restart_requested.is_set():
        ws = None
        try:
            headers = []
            if config["access_token"]:
                headers.append(f"Authorization: Bearer {config['access_token']}")

            ws = websocket.create_connection(
                config["ws_url"],
                header=headers,
                timeout=15
            )
            websocket_holder["socket"] = ws
            ws.settimeout(config["heartbeat_seconds"])
            print("已连接 NapCat，等待群消息...")
            LOGGER.info("已连接NapCat: %s", config["ws_url"])
            runtime_status.update(
                connection="connected",
                websocket_connected_at=CLOCK.now().isoformat(timespec="seconds"),
                websocket_reconnect_count=reconnect_count,
                websocket_last_error="",
            )
            runtime_status.clear_errors()

            while not restart_requested.is_set():
                current_calendar_date = trusted_today()
                if current_calendar_date != last_calendar_date:
                    print(
                        f"检测到日期变化: {last_calendar_date} -> {current_calendar_date}，"
                        "开始转入到期缓存。"
                    )
                    for profile in group_profiles.values():
                        excel_coordinator.request(
                            profile["group_id"],
                            force_full=True,
                        )
                    last_calendar_date = current_calendar_date

                try:
                    raw_event = ws.recv()
                except websocket.WebSocketTimeoutException:
                    ws.ping("freight-heartbeat")
                    runtime_status.update(
                        websocket_last_ping_at=CLOCK.now().isoformat(timespec="seconds"),
                    )
                    continue

                if not raw_event:
                    raise ConnectionError("NapCat 已关闭连接。")

                try:
                    event = json.loads(raw_event)
                except (json.JSONDecodeError, TypeError):
                    continue

                runtime_status.update(
                    websocket_last_event_at=CLOCK.now().isoformat(timespec="seconds"),
                )

                message_processor.submit(event)

        except KeyboardInterrupt:
            print("\nQQ 实时提取已停止。")
            runtime_status.update(connection="stopped")
            stopped_by_user = True
            break
        except Exception as exc:
            if restart_requested.is_set():
                break
            reconnect_count += 1
            print(
                f"QQ 实时连接异常: {exc}；"
                f"{config['reconnect_seconds']:g} 秒后重连。"
            )
            LOGGER.exception("QQ实时连接异常")
            runtime_status.update(
                connection="disconnected",
                websocket_reconnect_count=reconnect_count,
                websocket_last_disconnect_at=CLOCK.now().isoformat(timespec="seconds"),
                websocket_last_error=str(exc)[:300],
            )
            runtime_status.add_error(str(exc))
            if ws is not None:
                try:
                    ws.close()
                except Exception:
                    pass
            if restart_requested.wait(config["reconnect_seconds"]):
                break
        finally:
            if websocket_holder.get("socket") is ws:
                websocket_holder["socket"] = None
            if ws is not None:
                try:
                    ws.close()
                except Exception:
                    pass

    runtime_status.update(
        connection="stopped" if stopped_by_user else "restarting",
        restart_stage="draining_queues" if not stopped_by_user else "",
    )
    if not lifecycle_manager.stop(timeout=30):
        runtime_status.update(connection="restart_failed")
        raise RuntimeError("数据生命周期线程无法安全停止，已中止热重载。")
    if not data_service.daily.wait_idle(timeout=30):
        raise RuntimeError('每日图表仍在生成，不能安全重新加载程序。')
    message_stopped = message_processor.stop(drain=True, timeout=60)
    if not message_stopped:
        runtime_status.update(
            connection="restarting",
            restart_stage="waiting_message_queue",
        )
        LOGGER.warning("消息队列60秒内未排空，禁止启动新会话并继续等待。")
        if not message_processor.wait_idle(300):
            runtime_status.update(connection="restart_failed")
            raise RuntimeError("消息队列超过360秒仍未排空，已中止热重载。")
        if not message_processor.stop(drain=False, timeout=30):
            raise RuntimeError("消息工作线程无法安全停止，已中止热重载。")

    excel_stopped = excel_coordinator.stop(drain=True, timeout=60)
    if not excel_stopped:
        runtime_status.update(
            connection="restarting",
            restart_stage="waiting_excel_queue",
        )
        LOGGER.warning("Excel队列60秒内未排空，禁止启动新会话并继续等待。")
        if not excel_coordinator.wait_idle(300):
            runtime_status.update(connection="restart_failed")
            raise RuntimeError("Excel队列超过360秒仍未排空，已中止热重载。")
        if not excel_coordinator.stop(drain=False, timeout=30):
            raise RuntimeError("Excel工作线程无法安全停止，已中止热重载。")
    if dashboard_server:
        dashboard_server.shutdown()
        dashboard_server.server_close()
    runtime_status.close()
    if stopped_by_user:
        return False
    runtime_status.update(connection="restarting", restart_stage="reloading")
    return True


def run_qq_live(config_path: str) -> None:
    """持续运行实时采集器；配置保存后在同一进程内完成优雅重载。"""
    while _run_qq_live_session(config_path):
        pass


def list_qq_groups(config_path: str) -> None:
    """通过 OneBot API 列出当前 QQ 账号加入的群，便于填写允许列表。"""
    try:
        import websocket
    except ImportError as exc:
        raise RuntimeError(
            "缺少 websocket-client，请按项目文档安装 requirements.txt。"
        ) from exc

    config = load_qq_live_config(config_path, require_group_ids=False)
    headers = []
    if config["access_token"]:
        headers.append(f"Authorization: Bearer {config['access_token']}")

    ws = None
    try:
        ws = websocket.create_connection(config["ws_url"], header=headers, timeout=15)
        echo = f"freight-group-list-{int(time.time())}"
        request = {"action": "get_group_list", "params": {}, "echo": echo}
        ws.send(json.dumps(request, ensure_ascii=False))

        while True:
            response = json.loads(ws.recv())
            if response.get("echo") != echo:
                continue
            if response.get("status") != "ok":
                raise RuntimeError(f"NapCat 获取群列表失败: {response.get('message', '未知错误')}")

            groups = response.get("data") or []
            print(f"共找到 {len(groups)} 个群：")
            for group in groups:
                print(f"{group.get('group_id')}\t{group.get('group_name', '')}")
            return
    except RuntimeError:
        raise
    except Exception as exc:
        raise RuntimeError(f"无法连接 NapCat: {exc}") from exc
    finally:
        if ws is not None:
            ws.close()


# =========================
# 22. 监控 txt 变化
# =========================

def watch_file(file_path):
    if not os.path.exists(file_path):
        print(f"未找到 {file_path}，请先创建该文件。")
        return

    last_mtime = os.path.getmtime(file_path)
    print(f"开始监控文件：{file_path}")
    print("当 txt 内容修改并保存后，会自动更新Excel历史日均价、周期统计和折线图。")
    print("按 Ctrl + C 可停止程序。")

    rebuild_all()
    create_daily_backup(
        os.getcwd(),
        "",
        extra_files=[INPUT_FILE, REPORT_XLSX, RULES_CONFIG_FILE],
        retention_days=30,
        backup_date=trusted_today(),
    )
    last_calendar_date = trusted_today()

    while True:
        try:
            time.sleep(2)
            current_mtime = os.path.getmtime(file_path)
            current_calendar_date = trusted_today()
            if current_mtime != last_mtime or current_calendar_date != last_calendar_date:
                last_mtime = current_mtime
                rebuild_all()
                create_daily_backup(
                    os.getcwd(),
                    "",
                    extra_files=[INPUT_FILE, REPORT_XLSX, RULES_CONFIG_FILE],
                    retention_days=30,
                    backup_date=trusted_today(),
                )
                last_calendar_date = current_calendar_date
        except KeyboardInterrupt:
            print("\n已停止监控。")
            break
        except Exception as e:
            print(f"监控时出错: {e}")
            time.sleep(2)


# =========================
# 23. 主程序
# =========================

def main() -> None:
    global LOGGER
    parser = argparse.ArgumentParser(description="物流运价统计与 QQ 实时提取")
    parser.add_argument(
        "--mode",
        choices=["file-watch", "file-once", "qq-groups", "qq-live"],
        default="file-watch",
        help="file-watch=监听文本；file-once=统计一次；qq-groups=列出群；qq-live=实时接入"
    )
    parser.add_argument(
        "--config",
        default=QQ_LIVE_CONFIG_FILE,
        help="qq-live 模式使用的 JSON 配置文件"
    )
    parser.add_argument(
        "--data-dir",
        default="",
        help="手动文本模式的数据输出目录"
    )
    args = parser.parse_args()

    app_dir = application_dir()
    config_path = args.config
    if not os.path.isabs(config_path):
        config_path = os.path.join(app_dir, config_path)

    if args.mode not in {"qq-groups", "qq-live"}:
        data_dir = os.path.abspath(args.data_dir or os.path.join(app_dir, "手动文本统计"))
        os.makedirs(data_dir, exist_ok=True)
        os.chdir(data_dir)
        input_path = os.path.join(data_dir, INPUT_FILE)
        if not os.path.exists(input_path):
            with open(input_path, "a", encoding="utf-8"):
                pass
        load_rules_config(os.path.join(app_dir, RULES_CONFIG_FILE))
        LOGGER = configure_logging(os.path.join(data_dir, "logs"), "freight-manual")
        clock_status = CLOCK.sync()
        LOGGER.info("手动模式启动联网校时: %s", clock_status)

    if args.mode == "file-once":
        rebuild_all()
        create_daily_backup(
            data_dir,
            "",
            extra_files=[INPUT_FILE, REPORT_XLSX, RULES_CONFIG_FILE],
            retention_days=30,
            backup_date=trusted_today(),
        )
    elif args.mode in {"qq-groups", "qq-live"}:
        try:
            if args.mode == "qq-groups":
                list_qq_groups(config_path)
            else:
                run_qq_live(config_path)
        except (FileNotFoundError, ValueError, RuntimeError, json.JSONDecodeError) as exc:
            print(f"无法启动 QQ 实时提取：{exc}")
            raise SystemExit(1) from None
    else:
        watch_file(INPUT_FILE)


if __name__ == "__main__":
    main()

